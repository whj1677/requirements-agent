"""Word-only layout and adopted-label presentation; source objects stay intact."""
import copy
import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

WORD_ACCENT='57514A'
WORD_BODY_FONT='DengXian'
WORD_TABLE_HEADER='EAE7E2'
WORD_TABLE_ROW='FAF9F7'
WORD_FLOW_PALETTE=dict(line='#827C73',edge_text='#57514A',decision='#EEEAE3',
                       node='#F6F4F1',border='#898278',text='#222222')


def prose_parts(text):
    """Break dense prose at complete sentences, retaining every original character.

    Semicolons keep their shared conditions; quoted UI messages, decimal values
    and parenthesized explanations must not become separate requirement points.
    """
    text=str(text)
    if len(text)<220:return [text]
    closing={'“':'”','‘':'’','「':'」','『':'』','（':'）','(':')','[':']','【':'】'}
    stack=[];parts=[];start=0
    for index,char in enumerate(text):
        if char=='"':
            if stack and stack[-1]=='"':stack.pop()
            else:stack.append('"')
        elif char in closing:stack.append(closing[char])
        elif stack and char==stack[-1]:stack.pop()
        if char in '。！？' and not stack:
            parts.append(text[start:index+1]);start=index+1
    if start<len(text):parts.append(text[start:])
    if parts and not parts[-1].strip() and len(parts)>1:
        trailing=parts.pop();parts[-1]+=trailing
    return parts or [text]


def _clause_numbering(doc):
    """One native Word list per requirement, restarted without typed-in numbers."""
    root=doc.part.numbering_part.element
    abstract_id=max((int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))),default=-1)+1
    number_id=max((int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))),default=0)+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(abstract_id))
    level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
    for name,value in (('start','1'),('numFmt','decimal'),('lvlText','%1.'),('suff','tab'),('lvlJc','left')):
        element=OxmlElement('w:'+name);element.set(qn('w:val'),value);level.append(element)
    properties=OxmlElement('w:pPr');tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab')
    tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'420');tabs.append(tab);properties.append(tabs)
    indent=OxmlElement('w:ind');indent.set(qn('w:left'),'420');indent.set(qn('w:hanging'),'420');properties.append(indent)
    level.append(properties);abstract.append(level)
    first_num=root.find(qn('w:num'))
    if first_num is not None:first_num.addprevious(abstract)
    else:root.append(abstract)
    number=OxmlElement('w:num');number.set(qn('w:numId'),str(number_id))
    reference=OxmlElement('w:abstractNumId');reference.set(qn('w:val'),str(abstract_id));number.append(reference);root.append(number)
    return number_id


def add_body_prose(doc,node):
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    text=node['text']
    readable=node.get('block_kind') in ('requirement','rule','acceptance','behavior','narrative')
    parts=prose_parts(text) if readable else [text]
    if len(parts)==1:
        paragraph=doc.add_paragraph(style='Normal')
        _add_prose_runs(paragraph,text)
        if readable:
            paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent=Pt(22)
        return paragraph
    if 'Requirement Clause' not in doc.styles:
        style=doc.styles.add_style('Requirement Clause',WD_STYLE_TYPE.PARAGRAPH)
        style.base_style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11)
        style.font.bold=False;style.font.cs_bold=False;style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),WORD_BODY_FONT)
        fmt=style.paragraph_format;fmt.line_spacing=1.5;fmt.space_after=Pt(8)
        fmt.left_indent=Pt(21);fmt.first_line_indent=Pt(-21)
        fmt.widow_control=True;fmt.keep_with_next=False
    number_id=_clause_numbering(doc)
    for part in parts:
        paragraph=doc.add_paragraph(style='Requirement Clause');paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        numbering=paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val=0;numbering.get_or_add_numId().val=number_id
        _add_prose_runs(paragraph,part)
    return paragraph


def _add_prose_runs(paragraph,text):
    lead=re.match(r'^([^\n：:。！？；]{2,16}[：:])',text)
    values=[(lead[0],True),(text[len(lead[0]):],False)] if lead else [(text,False)]
    for value,bold in values:
        run=paragraph.add_run(value);run.bold=bold;run.font.cs_bold=bold


def cell_prose(cell,text):
    parts=prose_parts(text)
    cell.text=parts[0]
    for part in parts[1:]:cell.add_paragraph(part)


def clean_style_gallery(doc):
    """Keep used styles and their dependencies, with a small editing gallery."""
    from lxml import etree
    styles=doc.styles.element
    definitions={s.get(qn('w:styleId')):s for s in styles.findall(qn('w:style'))}
    needed={'Normal','Title'}|{sid for sid,s in definitions.items() if s.get(qn('w:default'))=='1'}
    used_paragraphs=set()
    for part in doc.part.package.parts:
        if str(part.partname) in ('/word/styles.xml','/word/stylesWithEffects.xml'):continue
        element=getattr(part,'element',None)
        if element is None:continue
        for name in ('pStyle','rStyle','tblStyle'):
            for ref in element.iter(qn('w:'+name)):
                needed.add(ref.get(qn('w:val')))
                if name=='pStyle':used_paragraphs.add(ref.get(qn('w:val')))
    while True:
        previous=needed.copy()
        for sid in previous:
            style=definitions.get(sid)
            if style is not None:
                for name in ('basedOn','next','link'):
                    for ref in style.findall(qn('w:'+name)):needed.add(ref.get(qn('w:val')))
        if previous==needed:break
    gallery={'Normal','Title'}|{sid for sid in used_paragraphs if sid=='RequirementClause' or sid.startswith('Heading')}
    def flag(element,name,value):
        child=element.find(qn('w:'+name))
        if child is None:child=OxmlElement('w:'+name);element.append(child)
        child.set(qn('w:val'),'1' if value else '0')
    def prune(root):
        for style in list(root.findall(qn('w:style'))):
            sid=style.get(qn('w:styleId'))
            if sid not in needed:root.remove(style);continue
            flag(style,'qFormat',sid in gallery)
            flag(style,'semiHidden',sid not in gallery)
            flag(style,'unhideWhenUsed',sid in gallery)
        latent=root.find(qn('w:latentStyles'))
        if latent is not None:
            for key,value in (('defQFormat','0'),('defSemiHidden','1'),('defUnhideWhenUsed','1')):latent.set(qn('w:'+key),value)
            for exception in list(latent):latent.remove(exception)
    prune(styles)
    for part in doc.part.package.parts:
        if str(part.partname)=='/word/stylesWithEffects.xml':
            root=etree.fromstring(part.blob);prune(root)
            part._blob=etree.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    return sorted(gallery&needed)


def adopted_text(text):
    """Remove editorial prefixes, never modal verbs inside a business rule."""
    prefix=r'(?:AI\s*)?(?:建议|推荐)(?:方案)?\s*[：:]\s*'
    value=re.sub(r'(?m)^(\s*)(?:'+prefix+r')+',r'\1',str(text))
    # These labels are emitted by the document renderer, not user quotations.
    labels=('现有产品','相关页面或模块','本次改动意图','当前页面与流程','目标使用者',
            '问题与改动价值','本期新增或修改','保持不变','暂不涉及','优先级与依据',
            '谁操作','入口及前置条件','主要流程','业务数据','权限约束','操作结果','边界和异常',
            '类型','来源','必填条件','默认值','约束','枚举选项','可见范围','编辑权限',
            '补充权限','操作者','入口','前置条件','触发动作','输入','校验','处理中反馈',
            '重复操作','取消或离开','成功结果','失败结果','数据影响')
    value=re.sub(r'(?m)^(\s*(?:'+'|'.join(labels)+r')\s*[：:]\s*)(?:'+prefix+r')+',r'\1',value)
    # Priority rationale is an editorial field; "建议用户..." in a business rule
    # remains a prescribed recommendation to the end user.
    return re.sub(r'(?m)^(\s*优先级与依据\s*[：:]\s*)建议(?:为)?',r'\1',value)


def reading_node(node):
    value=copy.deepcopy(node)
    if value.get('block_kind') in ('requirement','rule','acceptance','behavior','narrative'):
        value['text']=adopted_text(value['text'])
    if value.get('kind')=='table' and (value.get('table_role')=='acceptance' or value.get('detail_hash')):
        value['rows']=[[adopted_text(cell) for cell in row] for row in value['rows']]
    return value


def add_navigation(doc,outline):
    if not any(n['level']==1 for n in outline):return {}
    from docx.enum.style import WD_STYLE_TYPE
    style=doc.styles.add_style('Document Navigation',WD_STYLE_TYPE.PARAGRAPH)
    style.base_style=doc.styles['Normal'];style.font.name='Calibri'
    style.font.size=Pt(11);style.font.color.rgb=RGBColor.from_string(WORD_ACCENT)
    style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.2
    title=doc.add_paragraph('目录','Document Table Caption')
    title.paragraph_format.space_before=Pt(18);title.paragraph_format.space_after=Pt(10)
    anchors={}
    for index,node in enumerate(n for n in outline if n['level']==1):
        name='ReqSection'+str(index+1);anchors[node['number']]=(name,index+1)
        paragraph=doc.add_paragraph(style='Document Navigation')
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),name);link.set(qn('w:history'),'1')
        run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
        font=OxmlElement('w:rFonts')
        for key,val in (('ascii','Calibri'),('hAnsi','Calibri'),('eastAsia','Microsoft YaHei')):font.set(qn('w:'+key),val)
        properties.append(font);color=OxmlElement('w:color');color.set(qn('w:val'),WORD_ACCENT);properties.append(color)
        run.append(properties);text=OxmlElement('w:t');text.text=node['number']+'  '+node['text'];run.append(text);link.append(run);paragraph._p.append(link)
    doc.add_page_break()
    return anchors


def bookmark(paragraph,identity):
    name,number=identity
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0,start);paragraph._p.append(end)


def running_furniture(doc,name,kind,has_cover):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    first=doc.sections[0]
    first.different_first_page_header_footer=has_cover
    header=first.header.paragraphs[0]
    header.paragraph_format.space_after=Pt(3)
    run=header.add_run(name+'  '+kind.upper())
    run.font.name='Calibri';run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0);run.bold=False
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    footer=first.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before=Pt(3)
    footer.paragraph_format.line_spacing=1
    def label(text):
        run=footer.add_run(text);run.font.name='Calibri';run.font.size=Pt(9);run.bold=False;run.font.cs_bold=False
        run.font.color.rgb=RGBColor.from_string('68645E');run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    label('需求内容评审  ·  第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);label(' 页')
    return footer
