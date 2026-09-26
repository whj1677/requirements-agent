"""OOXML text/table readers; no Office launch, formulas or external connections."""
import io
import zipfile
from .core import Problem, require


def check_package(data):
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        require(sum(x.file_size for x in package.infolist()) <= 40 * 1024 * 1024,
                'SOURCE_UNSAFE', 'Office 文件展开大小超限')
        require(not any('vbaproject' in x.lower() or '/embeddings/' in x.lower() for x in package.namelist()),
                'SOURCE_UNSAFE', '当前读取方式不处理宏或嵌入可执行对象；请提供不含这些对象的材料')
        return package.namelist()


def parse_workbook(data):
    from openpyxl import load_workbook
    names = check_package(data)
    rows, limits = [], []
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    total_cells = 0
    try:
        for sheet in workbook.worksheets:
            count = (sheet.max_row or 0) * (sheet.max_column or 0)
            total_cells += count
            require(total_cells <= 100000, 'SOURCE_LIMIT', '工作簿超过十万个单元格，请拆分材料')
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = str(cell.value)
                    if cell.data_type == 'f':
                        value = '公式（未计算）：' + value
                        limits.append('公式原文已保留；未计算结果或更新外链')
                    rows.append((f'工作表「{sheet.title}」 / {cell.coordinate}', value))
        if any('/drawings/' in x or '/charts/' in x or 'comments' in x.lower() for x in names):
            limits.append('图片、图表、批注和其他绘图对象未提取')
        limits.append('只提取单元格原始值和公式；格式、合并单元格布局及页面视觉未核验')
    finally:
        workbook.close()
    return rows, 'partial', '；'.join(dict.fromkeys(limits)), None


def parse_presentation(data):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    check_package(data)
    document = Presentation(io.BytesIO(data))
    require(len(document.slides) <= 200, 'SOURCE_LIMIT', '演示文稿超过 200 页，请拆分材料')
    rows, limits = [], []

    def shapes(items, location):
        for n, shape in enumerate(items, 1):
            loc = f'{location} / 对象 {n}'
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes(shape.shapes, loc)
            elif shape.has_table:
                for r, row in enumerate(shape.table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        if cell.text.strip():
                            rows.append((f'{loc} / 表格 R{r}C{c}', cell.text))
            elif shape.has_text_frame and shape.text.strip():
                rows.append((loc, shape.text))
            else:
                limits.append('图片、图表和其他非文本对象未提取')

    for n, slide in enumerate(document.slides, 1):
        shapes(slide.shapes, f'幻灯片 {n}')
    limits.append('已提取幻灯片文字与表格；备注、动画及页面视觉未核验')
    return rows, 'partial', '；'.join(dict.fromkeys(limits)), None
