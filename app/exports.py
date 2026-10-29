import hashlib
import io
import os
import re
import tempfile
from pathlib import Path
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .core import brief_hash, digest, dumps, ident, now, require
from .preview import prototype
from .document_schema import profile
from .document_reader import reader_document
from .requirements import delivery_items
from .word_presentation import reading_node,add_navigation,bookmark,running_furniture,add_body_prose,cell_prose,WORD_ACCENT,WORD_FLOW_PALETTE,WORD_BODY_FONT,WORD_TABLE_HEADER,WORD_TABLE_ROW,clean_style_gallery


def block_text(block, p):
    items = {i['id']:i for i in p['items']}
    if block['kind'] in ('requirement','rule','acceptance'):
        return '\n'.join(f'{r} — {items[r]["statement"]}' for r in block['ref_ids'])
    return ('【建议／待确认】' if block['kind'] == 'ui_suggestion' else '') + (block['text'] or '')


def readable_picture_parts(data):
    """Paginate tall screenshots at quiet scanlines, preserving full image coverage."""
    from PIL import Image
    with Image.open(io.BytesIO(data)) as image:
        height_limit=int(image.width*1.15)
        if image.height<=height_limit:return [data]
        gray=image.convert('L');parts=[];top=0
        while top<image.height:
            bottom=min(top+height_limit,image.height)
            if bottom<image.height:
                start=top+int(height_limit*.8)
                # Avoid bisecting text when a whitespace row is available nearby.
                bottom=min(range(start,bottom+1),key=lambda y:(sum(gray.crop((0,y,image.width,y+1)).histogram()[:110]),-y))
            out=io.BytesIO();image.crop((0,top,image.width,bottom)).save(out,format='PNG')
            parts.append(out.getvalue());top=bottom
        return parts


def add_full_flow_page(doc, picture, title):
    """Place one complete flow on a landscape page and restore body geometry."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image

    geometry_keys=('orientation','page_width','page_height','top_margin','bottom_margin',
                   'left_margin','right_margin','header_distance','footer_distance','gutter')
    geometry={key:getattr(doc.sections[-1],key) for key in geometry_keys}
    # A diagram chapter normally starts with its heading. Move that contiguous
    # heading group with the image so a section break cannot strand its title.
    headings=[]
    for child in reversed(list(doc._body._body)[:-1]):
        if child.tag!=qn('w:p'):
            break
        paragraph=next((p for p in doc.paragraphs if p._p is child),None)
        if paragraph is None or not (paragraph.style.name.startswith('Heading') or
                paragraph.style.name=='Document Metadata' and paragraph.paragraph_format.keep_with_next):
            break
        headings.insert(0,paragraph)
    for paragraph in headings:
        paragraph._p.getparent().remove(paragraph._p)

    def new_page_section():
        section=doc.add_section(WD_SECTION.NEW_PAGE)
        # The paragraph carrying sectPr has no visible content. Keep it tiny so
        # it does not push a fitted landscape figure onto an additional page.
        paragraph=doc.paragraphs[-1]
        paragraph.paragraph_format.space_before=Pt(0)
        paragraph.paragraph_format.space_after=Pt(0)
        paragraph.paragraph_format.line_spacing=Pt(1)
        paragraph.paragraph_format.keep_with_next=False
        return section

    landscape=new_page_section()
    landscape.different_first_page_header_footer=False
    landscape.orientation=WD_ORIENT.LANDSCAPE
    landscape.page_width=max(geometry['page_width'],geometry['page_height'])
    landscape.page_height=min(geometry['page_width'],geometry['page_height'])
    for paragraph in headings:
        doc._body._body.insert_element_before(paragraph._p,'w:sectPr')
    if headings:
        headings[0].paragraph_format.space_before=Pt(0)

    width=(landscape.page_width-landscape.left_margin-landscape.right_margin)/Inches(1)
    height=(landscape.page_height-landscape.top_margin-landscape.bottom_margin)/Inches(1)
    # Reserve room for headings, the caption, the inline-image baseline and the
    # closing section-break paragraph; never crop or split the image to fit.
    height-=.7+.42*len(headings)
    with Image.open(io.BytesIO(picture)) as image:
        picture_width=min(width,height*image.width/image.height)
    doc.add_picture(io.BytesIO(picture),width=Inches(picture_width))
    paragraph=doc.paragraphs[-1]
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before=Pt(0)
    paragraph.paragraph_format.space_after=Pt(3)
    paragraph.paragraph_format.line_spacing=1
    paragraph.paragraph_format.keep_with_next=True
    caption=doc.add_paragraph(title+' · 完整流程图','Document Metadata')
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before=Pt(0)
    caption.paragraph_format.space_after=Pt(0)
    caption.paragraph_format.keep_with_next=False

    restored=new_page_section()
    restored.different_first_page_header_footer=False
    for key,value in geometry.items():
        setattr(restored,key,value)


async def capture(spec):
    from playwright.async_api import async_playwright
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        page = await browser.new_page(viewport={'width':1100,'height':780}, device_scale_factor=1)
        images = []
        try:
            for selected in spec['pages']:
                await page.set_content(prototype(dict(spec, pages=[selected])), wait_until='load')
                images.append((selected['page_id'], await page.screenshot(full_page=True)))
        finally:
            await browser.close()
    return images


async def document_files(p, kind, status='草稿／待产品经理内容确认', *, reader=True, mode='formal'):
    require(kind in p['documents'], 'NOT_FOUND', '请先生成此类型文档', 404)
    artifact = p['documents'][kind]
    require(mode in ('formal','review'),'SCHEMA_INVALID','文档导出模式无效。')
    require(artifact['brief_hash'] == brief_hash(p), 'STALE_REVISION', '底稿已改变，请重新成文', 409)
    require(kind not in p.get('stale_document_kinds', []), 'STALE_REVISION', '页面已改变，请重新生成此文档后导出', 409)
    structured = kind=='prd' and any(item.get('detail_contract') or item.get('detail_required') for item in delivery_items(p))
    review_warnings=[]
    if structured:
        from .document_reader import export_readiness
        readiness=export_readiness(p,kind,mode=mode)
        require(readiness['ready'],'DOCUMENT_CONTENT_REQUIRED',
                '；'.join(issue['message'] for issue in readiness['issues']),409)
        review_warnings=readiness.get('warnings',[])
    if mode=='review':
        status='需求评审稿／非正式交付／含待确认项'
    content = reader_document(artifact) if reader else artifact['content']
    # Keep machine-level source/profile mapping available to export internals
    # without placing its workflow explanations in the product reading view.
    reference_mapping = artifact['content'].get('reference_mapping', []) if reader else content.get('reference_mapping', [])
    meta = f'{status} · 底稿版本 {artifact["draft_revision"]} · 内容哈希 {digest(content)[:16]}'
    if reader:
        meta=f'{status} · 文档 v{artifact.get("document_version","旧版未记录")} · 底稿 v{artifact["draft_revision"]}'
    rendered = [(s, [b['text'] if reader else block_text(b,p) for b in s['blocks']]) for s in content['sections']]
    sketches_included=artifact.get('sketch_policy')!='excluded'
    md = [f'# {content["title"]}', meta, '本文用于需求内容评审，不代表业务系统已经实现。' if not sketches_included else '本文用于需求内容评审；页面及数据为原型模拟，不能证明业务系统已经实现。']
    if not reader and content['content_profile_id'].startswith('builtin-'):
        md[2]+=' 已明确选择内置章节回退，未按用户原始 Word 参考生成。'
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    if structured and reader:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    table_width = (section.page_width-section.left_margin-section.right_margin)/Inches(1)
    for style_name in ('Normal','Title',*[f'Heading {n}' for n in range(1,10)]):
        st = doc.styles[style_name]
        st.font.name = 'Calibri' if style_name=='Normal' else 'Microsoft YaHei'
        st.font.color.rgb = RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),WORD_BODY_FONT if style_name=='Normal' else 'Microsoft YaHei')
        for border in st.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
        for attr in ('asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme'):
            st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
        if style_name == 'Normal':
            st.font.size = Pt(11)
            st.font.bold = False
            st.font.cs_bold = False
            st.paragraph_format.space_after = Pt(8)
            st.paragraph_format.line_spacing = 1.5
            st.paragraph_format.widow_control = True
        elif style_name == 'Title':
            st.font.size = Pt(28)
            st.font.bold = True
            st.paragraph_format.space_after = Pt(12)
            st.paragraph_format.keep_with_next = True
        else:
            depth = int(style_name.split()[-1])
            st.font.size = Pt({1:16,2:14}.get(depth,12))
            st.font.bold = True
            st.paragraph_format.space_before = Pt(18 if depth==1 else 12)
            st.paragraph_format.space_after = Pt(8 if depth==1 else 6)
            st.paragraph_format.keep_with_next = True
    from docx.enum.style import WD_STYLE_TYPE
    metadata_style = doc.styles.add_style('Document Metadata', WD_STYLE_TYPE.PARAGRAPH)
    metadata_style.base_style = doc.styles['Normal']
    metadata_style.font.size = Pt(9)
    metadata_style.font.color.rgb = RGBColor.from_string('68645E')
    metadata_style.paragraph_format.space_after = Pt(7)
    cover_metadata_style = doc.styles.add_style('Document Cover Metadata', WD_STYLE_TYPE.PARAGRAPH)
    cover_metadata_style.base_style = doc.styles['Normal']
    cover_metadata_style.font.size = Pt(10)
    cover_metadata_style.font.color.rgb = RGBColor.from_string('57514A')
    cover_metadata_style.paragraph_format.line_spacing = 1.4
    cover_metadata_style.paragraph_format.space_after = Pt(10)
    table_caption_style = doc.styles.add_style('Document Table Caption', WD_STYLE_TYPE.PARAGRAPH)
    table_caption_style.base_style = doc.styles['Normal']
    table_caption_style.font.bold=True
    table_caption_style.font.size=Pt(10)
    table_caption_style.font.color.rgb=RGBColor(0,0,0)
    table_caption_style.paragraph_format.keep_with_next=True
    table_caption_style.paragraph_format.space_before=Pt(8)
    table_caption_style.paragraph_format.space_after=Pt(5)
    if structured and reader:
        from .document_reader import document_name
        doc.add_paragraph(document_name(artifact), 'Title')
        cover=doc.add_paragraph(style='Document Cover Metadata')
        cover.add_run('产品需求文档 PRD\n').font.color.rgb=RGBColor.from_string(WORD_ACCENT)
        cover.add_run(meta.replace(' · ','\n'))
    else:
        doc.add_paragraph(content['title'], 'Title')
        doc.add_paragraph(meta.replace(' · ', '\n'), 'Document Cover Metadata')
    doc.add_paragraph(md[2])
    if mode=='review':
        notice='本文件仅供需求评审，不能作为已确认的研发交付依据。'
        doc.add_paragraph(notice);md.append(notice)
        if structured and reader and review_warnings:
            line='本版保留待确认内容，请结合文末「未决产品事项」与「评审核对清单」审核。'
            doc.add_paragraph(line);md.append(line)
            question_ids={b.get('question_id') for s in content['sections'] for b in s['blocks'] if b.get('question_id')}
            labels={v['item_id']:v['label'] for v in content.get('reference_labels',[])}
            property_names={'enum_values':'枚举选项','condition':'生效条件','validation':'校验规则',
                            'inputs':'输入','constraints':'约束','required':'必填条件',
                            'empty_value':'空值处理','default_value':'默认值','time_basis':'时间口径'}
            rows=[]
            for warning in review_warnings:
                if warning.get('question_id') in question_ids:
                    continue  # The complete question, impact and state are already in the reader.
                message=warning['message']
                for raw,label in {**labels,**property_names}.items():
                    message=re.sub(r'(?<![\w-])'+re.escape(raw)+r'(?![\w-])',lambda _:label,message)
                rows.append([f'P{len(rows)+1:02d}',message])
            if rows:
                last_number=max((int(n['number']) for n in content['outline'] if n['level']==1),default=0)+1
                warning_section=dict(section_id='REVIEW-CHECKS',title='评审核对清单',reading_nodes=[
                    dict(kind='heading',number=str(last_number),text='评审核对清单',level=1),
                    dict(kind='table',title='待核对定义',text='待核对定义',columns=['编号','核对内容'],
                         rows=rows,column_weights=[.10,.90],table_role='review_warnings')])
                rendered.append((warning_section,[]))
        else:
            for warning in review_warnings:
                line='待核对：'+warning['message']
                doc.add_paragraph(line);md.append(line)
    ui_current = bool(sketches_included and p.get('ui') and p['ui']['brief_hash'] == brief_hash(p))
    images = await capture(p['ui']['spec']) if ui_current else []
    image_map = dict(images)
    files, bindings = {}, []
    pictured={}
    if sketches_included and not ui_current:
        notice = '未提供当前底稿可用的页面方案；本文未插入原型图片。' if not p.get('ui') else '页面方案对应旧底稿；本文未插入过期原型图片。'
        md.append(notice)
        doc.add_paragraph(notice)
    # Review-only chapters are added after the canonical outline is compiled.
    # Include their actual headings too, so every visible chapter is reachable.
    navigation=[n for section,_ in rendered for n in section.get('reading_nodes',[])
                if n['kind']=='heading'] if reader else []
    anchors=add_navigation(doc,navigation) if reader else {}
    if anchors:doc.sections[0].different_first_page_header_footer=True
    for s, texts in rendered:
        if reader:
            for raw_node in s['reading_nodes']:
                node=reading_node(raw_node)
                if node['kind']=='heading':
                    title=node['number']+' '+node['text']
                    md.append('#'*(node['level']+1)+' '+title)
                    heading=doc.add_heading(title, node['level'])
                    if node['level']==1 and node['number'] in anchors:bookmark(heading,anchors[node['number']])
                elif node['kind']=='table':
                    from .detail_rendering import table_markdown
                    if node['title']!=s.get('title'):
                        doc.add_paragraph(node['title'],'Document Table Caption')
                        md.append('**'+node['title']+'**')
                    md.append(table_markdown(node))
                    table=doc.add_table(rows=1,cols=len(node['columns']))
                    table.style='Table Grid'
                    table.autofit=False
                    weights=(node.get('column_weights') or {'字段属性':[.19,.30,.51] if len(node['columns'])==3 else [.14,.19,.18,.30,.19],
                              '流程步骤':[.10,.15,.35,.11,.29],
                              '分支与流转条件':[.09,.27,.27,.37],
                              '状态变化':[.08,.22,.22,.26,.22]}.get(node['title'])
                             or ([.12,.44,.44] if len(node['columns'])==3 and node.get('detail_hash') else None)
                             or ([.22,.78] if len(node['columns'])==2 else [1/len(node['columns'])]*len(node['columns'])))
                    for column,weight in zip(table.columns,weights):
                        column.width=Inches(table_width*weight)
                    borders=OxmlElement('w:tblBorders')
                    for edge in ('top','left','bottom','right','insideH','insideV'):
                        border=OxmlElement('w:'+edge)
                        for key,value in (('val','single'),('sz','3'),('color','D9D9D9')):
                            border.set(qn('w:'+key),value)
                        borders.append(border)
                    table._tbl.tblPr.append(borders)
                    for cell,value in zip(table.rows[0].cells,node['columns']):
                        cell.text=str(value)
                    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
                    for values in node['rows']:
                        row=table.add_row()
                        for cell,value in zip(row.cells,values):
                            cell_prose(cell,str(value))
                    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
                    for row_number,row in enumerate(table.rows):
                        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
                        for cell,weight in zip(row.cells,weights):
                            cell.width=Inches(table_width*weight)
                            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                            props=cell._tc.get_or_add_tcPr()
                            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),WORD_TABLE_HEADER if row_number==0 else WORD_TABLE_ROW if row_number%2 else 'FFFFFF');props.append(shade)
                            margins=OxmlElement('w:tcMar')
                            for edge in ('top','bottom','left','right'):
                                margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'100' if edge in ('top','bottom') else '120');margin.set(qn('w:type'),'dxa');margins.append(margin)
                            props.append(margins)
                            for paragraph in cell.paragraphs:
                                paragraph.paragraph_format.space_after=Pt(3)
                                paragraph.paragraph_format.line_spacing=1.25
                                paragraph.paragraph_format.keep_with_next=row_number==0
                                for run in paragraph.runs:
                                    run.font.size=Pt(10)
                                    run.font.bold=row_number==0
                                    run.font.color.rgb=RGBColor.from_string('222222')
                                    run.font.cs_bold=row_number==0
                                    run.font.name='Calibri'
                                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),WORD_BODY_FONT)
                    bindings.append(dict(section_id=s['section_id'],item_id=node.get('item_id'),
                                         item_ids=node.get('item_ids',[]),detail_hash=node.get('detail_hash'),kind='table',title=node['title']))
                    spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(3)
                elif node['kind']=='diagram':
                    from .detail_rendering import flow_png,flow_svg
                    try:
                        picture=flow_png(node['flow'],palette=WORD_FLOW_PALETTE,fit_labels=True)
                    except (ValueError,OSError,KeyError,TypeError) as error:
                        require(False,'DETAIL_RENDER_FAILED','流程图导出失败：'+str(error),409)
                    asset='assets/flow-'+node['detail_hash'][:20]+'.png'
                    svg_asset=asset.removesuffix('.png')+'.svg'
                    files[asset]=picture;files[svg_asset]=flow_svg(node['flow'],node['title'],palette=WORD_FLOW_PALETTE,fit_labels=True)[0].encode('utf-8')
                    # Self-contained Markdown also works when downloaded alone.
                    import base64
                    md.append('!['+node['title'].replace('[','').replace(']','')+'](data:image/png;base64,'+base64.b64encode(picture).decode('ascii')+')')
                    add_full_flow_page(doc,picture,node['title'])
                    bindings.append(dict(section_id=s['section_id'],item_id=node['item_id'],
                                         detail_hash=node['detail_hash'],kind='diagram',
                                         asset=asset,asset_hash=hashlib.sha256(picture).hexdigest()))
                elif node['kind']=='detail_error':
                    require(False,'DETAIL_RENDER_FAILED',node['text'],409)
                else:
                    md.append(node['text'])
                    paragraph = doc.add_paragraph(node['text'],'Document Metadata') if node['kind']=='metadata' else add_body_prose(doc,node)
                    if node['kind']=='metadata' and node.get('item_id'):
                        # Keep the heading and full identifier with the actual clause.
                        paragraph.paragraph_format.keep_with_next = True
        else:
            md.append('#'*(s['level']+1) + ' ' + s['title'])
            doc.add_heading(s['title'], s['level'])
            for text in texts:
                md.append(text)
                doc.add_paragraph(text)
        picture_dimension = 'PRD-4.F.1' if kind == 'prd' else 'MRD-5.1.F.3'
        maps = [m for m in reference_mapping if s['section_id'] in m['output_section_ids'] and m['profile_section_id'] == picture_dimension and m['disposition'] in ('included','merged')]
        if maps and images:
            allowed_refs = {m['scope_ref'] for m in maps if m['scope_ref']}
            if not allowed_refs:
                continue
            for page in p['ui']['spec']['pages']:
                if not allowed_refs.intersection(page['requirement_refs']):
                    continue
                image = image_map[page['page_id']]
                filename = 'assets/' + hashlib.sha256(page['page_id'].encode()).hexdigest()[:16] + '.png'
                caption = f'需求草图，非最终UI设计 · {page["title"]} · 草图版本 {p["ui"]["spec"]["draft_revision"]}'
                from PIL import Image
                if reader and page['page_id'] in pictured:
                    reference='本功能使用同一需求草图，参见「'+pictured[page['page_id']]+'」中的图示。'
                    doc.add_paragraph(reference);md.append(reference)
                else:
                    parts=readable_picture_parts(image) if reader else [image]
                    for n,part in enumerate(parts,1):
                        asset=filename if len(parts)==1 else filename.removesuffix('.png')+f'-{n}.png'
                        part_caption=caption+(f'（{n}/{len(parts)}）' if len(parts)>1 else '')
                        with Image.open(io.BytesIO(part)) as preview:
                            width=min(6.4,8.2*preview.width/preview.height)
                        doc.add_picture(io.BytesIO(part), width=Inches(width))
                        doc.paragraphs[-1].paragraph_format.keep_with_next = True
                        doc.add_paragraph(part_caption)
                        md.append(f'![{part_caption}]({asset})')
                        files[asset]=part
                    pictured[page['page_id']]=s['title']
                files[filename] = image
                bindings.append(dict(section_id=s['section_id'], ui_page_id=page['page_id'], ui_spec_hash=digest(p['ui']['spec']), asset_hash=hashlib.sha256(image).hexdigest(), caption=caption))
    if not reader:
        doc.add_heading('参考维度处置与待确认事项', 1)
        md.append('## 参考维度处置与待确认事项')
        labels={d['id']:d['reference_heading'] for d in profile(kind,content['content_profile_id'].startswith('builtin-'))['sections']}
        groups={}
        states={'pending':'待确认','not_applicable':'不适用','merged':'合并表达'}
        for m in reference_mapping:
            if m['disposition'] in states:
                key=(m['disposition'],m['reason'],m['scope_ref'])
                groups.setdefault(key,[]).append(labels.get(m['profile_section_id'],m['profile_section_id']))
        for (state,reason,scope),dimensions in groups.items():
            line=f'{states[state]}'+(f' · {scope}' if scope else '')+'：'+'；'.join(dimensions)+'。\n'+reason
            doc.add_paragraph(line)
            md.append(line)
        for q in p['questions']:
            if q['status'] != 'answered':
                line = f'待确认 {q["id"]}：{q["question"]}；影响：{q["why"]}'
                doc.add_paragraph(line)
                md.append(line)
    from .document_reader import document_name
    footer=running_furniture(doc,document_name(artifact),kind,bool(anchors))
    if structured and reader:
        # Explicit body runs keep local Word theme/template defaults from
        # turning ordinary Chinese prose into a second heading weight.
        for paragraph in [*doc.paragraphs,footer]:
            if paragraph.style.name in ('Normal','Document Metadata','Document Cover Metadata'):
                for run in paragraph.runs:
                    run.bold=run.bold is True
                    run.font.cs_bold=run.bold
                    run.font.name='Calibri'
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),WORD_BODY_FONT)
    clean_style_gallery(doc)
    out = io.BytesIO(); doc.save(out)
    files[kind.upper()+'.docx'] = out.getvalue()
    files[kind.upper()+'.md'] = ('\n\n'.join(md)+'\n').encode('utf-8')
    files['document_asset_bindings.json'] = dumps(bindings).encode('utf-8')
    files['reference_mapping.json'] = dumps(reference_mapping).encode('utf-8')
    if content.get('reference_labels'):
        files[kind.upper()+'.references.json']=dumps(dict(document_id=artifact['id'],
            content_hash=digest(artifact['content']),references=content['reference_labels'])).encode('utf-8')
    return files


def zip_files(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def handoff(store, pid, baseline_id, idempotency_key):
    # Capture one read snapshot; rendering and compression must not retain a
    # SQLite write transaction or block unrelated workbench edits.
    with store.connect() as db:
        db.execute('BEGIN')
        prior = [r for r in store.records(pid,'export',db) if r['idempotency_key'] == idempotency_key]
        if prior:
            require(prior[0]['baseline_id'] == baseline_id, 'IDEMPOTENCY_CONFLICT', '导出幂等键冲突', 409)
            return prior[0]
        baseline = next((b for b in store.records(pid,'baseline',db) if b['id'] == baseline_id),None)
        require(baseline is not None, 'CONFIRMATION_REQUIRED', '需要有效人工确认基线', 409)
        changes = store.records(pid,'requirement_change',db)
    baseline_identity = digest(baseline)
    files, manifest = handoff_files(baseline, changes)
    archive = zip_files(files)
    eid = ident('EXPORT')
    folder = store.folder / 'exports'
    folder.mkdir(exist_ok=True)
    destination = folder / (eid + '.zip')
    temporary = None
    published = committed = False
    try:
        # Only a committed export record grants download access. The temporary
        # path is never exposed, and rename publishes a complete archive at once.
        with tempfile.NamedTemporaryFile(dir=folder, prefix='.handoff-', suffix='.tmp', delete=False) as out:
            temporary = out.name
            out.write(archive)
            out.flush()
            os.fsync(out.fileno())
        with store.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            prior = [r for r in store.records(pid,'export',db) if r['idempotency_key'] == idempotency_key]
            if prior:
                require(prior[0]['baseline_id'] == baseline_id, 'IDEMPOTENCY_CONFLICT', '导出幂等键冲突', 409)
                return prior[0]
            current = next((b for b in store.records(pid,'baseline',db) if b['id'] == baseline_id),None)
            require(current is not None and digest(current) == baseline_identity,
                    'STALE_BASELINE', '确认基线在导出期间发生变化，请重新核对后导出', 409)
            os.replace(temporary, destination)
            published = True
            result = dict(baseline_id=baseline_id, confirmation_id=baseline['confirmation_id'],
                          idempotency_key=idempotency_key, manifest=manifest)
            store.record(pid,'export',result,eid,db)
        committed = True
        return dict(id=eid,**result)
    finally:
        if temporary is not None:
            Path(temporary).unlink(missing_ok=True)
        if published and not committed:
            destination.unlink(missing_ok=True)


def handoff_files(baseline, changes):
    baseline_id = baseline['id']
    p = baseline['project']
    items = delivery_items(p)
    canonical = '\n\n'.join(f'## {i["id"]}｜{i["title"]} · v{i.get("content_version",1)}\n\n{i["statement"]}' for i in items)
    refs = {i['id']: i for i in items}
    traces = []
    for req in [i for i in items if i['kind']=='requirement']:
        related = [i for i in items if i['id'] in req['related_refs'] or req['id'] in i['related_refs']]
        traces.append(dict(requirement_id=req['id'], rule_ids=[i['id'] for i in related if i['kind']=='rule'], ac_ids=[i['id'] for i in related if i['kind']=='acceptance'], ui_page_ids=[page['page_id'] for page in p['ui']['spec']['pages'] if req['id'] in page['requirement_refs']] if p['ui'] else [], frontend_sections=[req['id']], backend_sections=[req['id']], open_question_ids=[q['id'] for q in p['questions'] if req['id'] in q['related_refs'] and q['status']!='answered']))
    files = {
        'README.md': f'# 研发交接\n\n基线 {baseline_id}\n\n产品经理内容确认记录 {baseline["confirmation_id"]}。内容确认不代表技术设计或测试验收通过。\n\n本包为中性交接；ai-engineer-context 私有格式适配未验证。',
        'requirements.md': '# 共同需求\n\n'+canonical,
        'frontend_spec.md': '# 前端交接\n\n以下为同一基线原文；页面结构与模拟交互见 ui_spec.json / prototype.html。接口及技术实现由开发确认。\n\n'+canonical,
        'backend_spec.md': '# 后端交接\n\n以下为同一基线原文；不得以界面隐藏代替权限规则。接口路径、表结构和事务方案未由本包代定。\n\n'+canonical,
        'acceptance_criteria.md': '# 验收条件\n\n'+'\n\n'.join(f'{i["id"]} — {i["statement"]}' for i in items if i['kind']=='acceptance'),
        'open_questions.md': '# 未决事项\n\n'+'\n\n'.join(q['question']+'\n影响：'+q['why'] for q in p['questions'] if q['status']!='answered'),
        'ui_spec.json': dumps(p['ui']['spec'] if p['ui'] else None),
        'prototype.html': prototype(p['ui']['spec']) if p['ui'] else '<!doctype html><meta charset="utf-8"><p>本基线未包含 UI 方案</p>',
        'traceability.json': dumps(traces)}
    from .requirements import exchange, requirement_ref
    from .delivery_scope import formal_scope, SCOPE_SCHEMA
    files['delivery_scope.json'] = dumps(formal_scope(baseline))
    files['delivery_scope.schema.json'] = dumps(SCOPE_SCHEMA)
    files['README.md'] += '\n\n正式机器范围见 delivery_scope.json（requirements-delivery-scope-1），按 entities 的完整身份/版本/内容哈希读取，包含 constraints；旧 requirements-exchange-1 的 requirements 数组保持历史兼容，不等于正式范围。此附加契约尚未验证外部消费者。'
    payload=exchange(p,baseline_id)
    payload['change_records']=[r for r in changes if any(
        i['id']==r['requirement_id'] and r['version']<=i.get('content_version',1) for i in p['items'])]
    files['requirements-exchange.json']=dumps(payload)
    for kind,artifact in p['documents'].items():
        content=reader_document(artifact)
        files[kind.upper()+'.content.json']=dumps(content)
        files[kind.upper()+'.md']='# '+content['title']+'\n\n'+'\n\n'.join(
            '## '+s['title']+'\n\n'+'\n\n'.join(b['text'] for b in s['blocks']) for s in content['sections'])
    for role,keys in (('frontend',('actor','entry','flow','result','exceptions')),('backend',('data','permissions','result','exceptions'))):
        from .product_flow import BEHAVIOR
        details=[]
        for req in (i for i in items if i['kind']=='requirement'):
            details.append('\n\n### '+req['id']+' · v'+str(req.get('content_version',1))+' 的相关行为\n'+
                '\n'.join(BEHAVIOR[k]+'：'+req.get('behavior',{}).get(k,'未指定，需技术设计时核对') for k in keys))
        files[role+'_spec.md']+=''.join(details)
    for trace in traces:
        trace['requirement']=requirement_ref(p,refs[trace['requirement_id']])
        trace['test_case_refs']=[]
        trace['test_execution_status']='not_run'
    files['traceability.json']=dumps(traces)
    files = {k:v.encode('utf-8') for k,v in files.items()}
    manifest = dict(schema_version='1.2', delivery_scope_version='requirements-delivery-scope-1', baseline_id=baseline_id, baseline_hash=digest(baseline['hashes']), confirmation_id=baseline['confirmation_id'], generator_version='1.1', prompt_version='1.1', created=now(), artifacts={k:dict(sha256=hashlib.sha256(v).hexdigest(), role=k) for k,v in files.items()})
    files['manifest.json'] = dumps(manifest).encode()
    return files, manifest
