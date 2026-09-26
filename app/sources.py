import asyncio
import base64
import hashlib
import http.client
import io
import ipaddress
import re
import socket
import ssl
import zipfile
import subprocess
import sys
from pathlib import Path
from urllib.parse import urljoin, urlsplit
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pypdf import PdfReader
from .core import ROOT, Problem, ident, now, require
from .office import OFFICE_EXTENSIONS, extract_office
from .office_formats import check_package, parse_workbook, parse_presentation
from .material_images import word_images, image_mime

MAX_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000
MAX_EXCERPTS = 20_000
SOURCE_EXTENSIONS = ['.txt', '.md', '.docx', '.doc', '.xlsx', '.xls', '.pptx', '.ppt', '.pdf', '.png', '.jpg', '.jpeg', '.webp']


def validate_extraction(rows, status):
    """Apply the same text bounds to parsers, Office and supplied web text."""
    require(isinstance(rows, (list, tuple)) and all(
        isinstance(row, (list, tuple)) and len(row) == 2 and
        all(isinstance(value, str) for value in row) for row in rows),
        'SOURCE_FAILED', '提取结果格式无效')
    require(status in ('read', 'partial', 'awaiting_vision', 'office_required', 'permission_denied'),
            'SOURCE_FAILED', '读取状态无效')
    if status in ('read', 'partial', 'awaiting_vision'):
        require(any(value.strip() for _, value in rows), 'SOURCE_FAILED', '没有可读取内容；请提供可提取文本或截图')
    require(len(rows) <= MAX_EXCERPTS and sum(len(value) for _, value in rows) <= MAX_TEXT_CHARS,
            'SOURCE_LIMIT', '提取内容超过上限，请拆分材料')
    require(not any('\ufffd' in value or '\x00' in value for _, value in rows),
            'SOURCE_FAILED', '提取文本含无法识别的字符，尚不能标记为已读取')
    return rows


def public_target(url):
    u = urlsplit(url)
    require(u.scheme in ('http', 'https') and u.hostname and not u.username and not u.password,
            'SOURCE_FAILED', '仅支持无凭据的公开 HTTP(S) 网页')
    require(u.port in (None, 80, 443), 'SOURCE_FAILED', '网页端口不允许')
    try:
        addresses = list(dict.fromkeys(x[4][0] for x in socket.getaddrinfo(u.hostname, u.port or (443 if u.scheme == 'https' else 80), type=socket.SOCK_STREAM)))
    except OSError:
        raise Problem('SOURCE_FAILED', '域名解析失败')
    require(addresses and all(ipaddress.ip_address(a).is_global for a in addresses), 'SOURCE_FAILED', '拒绝回环、私网、保留地址及元数据地址')
    return u, addresses[0]


class PinnedHTTPS(http.client.HTTPSConnection):
    def __init__(self, host, ip, port, timeout):
        super().__init__(host, port=port, timeout=timeout, context=ssl.create_default_context())
        self.ip = ip

    def connect(self):
        raw = socket.create_connection((self.ip, self.port), self.timeout)
        self.sock = self._context.wrap_socket(raw, server_hostname=self.host)


def fetch_public(url):
    # Connect to the address that was validated, retaining TLS hostname checks.
    for _ in range(5):
        u, ip = public_target(url)
        conn = PinnedHTTPS(u.hostname, ip, u.port or 443, 10) if u.scheme == 'https' else http.client.HTTPConnection(ip, port=u.port or 80, timeout=10)
        try:
            conn.request('GET', (u.path or '/') + ('?' + u.query if u.query else ''), headers={'Host': u.netloc, 'User-Agent': 'RequirementsAgent/1.1', 'Accept-Encoding': 'identity'})
            response = conn.getresponse()
            if response.status in (301, 302, 303, 307, 308):
                url = urljoin(url, response.getheader('Location', ''))
                continue
            require(200 <= response.status < 300, 'SOURCE_FAILED', '网页响应 HTTP ' + str(response.status))
            payload = response.read(MAX_BYTES + 1)
            require(len(payload) <= MAX_BYTES, 'SOURCE_FAILED', '网页超过读取上限')
            return payload, response.getheader('Content-Type', 'text/html'), url
        finally:
            conn.close()
    raise Problem('SOURCE_FAILED', '重定向次数超限')


async def webpage(url, dynamic=False):
    if not dynamic:
        raw, mime, final = await asyncio.wait_for(asyncio.to_thread(fetch_public, url), 30)
        soup = BeautifulSoup(raw, 'html.parser')
        for e in soup(['script', 'style', 'iframe', 'noscript']):
            e.decompose()
        return raw, soup.get_text('\n', strip=True), 'http-text', final
    from playwright.async_api import async_playwright
    blocked = []
    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True, args=['--disable-background-networking', '--disable-quic'])
        context = await browser.new_context(service_workers='block', accept_downloads=False)
        async def route(request_route):
            req = request_route.request
            if req.method != 'GET' or req.resource_type not in ('document', 'script', 'stylesheet', 'xhr', 'fetch'):
                blocked.append(req.resource_type)
                return await request_route.abort()
            try:
                body, mime, final = await asyncio.to_thread(fetch_public, req.url)
                # Fulfillment prevents the browser from re-resolving the remote hostname.
                await request_route.fulfill(status=200, body=body, headers={'Content-Type': mime, 'Content-Security-Policy': "connect-src 'self' https:; frame-src 'none'; worker-src 'none'; object-src 'none'"})
            except Exception:
                blocked.append(req.resource_type)
                await request_route.abort()
        await context.route('**/*', route)
        await context.route_web_socket('**/*', lambda ws: ws.close())
        page = await context.new_page()
        page.on('dialog', lambda d: d.dismiss())
        try:
            await page.goto(url, wait_until='networkidle', timeout=25000)
            text = await page.locator('body').inner_text(timeout=3000)
            raw = (await page.content()).encode()
            return raw, text, 'controlled-browser' + (';partial-resources' if blocked else ''), url
        finally:
            await browser.close()


def parse_bytes(title, data):
    require(0 < len(data) <= MAX_BYTES, 'SOURCE_FAILED', '文件为空或超过 12 MiB')
    ext = Path(title).suffix.lower()
    rows, status, reason, image_mime = [], 'read', '', None
    if ext in ('.txt', '.md'):
        text = data.decode('utf-8-sig')
        rows = [(f'行 {n}', t) for n, t in enumerate(text.splitlines(), 1) if t.strip()]
    elif ext == '.docx':
        check_package(data)
        doc = Document(io.BytesIO(data))
        rows = [(f'段落 {i}', p.text) for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
        rows += [(f'表格 {t} 行 {r}', ' | '.join(c.text for c in row.cells)) for t, table in enumerate(doc.tables, 1) for r, row in enumerate(table.rows, 1)]
        for section_number, section in enumerate(doc.sections, 1):
            for name, part in (('页眉', section.header), ('页脚', section.footer),
                               ('首页页眉', section.first_page_header), ('首页页脚', section.first_page_footer),
                               ('偶数页页眉', section.even_page_header), ('偶数页页脚', section.even_page_footer)):
                rows += [(f'节 {section_number} / {name} / 段落 {i}', p.text)
                         for i, p in enumerate(part.paragraphs, 1) if p.text.strip()]
                rows += [(f'节 {section_number} / {name} / 表格 {t} 行 {r}', ' | '.join(c.text for c in row.cells))
                         for t, table in enumerate(part.tables, 1) for r, row in enumerate(table.rows, 1)
                         if any(c.text.strip() for c in row.cells)]
        limits = []
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            names = package.namelist()
            if any(name.startswith('word/') and any(marker in name.lower() for marker in
                   ('comments', 'footnotes', 'endnotes')) for name in names):
                limits.append('批注、脚注或尾注对象未提取')
            for name in names:
                if name == 'word/document.xml' or re.fullmatch(r'word/(header|footer)\d+\.xml', name):
                    xml = package.read(name)
                    if re.search(rb'<w:(?:ins|del|moveFrom|moveTo)(?:\s|>)', xml):
                        limits.append('修订文字未单独核对')
                    if re.search(rb'<w:(?:txbxContent|object|pict|altChunk|hyperlink)(?:\s|>)', xml):
                        limits.append('文本框、嵌入对象或链接文字可能未提取')
                    if name == 'word/document.xml' and re.search(rb'<w:drawing(?:\s|>)', xml):
                        limits.append('正文绘图对象尚未视觉分析')
                    if name != 'word/document.xml' and re.search(rb'<w:drawing(?:\s|>)', xml):
                        limits.append('页眉页脚图片未提取')
        if doc.inline_shapes:
            limits.append('正文内嵌图片尚未视觉分析')
        status = 'partial' if limits else 'read'
        reason = '已提取正文、表格、页眉及页脚文字；页面视觉未核验'
        if limits:
            reason += '；' + '；'.join(dict.fromkeys(limits))
        if len(doc.inline_shapes):
            if not rows:rows=[('读取限制','文档没有可提取的正文文字，内嵌图片尚未视觉分析。')]
    elif ext == '.xlsx':
        rows, status, reason, image_mime = parse_workbook(data)
    elif ext == '.pptx':
        rows, status, reason, image_mime = parse_presentation(data)
    elif ext in ('.doc', '.xls', '.ppt'):
        raise Problem('OFFICE_REQUIRED', '旧版 Office 文件需要本机 Office 读取')
    elif ext == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        require(len(reader.pages) <= 200, 'SOURCE_FAILED', 'PDF 超过 200 页')
        for i, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ''
            except Exception:
                text = ''
            if text.strip():
                rows.append((f'第 {i} 页', text))
            else:
                status, reason = 'partial', reason + f'第 {i} 页未提取文字；'
    elif ext in ('.png', '.jpg', '.jpeg', '.webp'):
        image = Image.open(io.BytesIO(data))
        require(image.format in ('PNG', 'JPEG', 'WEBP') and image.width * image.height <= 25_000_000, 'SOURCE_FAILED', '图片格式或像素数不支持')
        image.verify()
        image_mime = {'PNG': 'image/png', 'JPEG': 'image/jpeg', 'WEBP': 'image/webp'}[image.format]
        rows = [('整图', '图片原始像素；尚未视觉分析')]
        status = 'awaiting_vision'
    else:
        raise Problem('SOURCE_FAILED', '支持 Word、Excel、PowerPoint、TXT、MD、PDF、PNG、JPEG、WebP')
    validate_extraction(rows, status)
    return rows, status, reason, image_mime


def parse_bounded(title,data):
    if Path(title).suffix.lower() in ('.txt','.md'):
        return parse_bytes(title,data)
    try:
        result=subprocess.run([sys.executable,'-m','app.parse_worker',Path(title).suffix.lower()],input=data,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=30,cwd=ROOT,creationflags=subprocess.CREATE_NO_WINDOW if sys.platform=='win32' else 0)
    except subprocess.TimeoutExpired:
        raise Problem('SOURCE_FAILED','文件解析超过 30 秒，解析子进程已终止')
    import json
    if result.returncode != 0:
        try:
            failure = json.loads(result.stdout.decode('utf8'))
        except (ValueError, UnicodeError):
            failure = {}
        raise Problem(failure.get('code','SOURCE_FAILED'), failure.get('message','当前读取方式无法解析；原件已保留，可重新读取'))
    return json.loads(result.stdout.decode('utf8'))


def save_source(store, title, data, purpose='current', uri=None, parsed=None):
    sid = ident('SRC')
    sha = hashlib.sha256(data).hexdigest()
    rawdir = store.folder / 'sources'
    rawdir.mkdir(exist_ok=True)
    path = rawdir / sid
    path.write_bytes(data)
    source = dict(id=sid, title=title, purpose=purpose, uri=uri, sha256=sha, created=now(), version=1, excluded=False, excerpts=[], parse_status='failed', failure_reason='', image_mime=None)
    try:
        reading = dict(method='standard-parser', visual_review='not_run', manual_open='not_verified')
        pictures = []
        try:
            rows, status, reason, mime = parsed or parse_bounded(title, data)
            if Path(title).suffix.lower()=='.docx':
                pictures, image_limits = word_images(data)
                if pictures or image_limits:
                    status='partial'
                    reason += f'；已提取{len(pictures)}张正文图片，待视觉分析'
                    if image_limits:
                        reason += '；' + '；'.join(image_limits)
        except Exception as error:
            if Path(title).suffix.lower() not in OFFICE_EXTENSIONS or (isinstance(error, Problem) and error.code in ('SOURCE_UNSAFE','SOURCE_LIMIT')):
                raise
            result = extract_office(title, data, rawdir)
            rows, status, reason, mime = result.get('rows', []), result['status'], result['reason'], None
            pictures=result.get('images',[])
            reading.update({key:value for key,value in result.items() if key not in ('rows','status','reason','images')})
            reading.update(method='local-office', standard_parser='failed')
        validate_extraction(rows, status)
        source['reading'] = reading
        children=[]
        for picture in pictures:
            try:
                raw=base64.b64decode(picture['data_base64'],validate=True)
                child_mime=image_mime(raw)
                child=save_source(store, title+' · '+picture['locator']+'.png',raw,purpose,
                    parsed=([(picture['locator'],'图片像素已提取，尚未视觉分析')],'awaiting_vision','',child_mime))
                child.update(container_source_id=sid,container_hash=sha,container_locator=picture['locator'])
                children.append(child)
            except (ValueError, OSError):
                reason+='；'+picture.get('locator','图片')+' 无法读取，未进入分析'
        if pictures:
            source['embedded_image_ids']=[c['id'] for c in children]
            source['_embedded_sources']=children
        source.update(parse_status=status, failure_reason=reason, image_mime=mime)
        for locator, text in rows:
            # Bound each excerpt while retaining exact source positions.
            for offset in range(0, len(text), 5000):
                source['excerpts'].append(dict(id=ident('EX'), source_id=sid, source_hash=sha, locator=f'{locator} / 字符 {offset}', text=text[offset:offset+5000], method='image' if mime else ('local-office' if reading['method']=='local-office' else 'extracted')))
    except Exception as e:
        source['parse_status']='failed'
        source['excerpts']=[]
        source['failure_reason'] = e.message if isinstance(e, Problem) else '当前读取方式无法解析；原件已保留'
    return source
