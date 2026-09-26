"""Windows customer entry: one browser workbench and an activation-only gateway."""
import argparse
import ctypes
import hashlib
import io
import json
import os
from pathlib import Path
import socket
import sqlite3
import subprocess
import sys
import time
import urllib.request
import webbrowser

URL = 'http://127.0.0.1:8765/'
PORT = 8765


def user_home():
    base = Path(os.environ['LOCALAPPDATA'])
    if not base.is_absolute():
        raise RuntimeError('LOCALAPPDATA 必须为绝对路径')
    return base / 'RequirementsAgent'


def restore_pipes():
    """Windowed PyInstaller leaves sys streams None, even for piped workers."""
    if os.name != 'nt':
        return
    import msvcrt
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.GetStdHandle.argtypes = [ctypes.c_ulong]
    kernel.GetStdHandle.restype = ctypes.c_void_p
    for name, code, flags, mode in [('stdin', -10, os.O_RDONLY, 'rb'),
                                    ('stdout', -11, os.O_WRONLY, 'wb'),
                                    ('stderr', -12, os.O_WRONLY, 'wb')]:
        if getattr(sys, name) is None:
            handle = kernel.GetStdHandle(code & 0xffffffff)
            if handle and handle != ctypes.c_void_p(-1).value:
                stream = os.fdopen(msvcrt.open_osfhandle(handle, flags | os.O_BINARY), mode)
                setattr(sys, name, io.TextIOWrapper(stream, encoding='utf-8'))


def report(result, destination=None):
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if destination:
        # Explicit operator-selected output only; no implicit customer data initialization.
        Path(destination).write_text(text + '\n', encoding='utf-8')
    elif sys.stdout is not None:
        print(text, flush=True)


def active_tasks():
    database = user_home() / 'data/requirements.sqlite3'
    if not database.exists():
        return []
    with sqlite3.connect(database.as_uri() + '?mode=ro', uri=True) as db:
        return [rid for rid, body in db.execute("SELECT id,payload FROM records WHERE kind IN ('run','user_task')")
                if json.loads(body).get('status') in ('queued', 'running')]


def identity():
    with open(sys.executable, 'rb') as stream:
        executable = hashlib.file_digest(stream, 'sha256').hexdigest()
    profile = str(user_home().resolve()).casefold()
    return hashlib.sha256((executable + profile).encode('utf-8')).hexdigest()[:24]


def runtime_at_port():
    try:
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(URL + 'activation/api/instance', timeout=2) as response:
            return json.load(response)
    except Exception:
        return {}


def port_in_use():
    with socket.socket() as sock:
        sock.settimeout(0.4)
        return sock.connect_ex(('127.0.0.1', PORT)) == 0


def kernel_api():
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateMutexW.argtypes = [ctypes.c_void_p, ctypes.c_bool, ctypes.c_wchar_p]
    kernel.CreateMutexW.restype = ctypes.c_void_p
    kernel.CloseHandle.argtypes = [ctypes.c_void_p]
    kernel.WaitForSingleObject.argtypes = [ctypes.c_void_p, ctypes.c_ulong]
    kernel.ReleaseMutex.argtypes = [ctypes.c_void_p]
    return kernel


def server():
    # An unlicensed process serves ONLY the activation UI; no app.main import.
    from .activation import ActivationGateway
    from .core import ROOT
    import uvicorn
    kernel = kernel_api()
    mutex = kernel.CreateMutexW(None, False, 'Local\\RequirementsAgent-Desktop-Runtime-8765')
    if not mutex or ctypes.get_last_error() == 183:
        if mutex:
            kernel.CloseHandle(mutex)
        raise RuntimeError('客户工作台已在运行，不会另开实例')
    listener = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
            listener.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        listener.bind(('127.0.0.1', PORT))
        listener.listen(128)
        logdir = user_home() / 'logs'
        logdir.mkdir(parents=True, exist_ok=True)
        log = (logdir / 'server.log').open('a', encoding='utf-8', buffering=1)
        sys.stdout = sys.stderr = log
        os.environ['RA_ACCESS_TOKEN'] = 'off'
        os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(ROOT / 'browsers')
        asset_dir = ROOT / ('web/activation' if getattr(sys, 'frozen', False) else 'web/src/activation')
        gateway = ActivationGateway(active_tasks=active_tasks, asset_dir=asset_dir, instance_id=identity())
        instance = uvicorn.Server(uvicorn.Config(gateway, host='127.0.0.1', port=PORT,
                                                access_log=False, log_config=None))
        gateway.shutdown = lambda: setattr(instance, 'should_exit', True)
        print('客户激活入口：' + URL + 'activation/', flush=True)
        instance.run(sockets=[listener])
    finally:
        listener.close()
        kernel.CloseHandle(mutex)


def launcher():
    kernel = kernel_api()
    mutex = kernel.CreateMutexW(None, False, 'Local\\RequirementsAgent-Desktop-Launch-8765')
    if not mutex:
        raise RuntimeError('无法取得工作台启动锁')
    acquired = False
    child = None
    try:
        acquired = kernel.WaitForSingleObject(mutex, 30000) in (0, 0x80)
        if not acquired:
            raise RuntimeError('另一个启动操作仍在进行，请稍后重试')
        instance_id = identity()
        if port_in_use():
            status = runtime_at_port()
            if status.get('instance_id') != instance_id:
                raise RuntimeError('8765 已被另一工作台或其他程序占用。请先正常退出原工作台；不会改用其他端口或结束未知进程。')
            webbrowser.open(URL + 'activation/')
            return 0
        command = [sys.executable, '--serve'] if getattr(sys, 'frozen', False) else [sys.executable, str(Path(__file__).parents[1] / 'scripts/desktop.py'), '--serve']
        child = subprocess.Popen(command, creationflags=subprocess.CREATE_NO_WINDOW,
                                 stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        deadline = time.monotonic() + 40
        while time.monotonic() < deadline:
            if child.poll() is not None:
                raise RuntimeError('激活入口启动失败，请查看用户目录 logs/server.log。')
            if runtime_at_port().get('instance_id') == instance_id:
                webbrowser.open(URL + 'activation/')
                return 0
            time.sleep(0.3)
        # This is our own startup child, never an unrelated process at the port.
        child.terminate()
        child.wait(timeout=10)
        raise RuntimeError('激活入口启动超时，请查看用户目录 logs/server.log。')
    finally:
        if acquired:
            kernel.ReleaseMutex(mutex)
        kernel.CloseHandle(mutex)


def parser_worker(extension):
    restore_pipes()
    from .licensing import LicenseManager
    LicenseManager().verify_cached()
    from .sources import parse_bytes, MAX_BYTES
    from .core import Problem
    payload = sys.stdin.buffer.read(MAX_BYTES + 1)
    try:
        result = parse_bytes('input' + extension, payload)
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))
        sys.stdout.buffer.flush()
        return 0
    except Exception as error:
        result = dict(code=error.code if isinstance(error, Problem) else 'SOURCE_FAILED',
                      message=error.message if isinstance(error, Problem) else '当前读取方式无法解析')
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))
        sys.stdout.buffer.flush()
        return 2


def smoke_test():
    """Explicit licensed offline diagnostic, on this customer's selected user profile."""
    from .licensing import LicenseManager
    LicenseManager().verify_cached()
    from .core import ROOT
    os.environ['PLAYWRIGHT_BROWSERS_PATH'] = str(ROOT / 'browsers')
    os.environ['RA_ACCESS_TOKEN'] = 'off'
    from .sources import parse_bounded
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image
    checks = {}
    word = Document()
    word.add_paragraph('打包后的中文材料读取验证')
    buf = io.BytesIO()
    word.save(buf)
    checks['docx_subprocess'] = '中文材料' in json.dumps(parse_bounded('test.docx', buf.getvalue()), ensure_ascii=False)
    book = Workbook()
    book.active.append(['打包验收', 123])
    buf = io.BytesIO()
    book.save(buf)
    checks['xlsx_subprocess'] = '打包验收' in json.dumps(parse_bounded('test.xlsx', buf.getvalue()), ensure_ascii=False)
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[0])
    slide.shapes.title.text = '安装包幻灯片验收'
    buf = io.BytesIO()
    slides.save(buf)
    checks['pptx_subprocess'] = '幻灯片验收' in json.dumps(parse_bounded('test.pptx', buf.getvalue()), ensure_ascii=False)
    buf = io.BytesIO()
    Image.new('RGB', (20, 20), 'white').save(buf, format='PNG')
    checks['image_subprocess'] = parse_bounded('test.png', buf.getvalue())[1] == 'awaiting_vision'
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_content('<html><body>独立 Chromium 验收</body></html>')
        checks['bundled_chromium'] = page.locator('body').inner_text() == '独立 Chromium 验收'
        browser.close()
    from fastapi.testclient import TestClient
    from .main import app
    with TestClient(app) as client:
        checks['session'] = client.get('/api/session').status_code == 200
        checks['frontend'] = client.get('/').status_code == 200
        checks['projects'] = client.get('/api/projects').status_code == 200
    return {'status': 'OFFLINE_SMOKE_VERIFIED' if all(checks.values()) else 'OFFLINE_SMOKE_FAILED',
            'checks': checks, 'model_calls': 0}


def main(argv=None):
    parser = argparse.ArgumentParser()
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--serve', action='store_true')
    modes.add_argument('--parse-worker', metavar='EXT')
    modes.add_argument('--check-license', action='store_true')
    modes.add_argument('--make-request', metavar='NAME')
    modes.add_argument('--install-license', metavar='FILE')
    modes.add_argument('--smoke-test', action='store_true')
    modes.add_argument('--check-idle', action='store_true')
    parser.add_argument('--replace-license', action='store_true')
    parser.add_argument('--output')
    args = parser.parse_args(argv)
    if args.check_idle:
        try:
            idle = not port_in_use() and not active_tasks()
            report({'status': 'IDLE' if idle else 'BUSY'}, args.output)
            return 0 if idle else 2
        except Exception:
            return 2
    from .licensing import LicenseManager, LicenseError
    try:
        if args.parse_worker:
            return parser_worker(args.parse_worker)
        if args.serve:
            server()
            return 0
        if args.make_request:
            report(LicenseManager().create_request(args.make_request), args.output)
        elif args.install_license:
            report(LicenseManager().install_license(args.install_license, replace=args.replace_license), args.output)
        elif args.check_license:
            payload = LicenseManager().verify_cached()
            report({'status': 'LICENSE_VALID', 'licenseId': payload['licenseId'], 'product': payload['product']}, args.output)
        elif args.smoke_test:
            result = smoke_test()
            report(result, args.output)
            return 0 if all(result['checks'].values()) else 1
        else:
            return launcher()
        return 0
    except Exception as error:
        code = error.code if isinstance(error, LicenseError) else 'DESKTOP_FAILED'
        message = getattr(error, 'message', str(error))
        report({'status': 'REJECTED', 'code': code, 'message': message}, args.output)
        if not any((args.serve, args.parse_worker, args.make_request, args.install_license,
                    args.check_license, args.smoke_test, args.output)):
            ctypes.windll.user32.MessageBoxW(None, message, '需求 Agent 无法启动', 0x10)
        return 1
