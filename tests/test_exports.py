import asyncio
import copy
import hashlib
import io
import json
import zipfile
from docx import Document
from app.store import Store
from app.core import hashes,brief_hash
from app.exports import document_files,handoff
from app.workflow import confirm
from tests.helpers import prepared

def test_md_docx_canonical_and_actual_embedded_image(tmp_path):
    store=Store(tmp_path);p=prepared(store)
    for kind in ('prd','mrd'):
        files=asyncio.run(document_files(p,kind))
        text='\n'.join(x.text for x in Document(io.BytesIO(files[kind.upper()+'.docx'])).paragraphs)
        md=files[kind.upper()+'.md'].decode()
        for i in p['items']:
            assert i['statement'] in text and i['statement'] in md
        assert '待确认' in text and '低保真模拟' in text
        with zipfile.ZipFile(io.BytesIO(files[kind.upper()+'.docx'])) as z:
            assert any(name.startswith('word/media/') for name in z.namelist())
        assert '800X600' not in text and '撤销订单' not in text


def test_mrd_picture_follows_layout_mapping_not_description(tmp_path):
    p=prepared(Store(tmp_path))
    content=p['documents']['mrd']['content']
    content['sections'].append(dict(section_id='MRD-02',level=1,parent_section_id=None,title='交互与功能说明',blocks=[]))
    for mapping in content['reference_mapping']:
        if mapping['profile_section_id']=='MRD-5.1.F.3':
            mapping.update(disposition='merged',output_section_ids=['MRD-02'],reason='页面布局合并到交互与功能说明')
        elif mapping['profile_section_id']=='MRD-5.1.F.4':
            mapping.update(disposition='included',output_section_ids=['MRD-01'],reason='功能描述独立表达')
    files=asyncio.run(document_files(p,'mrd'))
    bindings=json.loads(files['document_asset_bindings.json'])
    assert bindings and all(b['section_id']=='MRD-02' for b in bindings)
    md=files['MRD.md'].decode()
    assert '![' not in md.split('## 交互与功能说明')[0]
    assert '![' in md.split('## 交互与功能说明')[1].split('## 参考维度')[0]
    doc=Document(io.BytesIO(files['MRD.docx']))
    headings=[(i,x.text) for i,x in enumerate(doc.paragraphs) if x.style.name.startswith('Heading')]
    layout=next(i for i,t in headings if t=='交互与功能说明')
    appendix=next(i for i,t in headings if t=='参考维度处置与待确认事项')
    pictures=[i for i,x in enumerate(doc.paragraphs) if x._p.xpath('.//w:drawing')]
    assert pictures and all(layout<i<appendix for i in pictures)


def test_two_functions_keep_distinct_images_and_sections(tmp_path):
    p=prepared(Store(tmp_path))
    second=copy.deepcopy(p['items'][0])
    second.update(id='REQ-0002',title='电价查询',statement='用户可查询电价。',related_refs=[])
    p['items'].append(second)
    page=copy.deepcopy(p['ui']['spec']['pages'][0])
    page['page_id']='UI-0002'
    page['title']='电价查询'
    page['requirement_refs']=['REQ-0002']
    p['ui']['spec']['pages'].append(page)
    content=p['documents']['mrd']['content']
    content['sections'][0]['title']='时段维护页面'
    content['sections'].append(dict(section_id='MRD-02',level=1,parent_section_id=None,title='电价查询页面',blocks=[dict(kind='requirement',ref_ids=['REQ-0002'],text=None)]))
    content['coverage'].append(dict(item_id='REQ-0002',section_ids=['MRD-02']))
    for mapping in list(content['reference_mapping']):
        if mapping['scope_ref']=='REQ-0001':
            other=copy.deepcopy(mapping)
            other['scope_ref']='REQ-0002'
            other['output_section_ids']=['MRD-02']
            content['reference_mapping'].append(other)
    current=brief_hash(p)
    p['ui']['brief_hash']=current
    p['documents']['mrd']['brief_hash']=current
    files=asyncio.run(document_files(p,'mrd'))
    bindings=json.loads(files['document_asset_bindings.json'])
    assert {(b['section_id'],b['ui_page_id']) for b in bindings}=={('MRD-01','UI-0001'),('MRD-02','UI-0002')}
    assert len({b['asset_hash'] for b in bindings})==2
    for binding in bindings:
        filename='assets/'+hashlib.sha256(binding['ui_page_id'].encode()).hexdigest()[:16]+'.png'
        assert hashlib.sha256(files[filename]).hexdigest()==binding['asset_hash']
    md=files['MRD.md'].decode()
    first=md.split('## 时段维护页面')[1].split('## 电价查询页面')[0]
    second=md.split('## 电价查询页面')[1].split('## 参考维度')[0]
    assert 'UI-0001' not in md and first.count('![')==1 and second.count('![')==1
    assert '电价查询' not in first and '电价时段管理' not in second
    doc=Document(io.BytesIO(files['MRD.docx']))
    current_heading=None
    image_sections=[]
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            current_heading=paragraph.text
        if paragraph._p.xpath('.//w:drawing'):
            image_sections.append(current_heading)
    assert image_sections==['时段维护页面','电价查询页面']


def test_no_current_ui_and_cropped_picture_are_explicit(tmp_path):
    p=prepared(Store(tmp_path))
    for ui_state in ('missing','stale'):
        candidate=copy.deepcopy(p)
        if ui_state=='missing':candidate['ui']=None
        else:candidate['ui']['brief_hash']='older-brief'
        files=asyncio.run(document_files(candidate,'mrd'))
        assert json.loads(files['document_asset_bindings.json'])==[]
        assert '未插入' in files['MRD.md'].decode()
        with zipfile.ZipFile(io.BytesIO(files['MRD.docx'])) as z:
            assert not any(x.startswith('word/media/') for x in z.namelist())
    cropped=copy.deepcopy(p)
    mapping=next(m for m in cropped['documents']['prd']['content']['reference_mapping'] if m['profile_section_id']=='PRD-4.F.1')
    mapping.update(disposition='not_applicable',reason='此轮只评审非页面规则，页面原型不适用。')
    files=asyncio.run(document_files(cropped,'prd'))
    assert json.loads(files['document_asset_bindings.json'])==[]
    md=files['PRD.md'].decode()
    assert '不适用' in md and '页面原型不适用' in md
    assert '相接边界如何处理' in md and '待确认' in md

def test_handoff_canonical_and_manifest(tmp_path):
    import hashlib,json
    store=Store(tmp_path);p=prepared(store)
    c=confirm(store,p['id'],dict(expected_revision=p['revision'],expected_hashes=hashes(p),scope_ids=[i['id'] for i in p['items']],idempotency_key='export-test'))
    e=handoff(store,p['id'],c['baseline_id'],'export-test')
    with zipfile.ZipFile(store.folder/'exports'/(e['id']+'.zip')) as z:
        manifest=json.loads(z.read('manifest.json'))
        for path,record in manifest['artifacts'].items():assert hashlib.sha256(z.read(path)).hexdigest()==record['sha256']
        for path in ('requirements.md','frontend_spec.md','backend_spec.md'):
            text=z.read(path).decode()
            assert all(i['statement'] in text for i in p['items'])
