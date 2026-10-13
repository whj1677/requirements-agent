"""Real HTTP route checks over isolated Stores; zero model/provider requests."""
import io
import zipfile

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app.main import create_app
from app.requirement_confirmation import row_fingerprint
from tests.test_detail_document import project_fixture, compile_fixture
from tests.test_document_publication import login


def fixture_app(tmp_path, confirmed=False):
    project=project_fixture(tmp_path,confirmed=confirmed)
    app=create_app(tmp_path,access_token='synthetic',env_path=tmp_path/'.env')
    store=app.state.store
    with store.edit(project['id'],project['revision'],'isolated detailed requirement fixture',bump=False) as (current,_):
        current.update(project)
    # Normal item reconciliation versions the changed definition at commit.
    # Confirm and compile only after that real persistence boundary.
    current=store.get(project['id'])
    with store.edit(project['id'],current['revision'],'isolated detailed document fixture',bump=False) as (current,db):
        if confirmed:
            current['items'][0]['requirement_confirmation']['content_hash']=row_fingerprint(current,current['items'][0])
        compile_fixture(current)
        store.record(project['id'],'document_artifact',current['documents']['prd'],db=db)
    return app,project['id']


@pytest.mark.parametrize('fmt',['md','docx','zip'])
def test_review_download_is_labelled_and_formal_cannot_bypass_confirmation(tmp_path,fmt):
    app,pid=fixture_app(tmp_path)
    with TestClient(app) as client:
        login(client)
        payload=client.get(f'/api/projects/{pid}').json()
        detail_blockers=[row for row in payload['confirmation_blockers'] if row['code']=='detail_contract']
        assert detail_blockers and all(isinstance(row.get('detail'),str) for row in detail_blockers)
        url=f'/api/projects/{pid}/documents/prd/{fmt}'
        response=client.get(url+'?mode=formal')
        assert response.status_code==409
        assert '确认' in response.json()['message']
        assert not app.state.store.records(pid,'document_export')
        review=client.get(url+'?mode=review')
        assert review.status_code==200,review.text[:200]
        if fmt=='md':
            text=review.content.decode('utf-8')
        elif fmt=='docx':
            word=Document(io.BytesIO(review.content))
            assert word.tables and word.inline_shapes
            text='\n'.join([*(p.text for p in word.paragraphs),
                *(cell.text for table in word.tables for row in table.rows for cell in row.cells)])
        else:
            archive=zipfile.ZipFile(io.BytesIO(review.content))
            text=archive.read(next(name for name in archive.namelist() if name.endswith('.md'))).decode('utf-8')
            assert any(name.endswith('.png') for name in archive.namelist())
        assert '非正式交付' in text and '详细需求未确认' in text
        records=app.state.store.records(pid,'document_export')
        assert len(records)==1 and records[0]['export_mode']=='review'
        assert 'detail_hash' in records[0]['asset_bindings']


def test_formal_download_after_confirmation_contains_current_tables_and_figure(tmp_path):
    app,pid=fixture_app(tmp_path,confirmed=True)
    with TestClient(app) as client:
        login(client)
        snapshot=app.state.store.get(pid);artifact=snapshot['documents']['prd']
        current=client.get(f'/api/projects/{pid}').json()
        assert current['document_export_readiness']['prd']['ready']
        for blocker in current.get('confirmation_blockers',[]):
            assert not isinstance(blocker.get('detail'),dict)
        response=client.get(f'/api/projects/{pid}/documents/prd/docx?mode=formal&document_id='+artifact['id'])
        assert response.status_code==200,response.text[:200]
        document=Document(io.BytesIO(response.content))
        assert len(document.tables)==5 and len(document.inline_shapes)>=1
        assert '非正式交付' not in '\n'.join(p.text for p in document.paragraphs)
        assert app.state.store.records(pid,'document_export')[0]['export_mode']=='formal'


@pytest.mark.parametrize('mode',['review','formal'])
@pytest.mark.parametrize('fault',['stale','binding','snapshot','old_document'])
def test_both_http_modes_reject_stale_or_mismatched_document(tmp_path,mode,fault):
    app,pid=fixture_app(tmp_path,confirmed=True)
    store=app.state.store;p=store.get(pid)
    with store.edit(pid,p['revision'],'isolated stale fixture',bump=False) as (current,_):
        artifact=current['documents']['prd']
        if fault=='stale':
            current['items'][0]['detail_contract']['fields'][0]['constraints']='修改为 1 至 50 个字符'
        elif fault=='binding':
            next(b for s in artifact['content']['sections'] for b in s['blocks'] if b['kind']=='detail')['detail_hash']='0'*64
        elif fault=='snapshot':
            artifact['item_snapshot'][0]['detail_contract']['fields'][0]['constraints']='偷偷修改的快照'
    with TestClient(app) as client:
        login(client)
        response=client.get(f'/api/projects/{pid}/documents/prd/md?mode={mode}'+('&document_id=DOC-old' if fault=='old_document' else ''))
        assert response.status_code==409,response.text[:200]
        assert not store.records(pid,'document_export')


def test_http_export_rechecks_snapshot_before_publishing(tmp_path,monkeypatch):
    app,pid=fixture_app(tmp_path,confirmed=True)
    import app.main as main
    original=main.document_files
    async def changed_during_export(*args,**kwargs):
        files=await original(*args,**kwargs)
        current=app.state.store.get(pid)
        with app.state.store.edit(pid,current['revision'],'isolated concurrent edit') as (state,_):
            state['items'][0]['detail_contract']['fields'][0]['constraints']='导出期间修改范围'
        return files
    monkeypatch.setattr(main,'document_files',changed_during_export)
    with TestClient(app) as client:
        login(client)
        response=client.get(f'/api/projects/{pid}/documents/prd/md?mode=formal')
        assert response.status_code==409 and response.json()['code']=='STALE_DOCUMENT'
        assert not app.state.store.records(pid,'document_export')
