"""Synthetic coverage for source states and isolated backup restoration."""
import hashlib
import io
import json
import shutil
import subprocess
import sys
from pathlib import Path
import sqlite3
import pytest
from contextlib import closing
import zipfile

from docx import Document

from app.sources import MAX_TEXT_CHARS, save_source
from app.store import Store
from scripts import backup
from scripts.backup import create, verify


def test_docx_header_footer_are_extracted_with_positions(tmp_path):
    document = Document()
    document.add_paragraph('正文')
    document.sections[0].header.paragraphs[0].text = '页眉独有文字'
    document.sections[0].footer.paragraphs[0].text = '页脚独有文字'
    raw = io.BytesIO()
    document.save(raw)
    result = save_source(Store(tmp_path), 'synthetic.docx', raw.getvalue())
    text = '\n'.join(row['text'] for row in result['excerpts'])
    assert result['parse_status'] == 'read'
    assert '页眉独有文字' in text and '页脚独有文字' in text
    assert any('页眉' in row['locator'] for row in result['excerpts'])


def test_docx_unsupported_object_is_partial_with_reason(tmp_path):
    document = Document()
    document.add_paragraph('正文')
    raw = io.BytesIO()
    document.save(raw)
    package = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as original, zipfile.ZipFile(package, 'w') as changed:
        for item in original.infolist():
            changed.writestr(item, original.read(item.filename))
        changed.writestr('word/comments.xml', '<comments/>')
    result = save_source(Store(tmp_path), 'synthetic.docx', package.getvalue())
    assert result['parse_status'] == 'partial'
    assert '批注' in result['failure_reason']


def test_preparsed_web_text_uses_shared_validation(tmp_path):
    store = Store(tmp_path)
    for text in ('', 'x' * (MAX_TEXT_CHARS + 1), '\ufffd'):
        result = save_source(store, 'https://example.test', b'<html></html>',
                             parsed=([('http-text', text)], 'read', '', None))
        assert result['parse_status'] == 'failed' and not result['excerpts']


def test_backup_restores_source_and_model_evidence_with_hashes(tmp_path):
    data = tmp_path / 'daily'
    store = Store(data)
    project = store.create('合成项目')
    source = save_source(store, 'synthetic.txt', '合成正文'.encode())
    with store.edit(project['id'], project['revision'], '添加合成来源') as (working, _):
        working['sources'].append(source)
    call_id = 'CALL-synthetic'
    store.record(project['id'], 'run', {'status': 'completed', 'attempts': [{'call_id': call_id}]})
    evidence = data / 'evidence' / 'model-calls' / (call_id + '.json')
    evidence.parent.mkdir(parents=True)
    evidence.write_text(json.dumps({'call_id': call_id, 'request_input': '合成请求', 'final_output': '合成响应'}), 'utf-8')
    export = data / 'exports' / 'synthetic.txt'
    export.parent.mkdir(parents=True)
    export.write_text('合成导出', 'utf-8')
    (data / '.env').write_text('SECRET=synthetic', 'utf-8')
    backup = create(data, tmp_path / 'backup')
    restored = tmp_path / 'restored'
    shutil.copytree(backup, restored)
    manifest = verify(restored)
    assert (restored / 'sources' / source['id']).read_bytes() == '合成正文'.encode()
    assert json.loads((restored / 'evidence' / 'model-calls' / (call_id + '.json')).read_text('utf-8'))['final_output'] == '合成响应'
    assert (restored / 'exports' / 'synthetic.txt').read_text('utf-8') == '合成导出'
    assert not (restored / '.env').exists()
    assert manifest['files']['sources/' + source['id']]['sha256'] == hashlib.sha256('合成正文'.encode()).hexdigest()


def test_backup_refuses_active_run(tmp_path):
    store = Store(tmp_path / 'daily')
    project = store.create('合成项目')
    store.record(project['id'], 'run', {'status': 'running'})
    try:
        create(store.folder, tmp_path / 'backup')
    except RuntimeError as error:
        assert '活动任务' in str(error)
    else:
        raise AssertionError('Active run must prevent backup')
    assert not (tmp_path / 'backup').exists()


def test_backup_refuses_missing_recorded_source(tmp_path):
    data = tmp_path / 'daily'
    store = Store(data)
    project = store.create('合成项目')
    source = save_source(store, 'synthetic.txt', '合成正文'.encode())
    with store.edit(project['id'], project['revision'], '添加合成来源') as (working, _):
        working['sources'].append(source)
    (data / 'sources' / source['id']).unlink()
    target = tmp_path / 'backup'
    with pytest.raises(RuntimeError, match='来源文件'):
        create(data, target)
    assert not target.exists()


def test_stop_guard_checks_active_tasks_read_only(tmp_path):
    store = Store(tmp_path / 'daily')
    project = store.create('合成项目')
    rid = store.record(project['id'], 'run', {'status': 'running'})
    checker = Path(__file__).resolve().parents[1] / 'scripts' / 'check_active_tasks.py'
    command = [sys.executable, '-X', 'utf8', str(checker), str(store.path)]
    active = subprocess.run(command, capture_output=True, text=True, encoding='utf-8')
    assert active.returncode == 2 and rid in active.stdout
    store.update_run(project['id'], rid, status='completed')
    assert subprocess.run(command, capture_output=True, text=True, encoding='utf-8').returncode == 0


def test_backup_refuses_unsafe_destination_and_divergent_db(tmp_path, monkeypatch):
    store = Store(tmp_path / 'daily')
    store.create('合成项目')
    with pytest.raises(RuntimeError, match='上级目录'):
        create(store.folder, tmp_path)
    assert store.path.exists()
    original = backup.database_copy

    def changed_copy(source, destination):
        original(source, destination)
        with closing(sqlite3.connect(destination)) as db:
            db.execute("INSERT INTO settings VALUES('synthetic','{}')")
            db.commit()

    monkeypatch.setattr(backup, 'database_copy', changed_copy)
    target = tmp_path / 'divergent'
    with pytest.raises(RuntimeError, match='复制前'):
        create(store.folder, target)
    assert not target.exists() and store.path.exists()
