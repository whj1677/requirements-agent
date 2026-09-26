"""Isolated regression counterexamples, not business decisions or live results."""
import copy
import io
import pytest
import jsonschema
from docx import Document
from PIL import Image
from app.core import Problem
from app.understanding import validate, classification, contract
from app.contracts import VALIDATOR, SCHEMA
from app import ingest
from app.workflow import apply_response
from app.product_flow import question_open
from tests.test_runtime import client, empty_ingest


def proposal(text='收益明细与现平台保持一致。',kind='preserved'):
    return dict(temp_id='TMP-one',action='add',target_item_id=None,kind='requirement',title=text,
        statement=text,applies_to='to_be',epistemic_status='reported',source_refs=[dict(source_id='S',excerpt_id='E')],
        related_refs=[],change_type=kind,scope_evidence=[dict(state='preserve',quote=text,source_ref=dict(source_id='S',excerpt_id='E'))])


def q(key,text,**extras):
    return dict(temp_id='QTMP-'+key,topic_key=key,question=text,why='合成决定',options=[],blocking=True,
                blocking_stage=3,source_refs=[],related_refs=['REQ-one'],**extras)


@pytest.mark.parametrize('validator',[VALIDATOR,jsonschema.Draft202012Validator(ingest.schema())])
def test_single_question_shape_and_concrete_classification_evidence(validator):
    r=empty_ingest()
    r['questions']=[q('one','谁可删除记录？',decision_points=[])]
    assert validator.is_valid(r)
    r['questions'][0]['decision_points']=[q('child','谁可删除记录？')]
    assert not validator.is_valid(r)  # one child is not a real split
    r['questions'][0]['decision_points'].append(q('other','记录保留多久？'))
    assert validator.is_valid(r)
    r['questions']=[q('one','谁可删除记录？')]
    assert validator.is_valid(r)  # optional field can be omitted

    item=proposal()
    item.pop('scope_evidence')
    r['proposals']=[item]
    assert not validator.is_valid(r)  # concrete preserved identity needs cited evidence
    item['scope_evidence']=[]
    assert not validator.is_valid(r)
    item['change_type']='unspecified'
    assert validator.is_valid(r)
    item.pop('change_type')
    assert validator.is_valid(r)  # existing fixtures may leave classification undecided
    item['kind']='rule';item['change_type']='new'
    assert validator.is_valid(r)  # condition applies only to requirements


def test_electric_counterexample_cannot_label_preserved_detail_as_new():
    r=empty_ingest();r['proposals']=[proposal(kind='new')]
    with pytest.raises(Problem,match='分类缺少依据'):
        validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text=r['proposals'][0]['statement'])])
    r['proposals'][0]['change_type']='preserved'
    validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text=r['proposals'][0]['statement'])])
    mixed=copy.deepcopy(r['proposals'][0]);mixed['scope_evidence'].append(dict(state='change',quote='新增导出',source_ref=dict(source_id='S',excerpt_id='E')))
    with pytest.raises(Problem,match='保持与改变'):classification(mixed)


@pytest.mark.parametrize('states,expected', [([], 'unspecified'),(['current'],'existing'),(['change'],'new'),(['current','change'],'modified'),(['preserve'],'preserved')])
def test_classification_is_derived_from_explicit_evidence(states,expected):
    assert classification(dict(temp_id='TMP-x',scope_evidence=[dict(state=s) for s in states]))==expected


def test_fabricated_classification_quote_rejected():
    r=empty_ingest();r['proposals']=[proposal()]
    with pytest.raises(Problem,match='不是所引片段原文'):
        validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text='完全不同的材料')])


def test_model_cannot_disguise_preserved_quote_as_change():
    r=empty_ingest();r['proposals']=[proposal(kind='new')]
    r['proposals'][0]['scope_evidence'][0]['state']='change'
    with pytest.raises(Problem,match='明确保持'):
        validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text=r['proposals'][0]['statement'])])


def test_changed_behavior_can_cite_original_mixed_sentence_without_claiming_preserved_rule_is_new():
    # Reduced structure of a saved synthetic response: only the statement's new behavior is classified.
    source=('关键词和状态同时填写时按AND组合筛选。每次查询从第1页展示。'
            '原有每页条数和排序规则保持不变；没有匹配结果时显示空结果提示。')
    statement='关键词和状态按AND组合筛选，每次查询从第1页展示，空结果显示提示。'
    item=proposal(text=statement,kind='new')
    item['scope_evidence']=[dict(state='change',quote=source,
                                 source_ref=dict(source_id='S',excerpt_id='E'))]
    r=empty_ingest();r['proposals']=[item]
    validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text=source)])

    # A source containing only a preserved rule still cannot support a new classification.
    item['scope_evidence'][0]['quote']='原有每页条数和排序规则保持不变'
    with pytest.raises(Problem,match='仅含明确保持'):
        validate(r,dict(items=[],questions=[]),[dict(source_id='S',id='E',text=source)])


def test_answer_ref_error_identifies_candidate_question_target_and_correct_structure():
    r=empty_ingest()
    r['proposals']=[dict(temp_id='TMP-rule-answer',action='revise',target_item_id='RULE-old',
        kind='rule',title='规则修订',statement='依据回答补充规则',applies_to='to_be',
        epistemic_status='reported',source_refs=[],related_refs=['REQ-old'],answer_refs=['Q-answered'])]
    project=dict(items=[dict(id='REQ-old',kind='requirement')],questions=[dict(
        id='Q-answered',status='answered',related_refs=['REQ-old','RULE-old'])])
    with pytest.raises(Problem) as caught:
        validate(r,project,[])
    assert caught.value.code=='REFERENCE_INVALID'
    for text in ('TMP-rule-answer','RULE-old','Q-answered','REQ-old','source_refs','related_refs'):
        assert text in caught.value.message
    assert 'kind=requirement、action=revise' in contract()['answer_binding']
    assert 'rule修订和新增acceptance不写answer_refs' in SCHEMA['$defs']['Proposal']['properties']['answer_refs']['description']


def test_legacy_rule_candidate_cannot_mark_requirement_answer_applied():
    from app.understanding import accept_answer_update
    from app.core import digest
    question=dict(id='Q',status='answered',answer='保留字段',related_refs=['REQ','RULE'],understanding_status='pending')
    project=dict(items=[dict(id='REQ',kind='requirement'),dict(id='RULE',kind='rule')],questions=[question])
    candidate=dict(kind='rule',target_item_id='RULE',answer_refs=['Q'],answer_versions={'Q':digest(question['answer'])})
    with pytest.raises(Problem,match='原需求修订'):
        accept_answer_update(project,candidate)
    assert question['understanding_status']=='pending' and 'applied_requirement_ids' not in question


def test_answer_revision_requires_explicit_acceptance_and_latest_answer(tmp_path):
    from app.understanding import annotate_candidates,accept_answer_update,record_answer
    from app.product_flow import missing,BEHAVIOR
    from tests.helpers import prepared
    from app.store import Store
    p=prepared(Store(tmp_path));p['product_context']={'scope_ids':['REQ-0001']}
    p['items'][0]['behavior']={k:'合成已知' for k in BEHAVIOR}
    question=p['questions'][0];question.update(blocking=True)
    record_answer(question);question.update(status='answered',answer='预先准备：允许相接。')
    assert any('修订仍待核对' in s for s in missing(p,3))
    candidate=dict(id='C',target_item_id='REQ-0001',answer_refs=[question['id']])
    annotate_candidates(p,[candidate])
    assert question['understanding_status']=='candidate_ready' and '允许相接' not in p['items'][0]['statement']
    accept_answer_update(p,candidate)
    assert not any('修订仍待核对' in s for s in missing(p,3))
    record_answer(question);question['answer']='预先准备的修正：不允许相接。'
    with pytest.raises(Problem,match='回答已变化'):accept_answer_update(p,candidate)


def test_question_key_is_scoped_to_requirement():
    from app.understanding import equivalent
    a=q('permission','谁可编辑备注？');b=q('permission','谁可导出报表？');b['related_refs']=['REQ-other']
    assert not equivalent(a,b)


def test_shared_answer_requires_each_affected_requirement_update():
    from app.understanding import annotate_candidates,accept_answer_update,record_answer
    question=dict(id='Q',related_refs=['R1','R2'],answer='合成回答',status='answered')
    p=dict(questions=[question],items=[dict(id=r,kind='requirement') for r in ('R1','R2')])
    candidates=[dict(target_item_id=r,answer_refs=['Q']) for r in ('R1','R2')]
    annotate_candidates(p,candidates)
    accept_answer_update(p,candidates[0])
    assert question['understanding_status']=='candidate_ready'
    assert question['applied_requirement_ids']==['R1'] and 'R2' in question['answer_effect']
    accept_answer_update(p,candidates[1])
    assert question['understanding_status']=='applied'
    record_answer(question)
    assert question['applied_requirement_ids']==[] and question['understanding_status']=='pending'


def test_compound_must_be_explicit_decision_points():
    r=empty_ingest();r['questions']=[q('compound','谁可编辑？最长多少字？')]
    with pytest.raises(Problem,match='多个独立问句'):validate(r,dict(items=[],questions=[]),[])


def test_split_preserves_original_and_answers_are_independent(tmp_path):
    c,app=client(tmp_path)
    p=app.state.store.create('隔离联系人')
    old=dict(id='Q-original',**{k:v for k,v in q('old','谁可编辑？备注最长多少字？').items() if k!='temp_id'},status='open',answer=None)
    with app.state.store.edit(p['id'],0,'synthetic setup') as (p,_):
        p['questions']=[old];p['items']=[dict(id='REQ-one',kind='requirement',title='联系人备注',statement='新增备注',related_refs=[],selection_status='candidate',applies_to='to_be')]
    response=empty_ingest();response['questions']=[q('group','联系人备注相关决定',target_question_id='Q-original',decision_points=[q('permission','谁可编辑备注？'),q('length','备注最长多少字？')])]
    validate(response,p,[])
    with app.state.store.edit(p['id'],1,'synthetic response') as (p,_):apply_response(p,response)
    assert len(p['questions'])==3 and p['questions'][0]['question']==old['question']
    assert not question_open(p['questions'][0]) and all(question_open(x) for x in p['questions'][1:])
    first,second=p['questions'][1:]
    with c:
        result=c.post(f'/api/projects/{p["id"]}/answers',json={'expected_revision':2,'answers':{first['id']:'预先准备的合成回答：管理员。'}})
        assert result.status_code==200,result.text
        current=app.state.store.get(p['id'])
        assert current['questions'][1]['status']=='answered' and current['questions'][2]['status']=='open'
        assert current['items'][0]['selection_status']=='candidate'
        again=c.post(f'/api/projects/{p["id"]}/answers',json={'expected_revision':3,'answers':{first['id']:'预先准备的合成修订回答：联系人创建人。'}})
        assert again.status_code==200
        assert app.state.store.get(p['id'])['questions'][1]['answer_history'][0]['answer'].endswith('管理员。')
        # A shared topic key cannot prove that a changed question has the old answer.
        r=empty_ingest();r['questions']=[q('permission','允许编辑的角色是什么？')]
        with app.state.store.edit(p['id'],4,'synthetic repeat') as (draft,_):apply_response(draft,r)
        assert len(draft['questions'])==4 and draft['questions'][1]['status']=='answered'
        assert draft['questions'][-1]['status']=='open'
        assert draft['questions'][-1]['prior_question_id']==draft['questions'][1]['id']
        # Exact normalized text can reuse the answered identity without adding a question.
        same=empty_ingest();same['questions']=[q('permission','谁可编辑备注？')]
        with app.state.store.edit(p['id'],5,'synthetic exact repeat') as (draft,_):apply_response(draft,same)
        assert len(draft['questions'])==4 and draft['questions'][1]['status']=='answered'


def word_with_picture():
    doc=Document();doc.add_paragraph('合成联系人：只增加选填备注，权限保持。')
    png=io.BytesIO();Image.new('RGB',(120,80),'white').save(png,format='PNG');png.seek(0)
    doc.add_picture(png);raw=io.BytesIO();doc.save(raw);return raw.getvalue()


def test_word_picture_http_preview_authorization_and_context(tmp_path):
    from app.provider import assemble, DEFAULT
    c,app=client(tmp_path)
    with c:
        p=c.post('/api/projects',json={'name':'合成图文'}).json();url='/api/projects/'+p['id']
        uploaded=c.post(url+'/sources/file',data={'expected_revision':0,'purpose':'current'},files={'file':('sample.docx',word_with_picture())})
        assert uploaded.status_code==200,uploaded.text
        p=app.state.store.get(p['id']);parent,child=p['sources']
        assert parent['embedded_image_ids']==[child['id']] and child['parse_status']=='awaiting_vision'
        assert c.get(url+'/sources/'+child['id']+'/image').headers['content-type']=='image/png'
        messages,excerpts,omitted=assemble(p,'vision','',DEFAULT,tmp_path)
        assert any(b['type']=='image_url' for b in messages[1]['content'])
        text_messages,_,omitted=assemble(p,'ingest','',DEFAULT,tmp_path)
        assert child['excerpts'][0]['id'] in omitted  # pixels never masquerade as understood text
        c.post(url+'/sources/'+parent['id']+'/exclude',json={'expected_revision':1})
        assert all(s['excluded'] for s in app.state.store.get(p['id'])['sources'])
