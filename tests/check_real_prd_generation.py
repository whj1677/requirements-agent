"""Opt-in real PRD generation with an explicitly frozen confirmed-input fixture.

This tests the normal generation/download stage, not real-model clarification
or review. The separate raw-material and conflict trials retain their failures.
"""
import argparse,asyncio,io,json,os,sqlite3,sys,uuid,zipfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
os.environ.setdefault('RA_DATA_DIR',str(ROOT/'evidence/full-acceptance-20261027/prd-import'))
import httpx
from app.main import create_app
from app.prewrite import apply_result
from app.provider import Provider,DEFAULT,origin
from tests.test_prewrite_consistency import isolated_store,review
from scripts.check_shared_draft_live import daily_fingerprint


async def run(args):
    assert args.execute_real_model
    output=args.output.resolve();assert output.is_relative_to(ROOT/'evidence') and not output.exists();output.mkdir(parents=True)
    before=daily_fingerprint();provider=Provider(ROOT/'.env',credential_dir=ROOT/'data/.credentials')
    app=create_app(output/'data',provider=provider);store,p=isolated_store(output/'data')
    with store.edit(p['id'],p['revision'],'冻结已核对的合成输入',bump=False) as (current,_):
        current['name']='合成验收：已核对需求的 PRD 生成'
        apply_result(current,review(current),'FROZEN-REVIEW-PRECONDITION')
    pid=p['id']
    with sqlite3.connect((ROOT/'data/requirements.sqlite3').as_uri()+'?mode=ro',uri=True) as db:
        config=dict(DEFAULT,**json.loads(db.execute("select payload from settings where id='model'").fetchone()[0]))
    assert origin(config)=='https://api.deepseek.com' and config['model']=='deepseek-flash' and provider.key(config)
    store.setting('model',config);report=dict(status='running',synthetic_only=True,prewrite_precondition='frozen test fixture, not real review')
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://testserver') as client:
            csrf=(await client.get('/api/session')).json()['csrf'];headers={'origin':'http://testserver','x-csrf-token':csrf};endpoint='/api/projects/'+pid
            # Match the normal UI click: extra instructions are new input and
            # intentionally require another review instead of reusing it.
            body=dict(expected_revision=store.get(pid)['revision'],action='document',message='',document_type='prd',option_id=None)
            response=await client.post(endpoint+'/actions/plan',headers=headers,json=body);assert response.status_code==200,response.text
            plan=response.json();assert not plan['missing'] and plan['stages']==['prd'],plan
            response=await client.post(endpoint+'/actions',headers=headers,json=dict(body,plan_hash=plan['plan_hash'],authorize=True,idempotency_key=str(uuid.uuid4())))
            assert response.status_code==200,response.text
            task=response.json();job=app.state.actions.jobs.get(task['id'])
            if job:await job
            task=store.get_record(pid,task['id'],'user_task');report.update(task_status=task['status'],real_model_calls=task['calls'])
            assert task['status'] in ('succeeded','partial'),task.get('message')
            artifact=store.get(pid)['documents']['prd'];assert artifact
            response=await client.get(endpoint+'/documents/prd/zip?mode=review');assert response.status_code==200,response.text[:500]
            (output/'prd-review.zip').write_bytes(response.content)
            with zipfile.ZipFile(io.BytesIO(response.content)) as package:
                assert package.testzip() is None
                word=next(n for n in package.namelist() if n.endswith('.docx'))
                from docx import Document
                document=Document(io.BytesIO(package.read(word)));assert len(document.paragraphs)>5 and document.tables
                report['files']=package.namelist()
            report.update(status='passed',document_id=artifact['id'],docx_verified=True)
    except Exception as error:
        import traceback
        report.update(status='failed',error=str(error)[:2000],traceback=traceback.format_exc())
    report['daily_data_unchanged']=daily_fingerprint()==before
    (output/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf8')
    print(json.dumps({k:v for k,v in report.items() if k!='traceback'},ensure_ascii=False),flush=True)
    assert report['status']=='passed' and report['daily_data_unchanged']


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--execute-real-model',action='store_true');asyncio.run(run(p.parse_args()))
