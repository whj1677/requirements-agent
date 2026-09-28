"""Product-document projection removes workflow/audit text without mutating records."""
import asyncio
import copy
import io
import json
import zipfile

import pytest
from docx import Document

from app.core import brief_hash, dumps
from app.document_reader import export_readiness, reader_document
from app.exports import document_files, handoff_files
from app.prd import compile_plan
from app.store import Store
from app.workflow import apply_response
from tests.helpers import prepared


@pytest.mark.parametrize('kind', ['prd', 'mrd'])
def test_reader_markdown_docx_and_handoff_json_share_clean_product_projection(
        tmp_path, monkeypatch, kind):
    async def no_images(_):
        return []

    monkeypatch.setattr('app.exports.capture', no_images)
    p = prepared(Store(tmp_path))
    answer = copy.deepcopy(p['questions'][0])
    answer.update(id='Q-PRIVATE-ANSWER', status='answered',
                  answer='PRIVATE_ANSWER_TEXT_NOT_FOR_PRODUCT_DOC',
                  understanding_status='candidate_ready', applied_item_ids=[])
    open_question = copy.deepcopy(p['questions'][0])
    open_question.update(id='Q-PRODUCT-OPEN', question='产品边界仍需决定？',
                         why='决定相邻时段是否允许。', status='open', answer=None)
    p['questions'] = [answer, open_question]

    rejected = copy.deepcopy(p['items'][0])
    rejected.update(id='REQ-PRIVATE-REJECTED', selection_status='rejected',
                    statement='PRIVATE_REJECTED_DISCUSSION_TEXT')
    p['items'].append(rejected)
    approved_goal = copy.deepcopy(p['items'][0])
    approved_goal.update(id='GOAL-INTERNAL-77', kind='goal', title='降低操作负担',
                         statement='值班人员能够快速完成日常核对。',
                         selection_status='selected', epistemic_status='reported')
    proposed_constraint = copy.deepcopy(approved_goal)
    proposed_constraint.update(id='CONSTRAINT-PROPOSED-88', kind='constraint',
                               title='PRIVATE_PROPOSED_CONSTRAINT',
                               statement='PRIVATE_PROPOSED_CONSTRAINT_TEXT',
                               epistemic_status='proposed')
    p['items'].extend([approved_goal, proposed_constraint])
    p['sources'][0].update(parse_status='partial',
                           failure_reason='SOURCE_PRIVATE_ID EX_PRIVATE_ID PRIVATE_SOURCE_EXCEPTION')

    for document_kind in ('prd', 'mrd'):
        artifact = p['documents'][document_kind]
        artifact['delivery_item_ids'] = [i['id'] for i in p['items']
                                         if i.get('selection_status') == 'selected'
                                         and i.get('applies_to') == 'to_be']
        artifact['item_snapshot'] = copy.deepcopy(p['items'])
        artifact['item_snapshot'].append(copy.deepcopy(rejected))
        artifact['item_snapshot'].extend([copy.deepcopy(approved_goal), copy.deepcopy(proposed_constraint)])
        artifact['question_snapshot'] = copy.deepcopy(p['questions'])
        artifact['source_snapshot'] = [dict(id='SRC_PRIVATE_ID', title='PRIVATE_SOURCE_NAME',
                                             parse_status='partial',
                                             failure_reason='SOURCE_PRIVATE_ID EX_PRIVATE_ID')]
        sections = artifact['content']['sections']
        sections.append(dict(section_id=document_kind.upper()+'-DISCUSS', title='附录：条目身份与历史讨论',
                             level=1, parent_section_id=None,
                             blocks=[dict(kind='goal', ref_ids=[rejected['id']],
                                          text='PRIVATE_REJECTED_DISCUSSION_TEXT')]))
        sections.append(dict(section_id=document_kind.upper()+'-ANSWERS', title='已有回答与未决问题',
                             level=1, parent_section_id=None, blocks=[
                                 dict(kind='open_question', ref_ids=[answer['id']],
                                      text='PRIVATE_ANSWER_TEXT_NOT_FOR_PRODUCT_DOC'),
                                 dict(kind='open_question', ref_ids=[open_question['id']],
                                      text='产品边界仍需决定？（尚未决定）'),
                             ]))
        sections.append(dict(section_id=document_kind.upper()+'-HISTORY', title='历史分析提示（非当前事实）',
                             level=1, parent_section_id=None, blocks=[
                                 dict(kind='narrative', ref_ids=[],
                                      text='PRIVATE_MODEL_STAGE_PROMPT_MARKER 第9轮 调用模型 EX_PRIVATE_ID')]))
        sections.append(dict(section_id=document_kind.upper()+'-LIMITS', title='文档边界与资料限制',
                             level=1, parent_section_id=None, blocks=[
                                 dict(kind='narrative', ref_ids=[],
                                      text='材料「PRIVATE_SOURCE_NAME」尚有读取限制：SOURCE_PRIVATE_ID EX_PRIVATE_ID PRIVATE_SOURCE_EXCEPTION'),
                                 dict(kind='open_question', ref_ids=[],
                                      text='本轮未完整送入的来源片段：EX_PRIVATE_ID'),
                                 dict(kind='narrative', ref_ids=[],
                                      text='章节建议未指定的已选规范条目：REQ-0001'),
                             ]))
        # A retained child must keep its structural parent, even if the parent
        # has no product prose after filtering.
        sections.append(dict(section_id=document_kind.upper()+'-EMPTY-PARENT', title='产品能力',
                             level=1, parent_section_id=None, blocks=[
                                 dict(kind='goal', ref_ids=['REQ-PRIVATE-REJECTED'],
                                      text='PRIVATE_REJECTED_PARENT_TEXT')]))
        sections.append(dict(section_id=document_kind.upper()+'-CHILD', title='已选能力',
                             level=2, parent_section_id=document_kind.upper()+'-EMPTY-PARENT', blocks=[
                                 dict(kind='requirement', ref_ids=['REQ-0001'], text='will be reconstructed')]))
        sections.append(dict(section_id=document_kind.upper()+'-BUSINESS', title='目标与价值',
                             level=1, parent_section_id=None, blocks=[
                                 dict(kind='goal', ref_ids=[approved_goal['id']], text='audit rendering'),
                                 dict(kind='constraint', ref_ids=[proposed_constraint['id']],
                                      text='PRIVATE_PROPOSED_CONSTRAINT_TEXT')]))
        artifact['content']['reference_mapping'] = [
            dict(profile_section_id='internal-profile-id', scope_ref='SRC_PRIVATE_ID',
                 disposition='pending', output_section_ids=[document_kind.upper()+'-DISCUSS'],
                 reason='PRIVATE_MAPPING_DISCUSSION_REASON')]
        artifact['content']['coverage'] = [
            dict(item_id='REQ-PRIVATE-REJECTED', section_ids=[document_kind.upper()+'-DISCUSS'])]
        artifact['brief_hash'] = brief_hash(p)

    canonical_before = copy.deepcopy(p['documents'])
    artifact = p['documents'][kind]
    reader = reader_document(artifact)
    reader_text = dumps(reader)
    assert 'REQ-0001' in reader_text and 'v1' in reader_text
    assert 'PRIVATE_REJECTED_DISCUSSION_TEXT' not in reader_text
    assert 'PRIVATE_ANSWER_TEXT_NOT_FOR_PRODUCT_DOC' not in reader_text
    assert 'PRIVATE_MODEL_STAGE_PROMPT_MARKER' not in reader_text
    assert 'PRIVATE_SOURCE_NAME' not in reader_text
    assert 'SOURCE_PRIVATE_ID' not in reader_text and 'EX_PRIVATE_ID' not in reader_text
    assert '产品边界仍需决定？' in reader_text
    assert '决定相邻时段是否允许。' in reader_text
    assert '本期未决' in reader_text
    assert '部分材料内容未能完整读取' in reader_text
    assert '降低操作负担' in reader_text and '值班人员能够快速完成日常核对。' in reader_text
    assert 'GOAL-INTERNAL-77' not in reader_text
    assert 'PRIVATE_PROPOSED_CONSTRAINT' not in reader_text
    assert 'reference_mapping' not in reader and 'coverage' in reader
    assert 'PRIVATE_MAPPING_DISCUSSION_REASON' not in reader_text
    assert 'SRC_PRIVATE_ID' not in reader_text and 'internal-profile-id' not in reader_text
    assert all(c['item_id'] != 'REQ-PRIVATE-REJECTED' for c in reader['coverage'])
    assert all(section_id != kind.upper()+'-DISCUSS'
               for c in reader['coverage'] for section_id in c['section_ids'])
    section_by_id = {s['section_id']: s for s in reader['sections']}
    assert kind.upper()+'-EMPTY-PARENT' in section_by_id
    assert section_by_id[kind.upper()+'-EMPTY-PARENT']['blocks'] == []
    assert section_by_id[kind.upper()+'-CHILD']['parent_section_id'] == kind.upper()+'-EMPTY-PARENT'
    assert not export_readiness(p, kind)['ready']
    assert any(issue['code'] == 'CLARIFICATION_REQUIRED'
               for issue in export_readiness(p, kind)['issues'])

    files = asyncio.run(document_files(p, kind))
    markdown = files[kind.upper()+'.md'].decode('utf-8')
    paragraphs = '\n'.join(paragraph.text for paragraph in
                           Document(io.BytesIO(files[kind.upper()+'.docx'])).paragraphs)
    for output in (markdown, paragraphs):
        for private in ('PRIVATE_REJECTED_DISCUSSION_TEXT', 'PRIVATE_ANSWER_TEXT_NOT_FOR_PRODUCT_DOC',
                        'PRIVATE_MODEL_STAGE_PROMPT_MARKER', 'PRIVATE_SOURCE_NAME',
                        'SOURCE_PRIVATE_ID', 'EX_PRIVATE_ID', 'PRIVATE_SOURCE_EXCEPTION',
                        'PRIVATE_ANSWER'):
            assert private not in output
        assert 'REQ-0001' in output and 'v1' in output
        assert '产品边界仍需决定？' in output
        assert '决定相邻时段是否允许。' in output
        assert '部分材料内容未能完整读取' in output
        assert '降低操作负担' in output and '值班人员能够快速完成日常核对。' in output
        assert 'GOAL-INTERNAL-77' not in output and 'PRIVATE_PROPOSED_CONSTRAINT' not in output

    baseline = dict(id='BASE-SYNTHETIC', project=p, hashes={}, confirmation_id='CONF-SYNTHETIC')
    handoff, _ = handoff_files(baseline, [])
    content_json = json.loads(handoff[kind.upper()+'.content.json'])
    handoff_text = dumps(content_json)
    for private in ('PRIVATE_REJECTED_DISCUSSION_TEXT', 'PRIVATE_ANSWER_TEXT_NOT_FOR_PRODUCT_DOC',
                    'PRIVATE_MODEL_STAGE_PROMPT_MARKER', 'PRIVATE_SOURCE_NAME',
                    'SOURCE_PRIVATE_ID', 'EX_PRIVATE_ID', 'PRIVATE_SOURCE_EXCEPTION'):
        assert private not in handoff_text
    assert '产品边界仍需决定？' in handoff_text
    assert '降低操作负担' in handoff_text and '值班人员能够快速完成日常核对。' in handoff_text
    assert 'GOAL-INTERNAL-77' not in handoff_text and 'PRIVATE_PROPOSED_CONSTRAINT' not in handoff_text
    assert 'reference_mapping' not in content_json
    assert 'PRIVATE_MAPPING_DISCUSSION_REASON' not in handoff_text
    assert all(c['item_id'] != 'REQ-PRIVATE-REJECTED' for c in content_json['coverage'])
    assert 'REQ-0001' in handoff_text and 'v1' in handoff_text
    assert p['documents'] == canonical_before


@pytest.mark.parametrize('kind', ['prd', 'mrd'])
def test_compile_plan_promotes_only_reported_in_scope_business_items_and_keeps_prose(tmp_path, kind):
    p = prepared(Store(tmp_path))
    goal = dict(id='GOAL-SELECTED-REAL-COMPILE', kind='goal', title='减少重复操作',
                statement='操作人员能够一次完成登记。', applies_to='to_be',
                epistemic_status='reported', source_refs=[], related_refs=[],
                selection_status='selected', revision=p['revision'])
    reference_assumption = dict(id='GOAL-REFERENCE-ASSUMPTION', kind='goal', title='PRIVATE_REFERENCE_GOAL',
                                statement='PRIVATE_REFERENCE_ASSUMPTION', applies_to='reference',
                                epistemic_status='reported', source_refs=[], related_refs=[],
                                selection_status='selected', revision=p['revision'])
    unknown_epistemic = dict(id='GOAL-UNKNOWN-EPISTEMIC', kind='goal', title='PRIVATE_UNKNOWN_GOAL',
                             statement='PRIVATE_UNKNOWN_EPISTEMIC_TEXT', applies_to='to_be',
                             epistemic_status='unknown', source_refs=[], related_refs=[], selection_status='selected',
                             revision=p['revision'])
    p['items'].extend([goal, reference_assumption, unknown_epistemic])
    plan = dict(plan_version='2', sections=[dict(section_key='function', context_refs=[],
                normative_refs=['REQ-0001'], discussion_refs=[])])
    apply_response(p, compile_plan(plan, p, kind))
    artifact = p['documents'][kind]
    # compile_plan puts omitted selected business items in the canonical audit appendix.
    appendix = next(s for s in artifact['content']['sections']
                    if s['title'] == '附录：条目身份与历史讨论')
    assert any(goal['id'] in b['ref_ids'] for b in appendix['blocks'])
    appendix_before = copy.deepcopy(appendix)
    artifact['content']['sections'][0]['blocks'].append(dict(
        kind='narrative', ref_ids=[], text='本轮改造应继续支持值班人员快速核对。'))
    canonical_before = copy.deepcopy(artifact['content'])

    view = reader_document(artifact)
    text = dumps(view)
    assert '减少重复操作' in text and goal['statement'] in text
    assert goal['id'] not in text and '来源：' not in text
    assert 'PRIVATE_REFERENCE_GOAL' not in text and 'PRIVATE_REFERENCE_ASSUMPTION' not in text
    assert 'PRIVATE_UNKNOWN_GOAL' not in text and 'PRIVATE_UNKNOWN_EPISTEMIC_TEXT' not in text
    assert '本轮改造应继续支持值班人员快速核对。' in text
    assert artifact['content'] == canonical_before
    assert appendix == appendix_before


@pytest.mark.parametrize('kind', ['prd', 'mrd'])
def test_product_layout_only_section_retains_its_image_without_audit_appendix(tmp_path, kind):
    p = prepared(Store(tmp_path))
    artifact = p['documents'][kind]
    artifact['item_snapshot'] = copy.deepcopy(p['items'])
    artifact['delivery_item_ids'] = [i['id'] for i in p['items']]
    sid = kind.upper()+'-PRODUCT-LAYOUT'
    artifact['content']['sections'].append(dict(section_id=sid, title='页面布局',
        level=1, parent_section_id=None, blocks=[]))
    dimension = 'PRD-4.F.1' if kind=='prd' else 'MRD-5.1.F.3'
    for mapping in artifact['content']['reference_mapping']:
        if mapping['profile_section_id']==dimension:
            mapping['output_section_ids']=[sid]
    before = copy.deepcopy(artifact)
    files = asyncio.run(document_files(p, kind))
    bindings = json.loads(files['document_asset_bindings.json'])
    assert bindings and {b['section_id'] for b in bindings}=={sid}
    assert '页面布局' in files[kind.upper()+'.md'].decode()
    assert '参考维度处置' not in files[kind.upper()+'.md'].decode()
    assert len(Document(io.BytesIO(files[kind.upper()+'.docx'])).inline_shapes)>0
    assert artifact==before
