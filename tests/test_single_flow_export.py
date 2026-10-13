"""A complete flow occupies one landscape sheet, without changing body layout."""
import asyncio
import copy
import io
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from PIL import Image
import pytest

from app.exports import document_files
from tests.test_detail_document import project_fixture, compile_fixture


def section_contents(document):
    sections=[[]]
    for child in document._body._body:
        sections[-1].append(child)
        if child.tag==qn('w:p') and child.find('./'+qn('w:pPr')+'/'+qn('w:sectPr')) is not None:
            sections.append([])
    return sections


def test_complete_flow_is_one_media_image_in_landscape_with_its_heading(tmp_path,monkeypatch):
    p=project_fixture(tmp_path);compile_fixture(p)
    before=copy.deepcopy(p)
    # A complete-image export must not silently revert to sliced chart pages.
    def obsolete_pagination(*args,**kwargs):
        raise AssertionError('Word flow must not be segmented')
    monkeypatch.setattr('app.detail_rendering.flow_png_pages',obsolete_pagination)
    files=asyncio.run(document_files(p,'prd'))
    document=Document(io.BytesIO(files['PRD.docx']))
    assert len(document.inline_shapes)==1
    assert [s.orientation for s in document.sections]==[
        WD_ORIENT.PORTRAIT,WD_ORIENT.LANDSCAPE,WD_ORIENT.PORTRAIT]
    first,last=document.sections[0],document.sections[-1]
    for key in ('page_width','page_height','top_margin','bottom_margin','left_margin','right_margin'):
        assert getattr(first,key)==getattr(last,key)
    contents=section_contents(document)
    assert len(contents)==3
    text=lambda nodes:' '.join(t.text or '' for n in nodes for t in n.iter(qn('w:t')))
    assert '业务流程说明' in text(contents[1]) and '业务流程说明' not in text(contents[0])
    assert sum(len(n.findall('.//'+qn('w:drawing'))) for n in contents[1])==1
    assert '流程步骤' in text(contents[2]) and '分支与流转条件' in text(contents[2])
    png=next(data for name,data in files.items() if name.startswith('assets/flow-') and name.endswith('.png'))
    with zipfile.ZipFile(io.BytesIO(files['PRD.docx'])) as archive:
        media=[n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media)==1 and archive.read(media[0])==png
    assert document.styles['Normal'].font.size.pt==11
    assert all(run.font.size.pt==9.5 for table in document.tables for row in table.rows
               for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs)
    assert p==before


@pytest.mark.parametrize('size',[(2000,1200),(800,2600)])
def test_image_keeps_full_pixels_and_aspect_ratio_inside_one_landscape_sheet(tmp_path,monkeypatch,size):
    p=project_fixture(tmp_path);compile_fixture(p)
    image=Image.new('RGB',size,'white')
    # Different corners make cropping or replacing the original bytes detectable.
    image.putpixel((0,0),(255,0,0));image.putpixel((size[0]-1,size[1]-1),(0,0,255))
    stream=io.BytesIO();image.save(stream,format='PNG');png=stream.getvalue()
    monkeypatch.setattr('app.detail_rendering.flow_png',lambda _:png)
    files=asyncio.run(document_files(p,'prd'))
    document=Document(io.BytesIO(files['PRD.docx']))
    assert len(document.inline_shapes)==1
    shape=document.inline_shapes[0];page=document.sections[1]
    assert shape.width<=page.page_width-page.left_margin-page.right_margin
    assert shape.height<page.page_height-page.top_margin-page.bottom_margin
    assert shape.width/shape.height==pytest.approx(size[0]/size[1],rel=.00001)
    with zipfile.ZipFile(io.BytesIO(files['PRD.docx'])) as archive:
        media=[n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media)==1 and archive.read(media[0])==png


def test_unstructured_export_keeps_its_existing_portrait_geometry(tmp_path):
    from app.store import Store
    from tests.helpers import prepared
    p=prepared(Store(tmp_path));p['ui']=None
    files=asyncio.run(document_files(p,'prd',reader=False))
    document=Document(io.BytesIO(files['PRD.docx']))
    assert len(document.sections)==1
    assert document.sections[0].orientation==WD_ORIENT.PORTRAIT
    assert document.sections[0].page_width.inches==8.5
    assert document.sections[0].page_height.inches==11
