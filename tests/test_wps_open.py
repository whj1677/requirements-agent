"""Real opt-in WPS coexistence: preserve an owned, visible, unsaved document."""
import hashlib
import io
import os

import pytest

from app import office

pytestmark = pytest.mark.skipif(os.name != 'nt' or os.environ.get('RA_TEST_WPS_OPEN') != '1',
                               reason='Explicit opt-in: visible synthetic WPS editing session')


@pytest.mark.parametrize('kind,extension', [('Word', '.docx'), ('Excel', '.xlsx')])
def test_reading_while_wps_is_open_preserves_unsaved_document(tmp_path, kind, extension):
    from app.office_worker import com_runtime
    from app.office_process import processes, cleanup_identity
    from docx import Document
    from openpyxl import Workbook
    selected = office.registered_application(kind)
    if not selected or not selected['application'].startswith('WPS'):
        pytest.skip('Requires actual WPS for this component')
    output = io.BytesIO()
    if kind == 'Word':
        original = Document()
        original.add_paragraph('合成储能云平台：已保存的站点资料。')
    else:
        original = Workbook()
        original.active['A1'] = '合成储能云平台：已保存的站点资料。'
    original.save(output)
    raw = output.getvalue()
    source = tmp_path / ('open-editing' + extension)
    source.write_bytes(raw)
    digest = hashlib.sha256(raw).hexdigest()
    pythoncom, dynamic = com_runtime()
    pythoncom.CoInitializeEx(pythoncom.COINIT_APARTMENTTHREADED)
    app = document = collection = None
    owned = []
    try:
        before = {(p['pid'], p['started']) for p in processes(selected['process'])}
        app = dynamic.Dispatch(pythoncom.CoCreateInstance(selected['clsid'], None,
                              pythoncom.CLSCTX_LOCAL_SERVER, pythoncom.IID_IDispatch))
        collection = getattr(app, 'Documents' if kind == 'Word' else 'Workbooks')
        created = [p for p in processes(selected['process']) if (p['pid'], p['started']) not in before]
        new_ids = {p['pid'] for p in created}
        assert len([p for p in created if p['parent_pid'] not in new_ids]) == 1 and collection.Count == 0
        owned = created
        app.AutomationSecurity = 3
        assert app.AutomationSecurity == 3
        document = collection.Open(str(source))
        app.Visible = True
        marker = '未保存编辑必须保留，不能混入上传资料'
        if kind == 'Word':
            document.Content.InsertAfter(marker)
            read_text = lambda: document.Content.Text
            prior_links = app.Options.UpdateLinksAtOpen
        else:
            document.Worksheets.Item(1).Cells(2, 1).Value = marker
            read_text = lambda: document.Worksheets.Item(1).Cells(2, 1).Value
            prior_links = app.AskToUpdateLinks
        assert document.Saved is False
        text_before, count_before, security = read_text(), collection.Count, app.AutomationSecurity
        visible_before = app.Visible
        folder = tmp_path / 'agent-copy'
        folder.mkdir()
        result = office.extract_office(source.name, raw, folder)
        assert result['status'] == 'partial', result
        assert result['application'].startswith('WPS') and result['owned_process_cleanup'] == 'verified'
        text = '\n'.join(row[1] for row in result['rows'])
        assert '已保存的站点资料' in text and marker not in text
        assert document.Saved is False and read_text() == text_before
        assert collection.Count == count_before and app.Visible == visible_before
        assert app.AutomationSecurity == security
        assert (app.Options.UpdateLinksAtOpen if kind == 'Word' else app.AskToUpdateLinks) == prior_links
        assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
        active = {(p['pid'], p['started']) for p in processes(selected['process'])}
        assert {(p['pid'], p['started']) for p in owned} <= active
        assert not list(folder.iterdir())
    finally:
        import gc
        if owned:
            if document is not None:
                document.Close(False)
        collection = document = app = None
        gc.collect()
        for identity in owned:
            assert cleanup_identity(identity)
        pythoncom.CoUninitialize()
