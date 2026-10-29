import asyncio
import copy
import io

from docx import Document
from docx.oxml.ns import qn
from app.exports import document_files
from app.word_presentation import adopted_text,reading_node,prose_parts
from app.requirement_confirmation import row_fingerprint
from tests.test_detail_document import project_fixture,compile_fixture


def test_adopted_labels_are_removed_without_rewriting_business_modality():
    assert adopted_text('建议：管理员可保存姓名。')=='管理员可保存姓名。'
    assert adopted_text('推荐方案：仅查询已有结果。')=='仅查询已有结果。'
    assert adopted_text('成功结果：建议：显示保存结果。')=='成功结果：显示保存结果。'
    assert adopted_text('优先级与依据：建议：建议按MVP优先。')=='优先级与依据：按MVP优先。'
    assert reading_node(dict(kind='paragraph',block_kind='narrative',text='相关页面或模块：建议：联系人详情'))['text']=='相关页面或模块：联系人详情'
    for text in ('不建议重复提交。','界面显示“建议：检查连接”。','建议用户先核对原记录。','待确认：超时时间尚未明确。'):
        assert adopted_text(text)==text
    pending=dict(kind='paragraph',block_kind='open_question',text='建议：由负责人决定是否纳入。')
    assert reading_node(pending)==pending


def test_style_cleanup_keeps_dependencies_and_hides_unused_word_gallery_entries():
    from docx.enum.style import WD_STYLE_TYPE
    from app.word_presentation import clean_style_gallery
    doc=Document();doc.styles.add_style('Unused custom style',WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading('需求说明',1);doc.add_paragraph('原文保留。')
    doc.add_table(rows=1,cols=1).style='Table Grid'
    before=len(doc.styles);gallery=clean_style_gallery(doc)
    stream=io.BytesIO();doc.save(stream);result=Document(io.BytesIO(stream.getvalue()))
    assert len(result.styles)<before and 'Unused custom style' not in result.styles
    assert result.paragraphs[0].text=='需求说明' and result.paragraphs[1].text=='原文保留。'
    definitions={s.style_id for s in result.styles}
    for style in result.styles:
        for name in ('basedOn','next','link'):
            ref=style.element.find(qn('w:'+name))
            assert ref is None or ref.get(qn('w:val')) in definitions
    assert set(gallery)=={'Normal','Title','Heading1'}
    assert 'TableGrid' in definitions and 'TableNormal' in definitions
    assert result.styles.element.find(qn('w:latentStyles')).get(qn('w:defQFormat'))=='0'


def test_export_keeps_source_content_and_outputs_clean_clauses_with_navigation(tmp_path):
    p=project_fixture(tmp_path);req=p['items'][0];req['statement']='建议：'+req['statement']
    ac=next(i for i in p['items'] if i['kind']=='acceptance');ac['statement']='推荐：'+ac['statement']
    req['requirement_confirmation']['content_hash']=row_fingerprint(p,req)
    compile_fixture(p);before=copy.deepcopy(p)
    files=asyncio.run(document_files(p,'prd'));doc=Document(io.BytesIO(files['PRD.docx']))
    assert p==before
    assert any(paragraph.text==req['statement'].removeprefix('建议：') for paragraph in doc.paragraphs)
    assert any(ac['statement'].removeprefix('推荐：') in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert not any(paragraph.text.startswith('建议：') for paragraph in doc.paragraphs)
    assert doc.paragraphs[0].style.name=='Title' and str(doc.styles['Title'].font.color.rgb)=='000000'
    links=doc._element.findall('.//'+qn('w:hyperlink'));assert links
    destinations={b.get(qn('w:name')) for b in doc._element.findall('.//'+qn('w:bookmarkStart'))}
    assert all(link.get(qn('w:anchor')) in destinations for link in links)
    assert doc.sections[0].different_first_page_header_footer
    assert all(not s.different_first_page_header_footer for s in list(doc.sections)[1:])
    footer=doc.sections[0].footer._element.xml
    assert 'PAGE' in footer and 'NUMPAGES' not in footer
    assert doc.tables[0].rows[0]._tr.find('./'+qn('w:trPr')+'/'+qn('w:tblHeader')) is not None


def test_long_prose_keeps_conditions_quotes_and_all_characters():
    paragraph=('租户管理员确认告警并记录备注，备注上限为200个Unicode码点。'
        '权限检查失败时保留备注、保持未确认；恢复权限后用户可手工重试。'
        '超时时显示“结果待核实。请查询已有记录。”，客户端不推断成功或失败。'
        '服务端如已处理请求，仍按实际结果记录日志。'
        '查询失败时继续保持待核实，不自动重发。'
        '校验失败且未发请求时，不生成服务端处理日志。'
        '单次记录耗时为1.5秒，阈值由已确认配置提供。'
        '并发确认以服务器先受理成功的有效记录为准，不覆盖原有确认人、时间和备注。\n  ')
    parts=prose_parts(paragraph)
    assert len(parts)>4 and ''.join(parts)==paragraph
    assert any('失败时保留备注、保持未确认；恢复权限后' in p for p in parts)
    assert any('“结果待核实。请查询已有记录。”' in p for p in parts)
    assert any('1.5秒' in p for p in parts)
    assert prose_parts('确认成功。显示已有结果。')==['确认成功。显示已有结果。']


def test_dense_requirement_exports_as_native_numbered_paragraphs_without_data_loss(tmp_path):
    p=project_fixture(tmp_path);req=p['items'][0]
    req['statement']=''.join(['用户保存姓名时必须校验当前电站权限，拒绝时保留输入；权限恢复后可手工重试。']*8)
    ac=next(i for i in p['items'] if i['kind']=='acceptance')
    ac['statement']=''.join(['提交失败处理：权限拒绝时原姓名不得变化，用户修改后重新提交。']*8)
    req['requirement_confirmation']['content_hash']=row_fingerprint(p,req)
    compile_fixture(p);before=copy.deepcopy(p)
    doc=Document(io.BytesIO(asyncio.run(document_files(p,'prd'))['PRD.docx']))
    clauses=[v for v in doc.paragraphs if v.style.name=='Requirement Clause']
    assert ''.join(v.text for v in clauses)==req['statement']
    assert all(v._p.pPr.numPr.numId.val>0 for v in clauses)
    assert len({v._p.pPr.numPr.numId.val for v in clauses})==1
    assert doc.styles['Requirement Clause'].paragraph_format.first_line_indent.pt<0
    cells=[cell for table in doc.tables for row in table.rows for cell in row.cells
           if ac['statement'].split('。')[0] in cell.text]
    assert len(cells)==1 and len(cells[0].paragraphs)>1
    # The acceptance cell also carries its existing descriptive title.
    assert ac['statement'] in ''.join(v.text for v in cells[0].paragraphs)
    assert p==before and not p.get('active_baseline_id')


def test_export_reserves_room_for_long_flow_labels_without_moving_the_flow():
    from PIL import Image,ImageFont
    from app.detail_rendering import flow_svg,flow_png,_flow_image
    from tests.test_detail_document import detail_fixture
    flow=detail_fixture()['flow'];flow['edges'][-1]['id']='ED-'+('long-connector-id-'*3)
    _,original=flow_svg(flow)
    svg,fitted=flow_svg(flow,fit_labels=True)
    assert fitted['width']>=original['width']
    assert fitted['boxes']==original['boxes'] and fitted['routes']==original['routes']
    assert flow['edges'][-1]['id'] in svg
    image,font_path,scale=_flow_image(fitted)
    font=ImageFont.truetype(str(font_path),round(fitted['edge_font_size']*scale))
    for route in fitted['routes']:
        for text in route['lines']:
            assert route['x']*scale+font.getlength(text)<image.width
    assert Image.open(io.BytesIO(flow_png(flow,fit_labels=True))).size==image.size
