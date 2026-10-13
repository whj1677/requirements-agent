"""Structured PRD projection, immutable flow images and formal export gates."""
import asyncio
import copy
import io
import xml.etree.ElementTree as ET

from docx import Document
from PIL import Image
import pytest

from app.contracts import document_gaps, validate_response
from app.core import Problem
from app.detail_rendering import detail_hash, detail_plaintext, flow_png, flow_svg, flow_layout, flow_png_pages
from app.document_reader import export_readiness, reader_document
from app.exports import document_files
from app.prd import compile_plan, plan_contract
from app.requirement_confirmation import row_fingerprint
from app.requirement_details import detail_issues
from app.workflow import apply_response
from tests.test_prd01 import state


def detail_fixture():
    return dict(schema_version=1,applicability={key:dict(applicable=key!='state_transitions',
        reason='联系人保存不引入独立业务生命周期' if key=='state_transitions' else '')
        for key in ('fields','interactions','flow','state_transitions')},
        fields=[dict(id='F-name',name='联系人姓名',meaning='电站联系人显示名称',data_type='text',
            source='用户输入',required='保存时必填',default_value='使用原姓名',constraints='1 至 30 个字符',
            visibility='电站成员可见',editability='仅管理员可编辑')],
        interactions=[dict(id='I-save',name='保存姓名',actor='管理员',entry='联系人详情',
            preconditions='具有当前电站编辑权限',trigger='点击保存',inputs='当前姓名',
            validation='姓名 1 至 30 个字符',pending='禁用保存按钮并显示处理中',
            repeat='同一请求只执行一次',cancel='提交前可取消，提交后等待结果',success='显示新姓名',
            failure='保留输入并提示原因',effects='只更新该联系人姓名',field_refs=['F-name'])],
        flow=dict(nodes=[dict(id='N-start',label='开始编辑',kind='start'),
                        dict(id='N-edit',label='修改联系人姓名',kind='action',actor='管理员'),
                        dict(id='N-check',label='姓名是否合法',kind='decision',actor='云平台'),
                        dict(id='N-save',label='保存并展示姓名',kind='action',actor='云平台'),
                        dict(id='N-end',label='完成',kind='end')],
                  edges=[dict(id='E1',**{'from':'N-start','to':'N-edit'}),
                         dict(id='E2',**{'from':'N-edit','to':'N-check'}),
                         dict(id='E3',condition='校验通过',**{'from':'N-check','to':'N-save'}),
                         dict(id='E4',condition='校验失败，保留输入并重新填写',**{'from':'N-check','to':'N-edit'}),
                         dict(id='E5',**{'from':'N-save','to':'N-end'})]),state_transitions=[])


def project_fixture(tmp_path,confirmed=True):
    _,p=state(tmp_path)
    p['product_context']={'product':'储能云平台'};p['questions']=[]
    req=p['items'][0];req['behavior']={};req['detail_contract']=detail_fixture()
    req['detail_required']=True
    if confirmed:
        req['requirement_confirmation']=dict(schema='requirement-confirmation-1',
            actor='authenticated_local_user',content_hash=row_fingerprint(p,req))
    return p


def compile_fixture(p):
    plan=dict(plan_version='4',sections=[dict(section_key='function',context_refs=[],
        normative_refs=plan_contract(p,'prd')['normative_item_ids'],discussion_refs=[],explanations=[])])
    result=compile_plan(plan,p,'prd',include_sketch=False)
    validate_response(result,'prd',p,p['sources'][0]['excerpts'])
    apply_response(p,result)
    p['documents']['prd']['sketch_policy']='excluded'
    return result


def test_structured_cloud_prd_uses_one_snapshot_without_legacy_seven_fields(tmp_path):
    p=project_fixture(tmp_path);before=copy.deepcopy(p)
    assert detail_issues(p['items'][0],p,require_confirmed=True)==[]
    assert plan_contract(p,'prd')['required_detail_requirement_ids']==[]
    assert plan_contract(p,'prd')['program_rendered_detail_ids']==['REQ-0001']
    compile_fixture(p)
    assert p['items']==before['items']
    assert not document_gaps(p['documents']['prd']['content'],p)
    assert export_readiness(p,'prd')['ready']
    view=reader_document(p['documents']['prd'])
    assert {'字段属性','交互逻辑','业务流程说明','状态变化'}<={s['title'] for s in view['sections']}
    nodes=[n for s in view['sections'] for n in s['reading_nodes']]
    tables=[n for n in nodes if n['kind']=='table'];diagram=next(n for n in nodes if n['kind']=='diagram')
    assert len([n for n in tables if n.get('detail_hash')])==4
    assert all(n['detail_hash']==detail_hash(p['items'][0]) for n in [*[n for n in tables if n.get('detail_hash')],diagram])
    assert diagram['image_uri'].startswith('data:image/svg+xml;base64,')
    assert any('重复操作' in '\n'.join(row) and '同一请求只执行一次' in '\n'.join(row) for t in tables for row in t['rows'])
    review=reader_document(p['documents']['prd'],include_locations=True)
    blocks=[b for s in review['sections'] for b in s['blocks'] if b['kind']=='detail']
    assert all(b.get('source_location') for b in blocks)
    assert any('1 至 30 个字符' in b['text'] and '电站成员可见' in b['text'] for b in blocks)
    assert any('校验失败，保留输入并重新填写' in b['text'] for b in blocks)


def test_word_real_tables_and_images_markdown_self_contained(tmp_path):
    p=project_fixture(tmp_path);compile_fixture(p)
    files=asyncio.run(document_files(p,'prd'))
    word=Document(io.BytesIO(files['PRD.docx']))
    assert len(word.tables)==5 and len(word.inline_shapes)>=1
    assert any('1 至 30 个字符' in cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert 'data:image/png;base64,' in files['PRD.md'].decode('utf-8')
    assert any(key.endswith('.svg') for key in files)
    png=next(value for key,value in files.items() if key.endswith('.png'))
    with Image.open(io.BytesIO(png)) as image:
        assert image.width>300 and image.height>300


@pytest.mark.parametrize('tamper',['snapshot','binding','missing'])
def test_changed_snapshot_or_missing_definition_cannot_export_in_either_mode(tmp_path,tamper):
    p=project_fixture(tmp_path);compile_fixture(p)
    artifact=p['documents']['prd']
    if tamper=='snapshot':
        artifact['item_snapshot'][0]['detail_contract']['fields'][0]['constraints']='改成 999 字符'
    elif tamper=='binding':
        next(b for s in artifact['content']['sections'] for b in s['blocks'] if b['kind']=='detail')['detail_hash']='0'*64
    else:
        artifact['content']['sections']=[s for s in artifact['content']['sections'] if s['title']!='字段属性']
    for mode in ('formal','review'):
        assert not export_readiness(p,'prd',mode)['ready']
        with pytest.raises(Problem):asyncio.run(document_files(p,'prd',mode=mode))


def test_unconfirmed_contract_can_only_be_exported_as_labelled_review(tmp_path):
    p=project_fixture(tmp_path,confirmed=False);compile_fixture(p)
    assert not export_readiness(p,'prd')['ready']
    assert export_readiness(p,'prd','review')['ready']
    with pytest.raises(Problem):asyncio.run(document_files(p,'prd'))
    files=asyncio.run(document_files(p,'prd',mode='review'))
    assert '非正式交付' in files['PRD.md'].decode('utf-8')
    assert '详细需求未确认' in files['PRD.md'].decode('utf-8')


def test_model_cannot_rewrite_structured_values(tmp_path):
    p=project_fixture(tmp_path)
    plan=dict(plan_version='4',sections=[dict(section_key='function',context_refs=[],normative_refs=['REQ-0001'],
        discussion_refs=[],explanations=[dict(text='姓名最多 999 个字符。',evidence_refs=dict(
            context_refs=[],normative_refs=['REQ-0001'],business_claim_refs=[]))])])
    with pytest.raises(Problem,match='程序直接成文'):compile_plan(plan,p,'prd')


def test_diagram_escapes_untrusted_labels_and_keeps_loop_conditions():
    flow=detail_fixture()['flow'];flow['nodes'][1]['label']='<script>alert(1)</script> & 处理'
    svg,layout=flow_svg(flow)
    ET.fromstring(svg)
    assert '<script>' not in svg and '&lt;script&gt;' in svg
    assert '重新填写' in svg and len(layout['routes'])==5
    assert flow_png(flow).startswith(b'\x89PNG')


def test_render_failure_blocks_review_instead_of_exporting_empty_picture(tmp_path):
    p=project_fixture(tmp_path)
    p['items'][0]['detail_contract']['flow']['edges'][0]['to']='MISSING'
    compile_fixture(p)
    result=export_readiness(p,'prd','review')
    assert not result['ready'] and any(i['code']=='DETAIL_RENDER_FAILED' for i in result['issues'])


def test_optional_business_properties_are_not_lost_in_review_projection():
    item=dict(id='REQ-1',title='字段定义',detail_contract=detail_fixture())
    item['detail_contract']['fields'][0].update(data_type='number',source='calculated',formula='输入值乘 2',
        unit='kW',precision='两位小数',empty_value='为空时不计算',uniqueness='本站范围',parent_id='F-parent')
    item['detail_contract']['flow']['nodes'][1]['interaction_id']='I-save'
    for expected in ('系统计算','数值','输入值乘 2','kW','两位小数','为空时不计算','本站范围','F-parent'):
        assert expected in detail_plaintext(item,'fields')
    assert 'I-save' in detail_plaintext(item,'flow')


def test_complex_branch_and_loop_routes_do_not_cross_node_interiors():
    from tests.test_detail_business_draft import sample_detail
    _,layout=flow_svg(sample_detail(complex_flow=True)['flow'])
    assert len(layout['routes'])==11
    for route in layout['routes']:
        for (x1,y1),(x2,y2) in zip(route['points'],route['points'][1:]):
            for box in layout['boxes'].values():
                left,right=box['x'],box['x']+box['width'];top,bottom=box['y'],box['y']+box['height']
                if y1==y2:
                    assert not (top<y1<bottom and max(min(x1,x2),left)<min(max(x1,x2),right))
                if x1==x2:
                    assert not (left<x1<right and max(min(y1,y2),top)<min(max(y1,y2),bottom))


def test_long_edge_conditions_do_not_expand_diagram_or_disappear_from_tables():
    import copy
    from tests.test_detail_business_draft import sample_detail
    item=dict(id='REQ-render-only',title='纯合成排版验证',detail_contract=sample_detail(complex_flow=True))
    flow=item['detail_contract']['flow']
    expected=flow_layout(flow)
    flow['edges'][0]['condition']='这是必须完整保留的实际业务条件，'*80
    before=copy.deepcopy(flow)
    svg,layout=flow_svg(flow)
    assert (layout['width'],layout['height'])==(expected['width'],expected['height'])
    assert all(route['lines']==[edge['id']] for route,edge in zip(layout['routes'],flow['edges']))
    assert flow['edges'][0]['condition'] in svg  # accessible title, not visible edge label
    assert flow['edges'][0]['condition'] in detail_plaintext(item,'flow')
    assert flow==before


def test_long_flow_stays_one_complete_image_with_every_node_and_edge(monkeypatch):
    import copy, io, re
    from PIL import Image, ImageDraw
    drawn_text=[];original_text=ImageDraw.ImageDraw.text
    def capture(draw,xy,text,*args,**kwargs):
        drawn_text.append(text)
        return original_text(draw,xy,text,*args,**kwargs)
    monkeypatch.setattr(ImageDraw.ImageDraw,'text',capture)
    flow=dict(nodes=[dict(id=f'N{i}',label='合成业务步骤'+str(i),kind='start' if i==1 else 'end' if i==13 else 'action',actor='系统')
                     for i in range(1,14)],
              edges=[dict(id=f'E{i}',**{'from':f'N{i}','to':f'N{i+1}'},condition='正常继续') for i in range(1,13)])
    flow['edges'].append(dict(id='E13',**{'from':'N12','to':'N3'},condition='返回已定义的步骤继续'))
    before=copy.deepcopy(flow)
    pages=flow_png_pages(flow,'合成流程')
    assert len(pages)==1
    rendered_nodes=[]
    for picture,caption in pages:
        rendered_nodes+=re.findall(r'\bN\d+\b',caption)
        image=Image.open(io.BytesIO(picture))
        assert image.height/image.width<1.35
    assert sorted(rendered_nodes)==sorted(n['id'] for n in flow['nodes'])
    assert all(edge['id'] in drawn_text for edge in flow['edges'])
    assert all(any(node['id']+' · ' in text for text in drawn_text) for node in flow['nodes'])
    assert flow==before
    layout=flow_layout(flow)
    assert {nid for layer in layout['layers'] for nid in layer['node_ids']}==set(rendered_nodes)
    # Every complete node stays inside the single canvas.
    assert all(layer['top']<=layout['boxes'][nid]['y'] and
               layout['boxes'][nid]['y']+layout['boxes'][nid]['height']<=layer['bottom']
               for layer in layout['layers'] for nid in layer['node_ids'])


def test_reciprocal_edges_remain_distinct_on_a_compact_single_canvas():
    flow=dict(nodes=[dict(id=f'N{i}',label='处理',kind='start' if i==1 else 'end' if i==8 else 'action') for i in range(1,9)],
              edges=[dict(id=f'E{i}',**{'from':f'N{i}','to':f'N{i+1}'}) for i in range(1,8)])
    flow['edges'] += [dict(id='R1',**{'from':'N3','to':'N2'}),dict(id='R2',**{'from':'N7','to':'N6'})]
    layout=flow_layout(flow)
    returns=[r for r in layout['routes'] if r['id'].startswith('R')]
    by_id={route['id']:route for route in layout['routes']}
    for reverse,forward in [('R1','E2'),('R2','E6')]:
        assert by_id[reverse]['points'] != list(reversed(by_id[forward]['points']))
    assert len(returns)==2 and len(layout['routes'])==len(flow['edges'])
    assert layout['width']<1300 and layout['height']/layout['width']<1
