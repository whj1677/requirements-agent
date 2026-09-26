"""Office source ingestion; all business text is an explicitly synthetic fixture."""
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app import sources
from app.core import Problem
from app.office import extract_office
from app.store import Store
from tests.test_runtime import client


def fixtures():
    word = Document()
    word.add_paragraph('合成样例：联系人备注。')
    word.add_table(rows=1, cols=2).rows[0].cells[1].text = '备注可为空。'
    excel = Workbook()
    excel.active.title = '合成联系人'
    excel.active['B2'] = '张三（合成）'
    excel.active['C2'] = '=1+2'
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = '合成演示：联系人备注'
    slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table.cell(0, 1).text = '备注可为空。'
    result = {}
    for ext, obj in [('.docx', word), ('.xlsx', excel), ('.pptx', deck)]:
        buf = io.BytesIO(); obj.save(buf); result[ext] = buf.getvalue()
    return result


@pytest.mark.parametrize('extension', ['.docx', '.xlsx', '.pptx'])
def test_normal_readers_use_real_contents_and_positions(tmp_path, monkeypatch, extension):
    monkeypatch.setattr(sources, 'extract_office', lambda *a: pytest.fail('Ordinary OOXML must not launch Office'))
    raw = fixtures()[extension]
    source = sources.save_source(Store(tmp_path), '合成' + extension, raw)
    assert source['parse_status'] in ('read', 'partial')
    assert source['sha256'] == hashlib.sha256(raw).hexdigest()
    assert all(x['source_hash'] == source['sha256'] and x['locator'] for x in source['excerpts'])
    text = '\n'.join(x['text'] for x in source['excerpts'])
    assert '合成' in text
    if extension == '.xlsx':
        assert '=1+2' in text and '未计算' in text
        assert any('B2' in x['locator'] and '合成联系人' in x['locator'] for x in source['excerpts'])
    if extension == '.pptx':
        assert '备注可为空。' in text and any('幻灯片 1' in x['locator'] for x in source['excerpts'])


def test_textless_image_deck_is_not_read(tmp_path, monkeypatch):
    from PIL import Image
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    image = io.BytesIO(); Image.new('RGB', (20, 20), 'white').save(image, format='PNG'); image.seek(0)
    slide.shapes.add_picture(image, 0, 0)
    raw = io.BytesIO(); deck.save(raw)
    monkeypatch.setattr(sources, 'extract_office', lambda *a: dict(status='office_required', rows=[], reason='没有正文'))
    source = sources.save_source(Store(tmp_path), 'image.pptx', raw.getvalue())
    assert source['parse_status'] == 'office_required' and not source['excerpts']


@pytest.mark.parametrize('extension', ['.doc','.docx','.xls','.xlsx','.ppt','.pptx'])
def test_office_fallback_preserves_source_and_reading_limits(tmp_path, monkeypatch, extension):
    def office(title, raw, folder):
        assert title.endswith(extension) and raw == b'synthetic-protected-content'
        return dict(status='partial', rows=[['工作表 / B2', '合成正文']], reason='图片未识别', method='local-office', visual_review='not_run')
    monkeypatch.setattr(sources, 'extract_office', office)
    source = sources.save_source(Store(tmp_path), 'sample'+extension, b'synthetic-protected-content')
    assert source['parse_status'] == 'partial'
    assert source['reading']['method'] == 'local-office'
    assert source['excerpts'][0]['method'] == 'local-office'
    assert source['excerpts'][0]['locator'].startswith('工作表 / B2')
    assert (tmp_path/'sources'/source['id']).read_bytes() == b'synthetic-protected-content'


@pytest.mark.parametrize('status', ['office_required','permission_denied'])
def test_office_failure_does_not_invent_excerpt(tmp_path, monkeypatch, status):
    monkeypatch.setattr(sources, 'extract_office', lambda *a: dict(status=status,rows=[],reason='合成限制'))
    source = sources.save_source(Store(tmp_path), 'protected.docx', b'synthetic')
    assert source['parse_status'] == status and source['excerpts'] == []


@pytest.mark.parametrize('rows', [[], [['段落1', '\ufffd乱码']], [['段落1', '']]])
def test_empty_or_broken_office_output_is_not_read(tmp_path, monkeypatch, rows):
    monkeypatch.setattr(sources, 'extract_office', lambda *a: dict(status='read',rows=rows,reason=''))
    s = sources.save_source(Store(tmp_path), 'protected.docx', b'synthetic')
    assert s['parse_status'] == 'failed' and not s['excerpts']


def test_rejected_package_does_not_bypass_limits_with_office(tmp_path, monkeypatch):
    monkeypatch.setattr(sources, 'extract_office', lambda *a: pytest.fail('Must not bypass package policy'))
    raw=io.BytesIO()
    with zipfile.ZipFile(raw, 'w') as package:
        package.writestr('word/vbaProject.bin', b'synthetic-only')
    s=sources.save_source(Store(tmp_path), 'macro.docx', raw.getvalue())
    assert s['parse_status'] == 'failed' and '宏' in s['failure_reason']


def test_non_office_failure_never_launches_office(tmp_path, monkeypatch):
    monkeypatch.setattr(sources, 'extract_office', lambda *a: pytest.fail('No Office for PDF'))
    assert sources.save_source(Store(tmp_path), 'bad.pdf', b'bad')['parse_status'] == 'failed'


def test_upload_retry_persists_old_source_and_requires_new_external_grant(tmp_path, monkeypatch):
    c, app = client(tmp_path)
    with c:
        p = c.post('/api/projects', json={'name':'合成 Office 导入'}).json()
        root = '/api/projects/' + p['id']
        monkeypatch.setattr(sources, 'extract_office', lambda *a: dict(status='office_required',rows=[],reason='需要 Office'))
        s = c.post(root+'/sources/file', data={'expected_revision':0}, files={'file':('sample.docx',b'synthetic')}).json()
        assert s['parse_status'] == 'office_required'
        monkeypatch.setattr(sources, 'extract_office', lambda *a: dict(status='partial',rows=[['段落 1','合成正文']],reason='图片未读'))
        assert c.post(root+'/sources/'+s['id']+'/retry', json={'expected_revision':0}).status_code == 409
        result=c.post(root+'/sources/'+s['id']+'/retry', json={'expected_revision':1})
        assert result.status_code == 200
        new=result.json(); p=c.get(root).json()
        assert new['parent_source_id'] == s['id'] and new['version'] == 2
        assert p['sources'][0]['parse_status'] == 'office_required'
        assert p['sources'][1]['parse_status'] == 'partial'
        assert not p['items'] and not p['baselines']
        assert not c.get(root+'/runs').json()
        assert set(['.doc','.docx','.xls','.xlsx','.ppt','.pptx']) <= set(p['source_capabilities']['extensions'])


def test_office_timeout_cleanup_is_scoped(tmp_path, monkeypatch):
    from app import office
    calls=[]
    def run(cmd, **kwargs):
        calls.append(cmd)
        if '-OwnerPath' in cmd:
            Path(cmd[cmd.index('-OwnerPath')+1]).write_text('{}')
            raise subprocess.TimeoutExpired(cmd, 90)
        return subprocess.CompletedProcess(cmd, 0)
    monkeypatch.setattr(office.subprocess, 'run', run)
    if os.name != 'nt': pytest.skip('Windows cleanup command')
    result=extract_office('a.docx', b'synthetic', tmp_path)
    assert result['code'] == 'OFFICE_TIMEOUT'
    assert '-CleanupOwner' in calls[-1] and len(calls) == 2
    assert not list(tmp_path.iterdir())


def test_office_fallback_finds_system_powershell_without_developer_path(tmp_path, monkeypatch):
    from app import office
    if os.name != 'nt': pytest.skip('Windows PowerShell discovery')
    monkeypatch.setenv('ProgramFiles', str(tmp_path / 'no-powershell-seven'))
    monkeypatch.setenv('SystemRoot', str(tmp_path / 'Windows'))
    monkeypatch.setattr(office.shutil, 'which', lambda _name: None)
    commands = []
    def run(command, **kwargs):
        commands.append(command)
        return subprocess.CompletedProcess(command, 1)
    monkeypatch.setattr(office.subprocess, 'run', run)
    result = extract_office('synthetic.docx', b'synthetic', tmp_path)
    assert commands[0][0] == str(tmp_path / 'Windows/System32/WindowsPowerShell/v1.0/powershell.exe')
    assert '-File' in commands[0] and '-ExecutionPolicy' not in commands[0]
    assert result['code'] == 'OFFICE_NO_RESULT'


@pytest.mark.skipif(os.environ.get('RA_TEST_LOCAL_OFFICE') != '1', reason='Explicit opt-in: starts local installed Office with synthetic documents')
@pytest.mark.parametrize('extension', ['.docx','.xlsx','.pptx'])
def test_installed_office_reads_synthetic_document(tmp_path, extension):
    raw=fixtures()[extension]
    result=extract_office('synthetic'+extension, raw, tmp_path)
    assert result['status'] == 'partial', {k:v for k,v in result.items() if k!='rows'}
    assert any('合成' in text for _,text in result['rows'])
    assert result['visual_review'] == 'not_run'
    assert not list(tmp_path.iterdir())


@pytest.mark.skipif(os.environ.get('RA_TEST_LOCAL_OFFICE') != '1', reason='Explicit opt-in: relative data-directory COM regression')
def test_installed_office_relative_data_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    Path('relative-data').mkdir()
    result=extract_office('relative.docx', fixtures()['.docx'], Path('relative-data'))
    assert result['status']=='partial', {k:v for k,v in result.items() if k!='rows'}
    assert any('合成样例' in text for _,text in result['rows'])
