"""Native COM safety and real opt-in synthetic document extraction."""
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace

import pytest

from app import office, office_worker, office_process

pytestmark = pytest.mark.skipif(os.name != 'nt', reason='Windows native COM')


def test_native_runtime_probe_isolated_and_has_no_office_requirement(tmp_path, monkeypatch):
    monkeypatch.setenv('PATH', str(tmp_path / 'no-programs'))
    assert office.probe_runtime(tmp_path)
    assert not list(tmp_path.iterdir())


def test_frozen_native_worker_reenters_same_executable(monkeypatch):
    monkeypatch.setattr(sys, 'frozen', True, raising=False)
    assert office.worker_command() == [sys.executable, '--office-worker']


def test_native_child_entry_never_starts_customer_server(tmp_path, monkeypatch):
    from app import desktop_runtime
    monkeypatch.setattr(desktop_runtime, 'server', lambda: pytest.fail('Worker must not start a server'))
    monkeypatch.setattr(desktop_runtime, 'launcher', lambda: pytest.fail('Worker must not open UI'))
    output = tmp_path / 'probe.json'
    assert desktop_runtime.main(['--office-worker', '--probe-runtime', '--output', str(output)]) == 0
    assert json.loads(output.read_text('utf-8'))['powershell_used'] is False


@pytest.mark.parametrize('readonly,denied', [(False, False), (True, True)])
def test_readonly_and_permission_failures_never_produce_excerpts(tmp_path, monkeypatch, readonly, denied):
    events = []
    class PolicyDenied(Exception):
        hresult = -2147352567
        excepinfo = (0, 'private-app', 'private-document-body', None, 0, -2147024891)
    document = SimpleNamespace(ReadOnly=readonly, Close=lambda *_: events.append('close'))
    def open_document(*args):
        assert args[1:4] == (False, True, False)
        if denied:
            raise PolicyDenied()
        return document
    app = SimpleNamespace(Documents=SimpleNamespace(Count=0, Open=open_document),
                          Options=SimpleNamespace(UpdateLinksAtOpen=True), Quit=lambda *_: events.append('quit'))
    com = SimpleNamespace(COINIT_APARTMENTTHREADED=2, CLSCTX_LOCAL_SERVER=4, IID_IDispatch='dispatch',
                          CoInitializeEx=lambda _: None, CoUninitialize=lambda: events.append('uninitialize'),
                          CoCreateInstance=lambda *_: app)
    monkeypatch.setattr(office_worker, 'com_runtime', lambda: (com, SimpleNamespace(Dispatch=lambda value: value)))
    snapshots = iter([[], [dict(pid=123, name='winword', image='synthetic', started='1')]])
    monkeypatch.setattr(office_worker, 'processes', lambda _: next(snapshots))
    result = office_worker.extract(tmp_path / 'synthetic.docx', tmp_path / 'owner.json', 'Word.Application', 'winword')
    assert result['code'] == ('OFFICE_PERMISSION_DENIED' if denied else 'OFFICE_NOT_READONLY')
    assert result['rows'] == [] and result['failure_stage'] == 'document_open_or_extract'
    assert 'private-' not in json.dumps(result)
    assert app.AutomationSecurity == 3 and app.Options.UpdateLinksAtOpen is True
    assert events[-2:] == ['quit', 'uninitialize']


def test_cleanup_requires_matching_owned_handle_creation_stamp_and_image(tmp_path, monkeypatch):
    executable = Path(getattr(sys, '_base_executable', sys.executable))
    process = subprocess.Popen([str(executable), '-c', 'import time; time.sleep(30)'],
                               creationflags=subprocess.CREATE_NO_WINDOW)
    try:
        name = executable.stem.casefold()
        identity = next(p for p in office_process.processes(name) if p['pid'] == process.pid)
        # Only the one process created here is eligible in this test process.
        monkeypatch.setattr(office_process, 'ALLOWED', {name})
        owner = tmp_path / 'owner.json'
        owner.write_text(json.dumps({**identity, 'started': '0'}), 'utf-8')
        assert office_process.cleanup_owned(owner) is False and process.poll() is None
        owner.write_text(json.dumps(identity), 'utf-8')
        assert office_process.cleanup_owned(owner) is True
        process.wait(timeout=5)
    finally:
        if process.poll() is None:
            process.terminate()
            process.wait(timeout=5)


@pytest.mark.parametrize('exits', [True, False])
def test_cleanup_identity_query_failure_only_succeeds_after_observed_exit(monkeypatch, exits):
    waits = iter([258, 0 if exits else 258])
    kernel = SimpleNamespace(OpenProcess=lambda *_: 123,
        WaitForSingleObject=lambda *_: next(waits), CloseHandle=lambda *_: None,
        TerminateProcess=lambda *_: pytest.fail('Never terminate when identity could not be verified'))
    monkeypatch.setattr(office_process, 'kernel_api', lambda: kernel)
    def denied(*_):
        raise PermissionError('Synthetic process is exiting or inaccessible')
    monkeypatch.setattr(office_process, 'handle_identity', denied)
    owner = dict(pid=12345, name='wps', image='synthetic', started='1')
    if exits:
        assert office_process.cleanup_identity(owner) is True
    else:
        with pytest.raises(PermissionError):
            office_process.cleanup_identity(owner)


@pytest.mark.skipif(os.environ.get('RA_TEST_LOCAL_OFFICE') != '1', reason='Explicit opt-in: synthetic Word image')
def test_native_word_retains_inline_image_extraction_and_original_bytes(tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    image = io.BytesIO()
    Image.new('RGB', (40, 30), 'green').save(image, 'PNG')
    image.seek(0)
    document = Document()
    document.add_paragraph('原生读取合成图片样例')
    document.add_picture(image)
    raw = io.BytesIO()
    document.save(raw)
    original = raw.getvalue()
    digest = hashlib.sha256(original).hexdigest()
    monkeypatch.setenv('PATH', str(tmp_path / 'no-powershell'))
    result = office.extract_office('synthetic.docx', original, tmp_path)
    assert result['status'] == 'partial', result
    assert result['engine'] == 'native-com' and result['owned_process_cleanup'] == 'verified'
    assert len(result['images']) == 1, result['reason']
    assert hashlib.sha256(original).hexdigest() == digest
    assert result['visual_review'] == 'not_run' and not list(tmp_path.iterdir())
