"""Read-view hierarchy fixtures; not generated business documents or approvals."""
import asyncio
import copy
import io

from docx import Document
from app.document_reader import reader_document
from app.exports import document_files
from app.prd import compile_plan
from app.store import Store
from app.workflow import apply_response
from tests.helpers import prepared


def hierarchy_project(tmp_path):
    p = prepared(Store(tmp_path))
    item = p['items'][0]
    item.update(title='明确标记的层级夹具：跨页面维护操作及权限保持的长需求标题',
                behavior=dict(actor='管理员', entry='既有列表入口', flow='先核对，再保存。',
                              data='原有字段保持。', permissions='权限规则保持。',
                              result='结果以原文为准。', exceptions='未知分支仍未定义。'))
    plan = dict(plan_version='1', title='层级工程夹具', limitations=[], sections=[dict(
        title='并列主题甲', normative_refs=[i['id'] for i in p['items']], discussion_refs=[],
        narration='REQ 是此处的普通短句。'),dict(title='并列主题乙', normative_refs=[], discussion_refs=[], narration='短句。')])
    apply_response(p, compile_plan(plan, p, 'prd', include_sketch=False))
    artifact=p['documents']['prd'];artifact['sketch_policy']='excluded'
    first=artifact['content']['sections'][0]
    child=dict(section_id='PRD-explicit-child',title='原已存在的子章',level=2,
               parent_section_id=first['section_id'],blocks=[dict(kind='narrative',ref_ids=[],text='普通说明，不是标题。')])
    artifact['content']['sections'].insert(1,child)
    return p


def test_hierarchy_uses_explicit_objects_and_never_promotes_short_prose(tmp_path):
    p=hierarchy_project(tmp_path);artifact=p['documents']['prd'];before=copy.deepcopy(artifact)
    view=reader_document(artifact);outline=view['outline'];sections=view['sections']
    assert artifact==before
    assert view['presentation_version']=='document-reading-2'
    req=next(n for n in outline if n.get('item_id')=='REQ-0001')
    assert req['level']==2 and req['parent_id']==sections[0]['section_id']
    actor=next(n for n in outline if n['text']=='谁操作')
    assert actor['level']==3 and actor['parent_id']==req['id']
    child=next(n for n in outline if n['id']=='PRD-explicit-child')
    assert child['level']==2 and child['parent_id']==sections[0]['section_id']
    sibling=next(n for n in outline if n['text']=='并列主题乙')
    assert sibling['level']==1 and sibling['parent_id'] is None
    assert len({n['number'] for n in outline})==len(outline)
    assert not any(n['text'] in ('短句。','REQ 是此处的普通短句。') for n in outline)
    nodes=[n for s in sections for n in s['reading_nodes']]
    for item in p['items']:
        assert any(n['kind']=='paragraph' and n['text']==item['statement'] for n in nodes)
    assert any(n.get('source_refs')==p['items'][0]['source_refs'] for n in nodes if n['kind']=='metadata')


def test_long_id_missing_snapshot_and_unknown_keep_their_identity(tmp_path):
    p=hierarchy_project(tmp_path);a=p['documents']['prd'];item=a['item_snapshot'][0]
    long_id='REQ-'+('LONG-'*18)+'0001';old=item['id'];item['id']=long_id
    for s in a['content']['sections']:
        for b in s['blocks']:b['ref_ids']=[long_id if r==old else r for r in b['ref_ids']]
    view=reader_document(a)
    nodes=[n for s in view['sections'] for n in s['reading_nodes']]
    assert any(n.get('item_id')==long_id and n['kind']=='metadata' for n in nodes)
    assert any(p['questions'][0]['question'] in n['text'] for n in nodes)
    a['item_snapshot']=[]
    missing=reader_document(a)
    assert not any(n.get('item_id') for n in missing['outline'])
    assert any(long_id in n['text'] and '未保存' in n['text'] for s in missing['sections'] for n in s['reading_nodes'])


def test_markdown_word_and_web_share_heading_tree_and_exact_clauses(tmp_path):
    p=hierarchy_project(tmp_path);before=copy.deepcopy(p);view=reader_document(p['documents']['prd'])
    files=asyncio.run(document_files(p,'prd',reader=True))
    md=files['PRD.md'].decode();doc=Document(io.BytesIO(files['PRD.docx']))
    headings=[(x.text,x.style.name) for x in doc.paragraphs if x.style.name.startswith('Heading')]
    assert headings==[(n['number']+' '+n['text'],'Heading '+str(n['level'])) for n in view['outline']]
    for n in view['outline']:assert '#'*(n['level']+1)+' '+n['number']+' '+n['text'] in md
    for item in p['items']:
        assert item['statement'] in md and any(x.text==item['statement'] for x in doc.paragraphs)
    assert doc.paragraphs[0].style.name=='Title'
    assert doc.styles['Heading 1'].font.size.pt==16
    assert doc.styles['Normal'].font.size.pt==11
    assert all(x.paragraph_format.keep_with_next for x in doc.paragraphs
               if x.style.name=='Document Metadata' and ' · v' in x.text)
    assert p==before and not p.get('active_baseline_id')
