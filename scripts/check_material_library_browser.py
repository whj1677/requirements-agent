"""Material grouping, screenshot reference and original preview/download UI.

Only synthetic data, an isolated temporary HTTP app and a short-lived browser.
"""
import argparse
import asyncio
import io
import json
import os
from pathlib import Path
import socket
import sys
import tempfile
import threading

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
ISOLATED = Path(tempfile.mkdtemp(prefix='ra-material-library-'))
os.environ['RA_DATA_DIR'] = str(ISOLATED / 'module-data')
import app.config
app.config.ROOT = ISOLATED
from app.main import create_app
from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
from PIL import Image
from playwright.async_api import async_playwright, expect
from starlette.staticfiles import StaticFiles
import uvicorn


async def run(dist, evidence):
    evidence.mkdir(parents=True, exist_ok=True)
    result = {'status': 'running', 'cases': [], 'page_errors': [], 'model_calls': 0}
    class NoModel:
        async def request(self, *args, **kwargs):
            result['model_calls'] += 1
            raise AssertionError('Material browsing must not call a model')
        def key_status(self, config):
            return {'key_configured': False, 'key_source': 'none', 'bound_origin': 'synthetic'}
    app = create_app(ISOLATED / 'case', access_token='synthetic-token', provider=NoModel(), env_path=ISOLATED / 'absent.env')
    app.router.routes[:] = [r for r in app.router.routes if getattr(r, 'name', '') != 'web']
    app.mount('/', StaticFiles(directory=dist, html=True), name='web')
    image_a = io.BytesIO(); Image.new('RGB', (500, 280), '#265b43').save(image_a, format='PNG')
    image_b = io.BytesIO(); Image.new('RGB', (320, 180), '#b45d32').save(image_b, format='PNG')
    document = Document()
    document.add_paragraph('合成现状：已有联系人列表。此原文必须可以在附件中核对。')
    document.add_paragraph('当前流程：使用者进入客户管理，在联系人列表中按姓名或手机号查找，再打开详情查看所属企业和联系记录。')
    document.add_paragraph('本次诉求：减少重复查找，希望在列表中直接看到联系人所属企业。现有查看与编辑权限保持不变。')
    document.add_paragraph('待核对事项：企业名称过长时如何展示、缺少企业信息时的提示，以及历史数据是否需要补录。')
    document.add_paragraph('第三段含有连续字符以检查窄屏换行：' + ('longtextwithoutspaces-' * 35))
    document.add_picture(io.BytesIO(image_a.getvalue()), width=Inches(4))
    document.add_paragraph('两张内嵌图之间的正文。')
    document.add_picture(io.BytesIO(image_b.getvalue()), width=Inches(3))
    for number in range(60):
        document.add_paragraph(f'补充材料第 {number + 1} 段：合成验收内容，保持原文顺序与来源位置。')
    raw = io.BytesIO(); document.save(raw)
    thread = server = sock = None
    try:
        with TestClient(app) as client:
            token = client.post('/api/session/login', json={'key': 'synthetic-token'}).json()['csrf']
            client.headers.update({'Origin': 'http://testserver', 'X-CSRF-Token': token})
            project = client.post('/api/projects', json={'name': '材料分类与预览验收'}).json()
            root = '/api/projects/' + project['id']
            goal = client.post(root + '/sources/text', json={'expected_revision': project['revision'], 'title': '本期目标', 'text': '合成目标：减少查找时间。', 'purpose': 'goal'}).json()
            revision = client.get(root).json()['revision']
            response = client.post(root + '/sources/file', data={'expected_revision': revision, 'purpose': 'current'}, files={'file': ('现状说明.docx', raw.getvalue())})
            assert response.status_code == 200, response.text
            word = response.json()
            assert len(word['excerpts']) > 55
            deep_excerpt = word['excerpts'][55]
            current = app.state.store.get(project['id'])
            with app.state.store.edit(project['id'], current['revision'], '合成深引用验收条目') as (draft, _):
                draft['items'].append(dict(id='REQ-0001', kind='requirement', title='深引用定位合成条目',
                    statement='仅用于验证点击查看来源定位到第五十段之后，不是模型结论。', applies_to='to_be',
                    epistemic_status='reported', change_type='new', selection_status='selected', revision=1,
                    source_refs=[dict(source_id=word['id'], excerpt_id=deep_excerpt['id'])], related_refs=[]))
            sock = socket.socket(); sock.bind(('127.0.0.1', 0))
            server = uvicorn.Server(uvicorn.Config(app, log_level='error', access_log=False))
            thread = threading.Thread(target=lambda: server.run(sockets=[sock]), daemon=True); thread.start()
            for _ in range(100):
                if server.started: break
                await asyncio.sleep(.05)
            assert server.started
            async with async_playwright() as pw:
                browser = await pw.chromium.launch(headless=True)
                page = await browser.new_page(viewport={'width': 1440, 'height': 1000})
                page.on('pageerror', lambda error: result['page_errors'].append(str(error)))
                await page.goto(f'http://127.0.0.1:{sock.getsockname()[1]}/')
                await page.get_by_label('本机访问令牌').fill('synthetic-token')
                await page.get_by_role('button', name='进入工作台').click()
                await page.get_by_role('button', name='材料分类与预览验收', exact=True).click()
                await expect(page.locator('.intake-material-group').filter(has_text='目标与诉求')).to_contain_text('本期目标')
                await expect(page.locator('.intake-material-group').filter(has_text='现状资料')).to_contain_text('现状说明.docx')
                assert await page.locator('.intake-material-group .material-file').count() == 2
                result['cases'].append('intake separates goals/current and nests embedded images')
                await page.get_by_role('button', name='项目资料 ·', exact=False).click()
                drawer = page.get_by_role('dialog', name='项目资料', exact=True)
                async with page.expect_file_chooser() as chooser:
                    await drawer.get_by_role('button', name='导入墨刀原型截图', exact=True).click()
                await (await chooser.value).set_files({'name': '墨刀-联系人.png', 'mimeType': 'image/png', 'buffer': image_a.getvalue()})
                await expect(drawer.get_by_role('button', name='查看内容：墨刀-联系人.png', exact=True)).to_be_visible()
                screenshot = next(s for s in client.get(root).json()['sources'] if s['title'] == '墨刀-联系人.png')
                assert screenshot['purpose'] == 'reference' and screenshot['parse_status'] == 'awaiting_vision'
                await expect(drawer.locator('.source-category').filter(has=page.get_by_role('heading', name='设计参考', exact=False))).to_contain_text('墨刀-联系人.png')
                result['cases'].append('Modao screenshot imports as reference without claiming visual understanding')
                opener = drawer.get_by_role('button', name='查看内容：现状说明.docx', exact=True)
                await opener.click()
                preview = page.get_by_role('dialog', name='附件内容', exact=True)
                await expect(preview).to_contain_text('此原文必须可以在附件中核对')
                await expect(preview.locator('.preview-reading-area')).to_be_visible()
                await expect(preview.locator('.material-excerpt')).to_have_count(50)
                await expect(preview.locator('.material-excerpt small')).to_have_count(0)
                await expect(preview.get_by_role('button', name='正文', exact=True)).to_be_visible()
                assert '个已提取片段' not in await preview.inner_text()
                assert word['excerpts'][0]['locator'] not in await preview.locator('.preview-paper').inner_text()
                long_excerpt = preview.locator('.material-excerpt').filter(has_text='longtextwithoutspaces-').locator('pre')
                await expect(long_excerpt).to_contain_text('longtextwithoutspaces-')
                body_tab = preview.locator('.preview-views button').first
                images_tab = preview.get_by_role('button', name='文档内图片 · 2', exact=True)
                await expect(body_tab).to_have_attribute('aria-pressed', 'true')
                await expect(images_tab).to_have_attribute('aria-pressed', 'false')
                await expect(preview.get_by_role('heading', name='文档内图片 · 2')).to_have_count(0)

                # Focus must remain in the top dialog, Escape closes only it, and focus returns to its opener.
                await expect(preview.get_by_role('button', name='关闭附件', exact=True)).to_be_focused()
                await page.keyboard.press('Shift+Tab')
                await expect(preview.locator('summary').last).to_be_focused()
                await page.keyboard.press('Tab')
                await expect(preview.get_by_role('button', name='关闭附件', exact=True)).to_be_focused()
                await page.keyboard.press('Escape')
                await expect(preview).to_have_count(0)
                await expect(opener).to_be_focused()
                await opener.click()
                preview = page.get_by_role('dialog', name='附件内容', exact=True)
                await preview.get_by_role('button', name='继续显示内容', exact=False).click()
                await expect(preview.locator('.material-excerpt')).to_have_count(len(word['excerpts']))
                await expect(preview.locator('.material-excerpt').last.locator('pre')).to_have_text(word['excerpts'][-1]['text'])
                await preview.get_by_role('region', name='已读取内容', exact=True).evaluate('(el)=>el.scrollTo(0,0)')
                async with page.expect_download() as download_info:
                    await preview.get_by_role('button', name='下载原始附件', exact=True).click()
                download = await download_info.value
                await download.save_as(evidence / 'original.docx')
                assert (evidence / 'original.docx').read_bytes() == raw.getvalue()
                await page.get_by_role('button', name='关闭提示', exact=True).click()
                await page.screenshot(path=str(evidence / 'word-preview-desktop.png'))
                await expect(images_tab).to_have_attribute('aria-pressed', 'false')
                await images_tab.click()
                await expect(images_tab).to_have_attribute('aria-pressed', 'true')
                await expect(preview.get_by_role('heading', name='文档内图片 · 2')).to_be_visible()
                embedded_images = preview.locator('.material-image-view img')
                await expect(embedded_images).to_have_count(2)
                await expect(embedded_images.nth(0)).to_be_visible()
                await expect(embedded_images.nth(1)).to_be_visible()
                assert await embedded_images.nth(0).evaluate('(img)=>img.naturalWidth') == 500
                assert await embedded_images.nth(1).evaluate('(img)=>img.naturalWidth') == 320
                child_id = next(s['id'] for s in client.get(root).json()['sources'] if s.get('container_source_id') == word['id'])
                failures = {'count': 0}
                async def fail_once(route):
                    failures['count'] += 1
                    if failures['count'] == 1:
                        await route.abort()
                    else:
                        await route.continue_()
                await page.route(f'**/sources/{child_id}/image*', fail_once)
                # Re-open once with a forced initial image failure, then retry and verify recovery.
                await preview.get_by_role('button', name='关闭附件').click()
                await drawer.get_by_role('button', name='查看内容：现状说明.docx', exact=True).click()
                preview = page.get_by_role('dialog', name='附件内容', exact=True)
                await preview.get_by_role('button', name='文档内图片 · 2', exact=True).click()
                failed_image = preview.locator('.material-image-view img').first
                await expect(failed_image).to_have_js_property('complete', True)
                await expect(preview.get_by_role('alert')).to_contain_text('图片暂未取得')
                await preview.get_by_role('button', name='重新加载图片', exact=True).click()
                await expect(failed_image).to_be_visible()
                await expect(failed_image).to_have_js_property('naturalWidth', 500)
                assert failures['count'] >= 2
                await page.unroute(f'**/sources/{child_id}/image*', fail_once)
                await page.screenshot(path=str(evidence / 'word-preview-images-desktop.png'))
                result['cases'].append('Word body/image tabs, two embedded images, failed image retry, nested keyboard focus, and byte-identical original download')
                await preview.get_by_role('button', name='关闭附件').click()
                # Reclassify a parent and verify the evidence identity and child purpose.
                card = drawer.locator('.material-card').filter(has=page.get_by_role('button', name='查看内容：现状说明.docx', exact=True))
                await card.get_by_text('分类与读取操作', exact=True).click()
                await card.get_by_label('资料分类：现状说明.docx', exact=True).select_option('goal')
                await expect(drawer.locator('.source-category').filter(has=page.get_by_role('heading', name='目标与诉求', exact=False))).to_contain_text('现状说明.docx')
                current = client.get(root).json()
                parent = next(s for s in current['sources'] if s['id'] == word['id'])
                assert parent['purpose'] == 'goal' and parent['sha256'] == word['sha256'] and parent['excerpts'] == word['excerpts']
                assert all(s['purpose'] == 'goal' for s in current['sources'] if s.get('container_source_id') == word['id'])
                result['cases'].append('manual reclassification synchronizes children and preserves evidence')
                await page.screenshot(path=str(evidence / 'library-desktop.png'))
                await drawer.get_by_role('button', name='查看内容：墨刀-联系人.png', exact=True).click()
                await expect(preview.get_by_role('img', name='墨刀-联系人.png', exact=True)).to_be_visible()
                assert await preview.get_by_role('img', name='墨刀-联系人.png', exact=True).evaluate('(img)=>img.naturalWidth') == 500
                await expect(preview.locator('.preview-views button').filter(has_text='正文')).to_have_count(0)
                await expect(preview.get_by_role('region', name='图片预览', exact=True)).to_be_visible()
                await expect(preview.locator('.material-excerpt')).to_have_count(0)
                await preview.get_by_role('button', name='查看原图尺寸').click()
                await expect(preview.get_by_role('button', name='适应窗口')).to_be_visible()
                result['cases'].append('screenshot opens and supports original-size viewing')
                await preview.get_by_role('button', name='关闭附件').click()
                await page.set_viewport_size({'width': 390, 'height': 844})
                if await page.get_by_role('button', name='关闭提示', exact=True).count():
                    await page.get_by_role('button', name='关闭提示', exact=True).click()
                await page.screenshot(path=str(evidence / 'library-mobile.png'))
                await drawer.get_by_role('button', name='关闭项目资料').click()
                # First-screen attachment itself opens the corresponding content.
                await page.get_by_role('button', name='墨刀-联系人.png', exact=False).click()
                await expect(preview.get_by_role('img', name='墨刀-联系人.png', exact=True)).to_be_visible()
                assert await page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                assert await preview.evaluate('(el)=>el.scrollWidth <= el.clientWidth')
                await page.screenshot(path=str(evidence / 'image-preview-mobile.png'))
                await preview.get_by_role('button', name='关闭附件').click()
                # Opening the first-screen attachment also opens its source drawer; it remains after the preview closes.
                drawer = page.get_by_role('dialog', name='项目资料', exact=True)
                await drawer.get_by_role('button', name='查看内容：现状说明.docx', exact=True).click()
                preview = page.get_by_role('dialog', name='附件内容', exact=True)
                assert await page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                assert await preview.evaluate('(el)=>el.scrollWidth <= el.clientWidth')
                assert await preview.locator('.preview-reading-area').evaluate('(el)=>el.scrollWidth <= el.clientWidth')
                assert await preview.locator('.material-excerpt').filter(has_text='longtextwithoutspaces-').locator('pre').evaluate('(el)=>el.scrollWidth <= el.clientWidth')
                await page.screenshot(path=str(evidence / 'word-preview-mobile.png'))
                result['cases'].append('390px image and Word previews fit without horizontal overflow; screenshots captured')
                await preview.get_by_role('button', name='关闭附件').click()
                await drawer.get_by_role('button', name='关闭项目资料').click()
                await page.set_viewport_size({'width': 1440, 'height': 1000})
                await page.get_by_role('button', name='核对现状、价值与改动范围', exact=False).click()
                requirement = page.locator('.scope-requirement').filter(has_text='深引用定位合成条目')
                await requirement.locator('summary').click()
                await requirement.get_by_role('button', name='查看来源', exact=True).click()
                preview = page.get_by_role('dialog', name='附件内容', exact=True)
                target = preview.locator('.material-excerpt.source-highlight')
                await expect(target).to_have_attribute('id', 'excerpt-' + deep_excerpt['id'])
                await expect(target.locator('pre')).to_have_text(deep_excerpt['text'])
                await expect(target).to_be_in_viewport()
                assert await preview.locator('.material-excerpt').count() >= 56
                await preview.get_by_role('button', name='文档内图片 · 2', exact=True).click()
                await expect(preview.get_by_role('region', name='图片预览', exact=True)).to_be_visible()
                await preview.locator('.preview-views button').first.click()
                await expect(target).to_be_in_viewport()
                await page.screenshot(path=str(evidence / 'deep-reference.png'))
                await preview.get_by_role('button', name='关闭附件', exact=True).click()
                await page.get_by_role('dialog', name='项目资料', exact=True).get_by_role('button', name='关闭项目资料').click()
                result['cases'].append('load more preserves all extracted text; genuine source-reference navigation reveals and scrolls to excerpt 56, including after image/body switching')
                assert not result['page_errors'] and result['model_calls'] == 0
                await browser.close()
                result['status'] = 'passed'
    except Exception as error:
        result['status'] = 'failed'; result['error'] = str(error)
        raise
    finally:
        if server: server.should_exit = True
        if thread: await asyncio.to_thread(thread.join, 5)
        if sock: sock.close()
        result['listener_closed'] = not thread or not thread.is_alive()
        (evidence / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--dist', type=Path, required=True); parser.add_argument('--evidence', type=Path, required=True)
    args = parser.parse_args(); asyncio.run(run(args.dist.resolve(), args.evidence.resolve()))
