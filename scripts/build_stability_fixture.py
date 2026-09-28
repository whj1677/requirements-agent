"""Build the synthetic contacts acceptance materials; never calls a model."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from playwright.sync_api import sync_playwright


ROOT = Path(__file__).resolve().parent.parent
FIXTURE = ROOT / "tests" / "fixtures" / "stability_contacts_v1.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def make_docx(path: Path, fixture: dict) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    doc.add_heading("联系人备注：当前状态与保持约束", 0)
    doc.add_paragraph("材料类型：合成验收示意输入；不代表真实客户、真实业务或已批准需求。")
    doc.add_heading("当前联系人管理事实", level=1)
    doc.add_paragraph("联系人列表现有字段：" + "、".join(fixture["current_state"]["contact_list_fields"]) + "。")
    doc.add_paragraph("管理员可新增、编辑联系人；只读用户仅可查看联系人。")
    doc.add_heading("需保持的现有行为", level=1)
    doc.add_paragraph("查找、排序、分页和删除行为保持原状。")
    doc.add_heading("输入说明", level=1)
    doc.add_paragraph("本文件只记录联系人管理的当前状态和保持约束，不包含本期变更答案或产品决定。")
    doc.core_properties.title = "联系人备注合成验收：当前状态"
    doc.core_properties.subject = "synthetic acceptance input"
    doc.save(path)


def make_reference_html(path: Path) -> None:
    html = '''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>联系人列表示意参考</title><style>
*{box-sizing:border-box}body{margin:0;background:#f2f5f8;color:#182230;font:15px "Microsoft YaHei",sans-serif}
.shell{max-width:1100px;margin:48px auto;padding:0 24px}.banner{padding:12px 16px;background:#fff3cd;border:1px solid #e8cd78;border-radius:8px;color:#604d0c;font-weight:700}
h1{font-size:27px;margin:24px 0 6px}.sub{color:#657386;margin:0 0 22px}.panel{background:white;border:1px solid #dde4ec;border-radius:12px;overflow:hidden;box-shadow:0 8px 24px #2132470d}
.toolbar{padding:16px 18px;border-bottom:1px solid #e8edf2;display:flex;justify-content:space-between;align-items:center}.button{background:#1769aa;color:white;padding:9px 15px;border-radius:7px;font-weight:600}
table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{text-align:left;padding:15px 16px;border-bottom:1px solid #edf0f4;vertical-align:top}th{background:#f8fafc;color:#536274;font-size:13px}th:nth-child(1){width:18%}th:nth-child(2){width:23%}th:nth-child(3){width:17%}th:nth-child(4){width:42%}.tag{display:inline-block;background:#eaf4fb;color:#25618a;padding:4px 9px;border-radius:20px}.note{line-height:1.6;color:#27384a}
@media(max-width:700px){.shell{margin:20px auto;padding:0 12px}.panel{overflow-x:auto}table{min-width:650px}th,td{padding:12px}}
</style></head><body><main class="shell"><div class="banner">合成验收示意，非真实客户页面</div>
<h1>联系人</h1><p class="sub">联系人列表 · 仅作布局参考</p><section class="panel"><div class="toolbar"><strong>联系人列表</strong><span class="button">＋ 新增联系人</span></div>
<table><thead><tr><th>姓名</th><th>电话</th><th>类型</th><th>备注</th></tr></thead><tbody>
<tr><td>林晓</td><td>138 0000 3476</td><td><span class="tag">合作伙伴</span></td><td class="note">项目交接联系人，工作日可联系。</td></tr>
<tr><td>陈宁</td><td>139 0000 3400</td><td><span class="tag">供应商</span></td><td class="note">负责日常对接；重要事项请同步团队。</td></tr>
<tr><td>周可</td><td>—</td><td><span class="tag">内部</span></td><td class="note">—</td></tr>
</tbody></table></section></main></body></html>'''
    path.write_text(html, encoding="utf-8")


def make_oracle(fixture: dict) -> dict:
    return {
        "fixture_id": fixture["fixture_id"],
        "audience": "操作员专用；不得放入首次模型输入、current-state.docx、reference.html 或 prototype-reference.png。",
        "expected_clarification_answer": {
            "initial_answer": "备注最多120个字符；超出时阻止保存并明确提示，保留已输入内容。",
            "later_change": "之后将长度上限从120改为200个字符，其它规则不变。",
            "counting_details": ["按Unicode码点计数", "不允许换行", "emoji计数", "空格计数"]
        },
        "final_document_oracle": {
            "must_include": ["最终长度上限为200个字符", "超出阻止保存并明确提示", "保留已输入内容", "单行", "空格保留", "emoji和空格均计数"],
            "must_not_include_as_final_rule": ["120个字符"],
            "must_not_invent": ["备注必填", "额外权限", "超长截断", "筛选", "导出", "批量编辑"],
            "must_not_treat_as_requirement": ["澄清讨论过程"]
        },
        "scope_guard": "仅用于合成验收判定，不是联系人产品需求或模型输入材料。"
    }


def build(output: Path) -> list[Path]:
    if output.exists():
        raise FileExistsError(f"输出目录必须不存在：{output}")
    fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
    output.mkdir(parents=True)
    products = {
        "current-state.docx": lambda p: make_docx(p, fixture),
        "reference.html": make_reference_html,
    }
    for name, creator in products.items():
        creator(output / name)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1280, "height": 900}, device_scale_factor=1)
        page.goto((output / "reference.html").resolve().as_uri(), wait_until="load")
        page.screenshot(path=str(output / "prototype-reference.png"), full_page=True)
        browser.close()
    intake = {key: fixture[key] for key in ("fixture_id", "label", "current_state", "initial_request")}
    write_json(output / "intake.json", intake)
    write_json(output / "operator-oracle.json", make_oracle(fixture))
    names = ["current-state.docx", "reference.html", "prototype-reference.png", "intake.json", "operator-oracle.json"]
    manifest = {
        "fixture_id": fixture["fixture_id"],
        "label": fixture["label"],
        "generated_on": date.today().isoformat(),
        "files": [{"path": name, "sha256": sha256(output / name)} for name in names],
        "manifest_note": "清单记录其余五个交付文件的SHA-256；不对清单自身做递归哈希。"
    }
    write_json(output / "manifest.json", manifest)
    return [output / name for name in names + ["manifest.json"]]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path, help="必须是尚不存在的输出目录")
    args = parser.parse_args()
    try:
        files = build(args.output.resolve())
    except (FileExistsError, OSError, ValueError, RuntimeError) as exc:
        print(f"生成失败：{exc}", file=sys.stderr)
        return 2
    for path in files:
        print(f"{path}  sha256={sha256(path)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
