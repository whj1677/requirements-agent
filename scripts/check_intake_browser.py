"""Built-UI + real ASGI ingestion/actions, with isolated synthetic model/data.

Ephemeral loopback test listener only; no daily data or external model calls.
Run with .venv Python; --dist must point to the candidate Vite build.
"""
import argparse
import asyncio
import io
import json
import mimetypes
import os
from pathlib import Path
import random
import sys
import tempfile
import socket
import threading
import uvicorn
from starlette.staticfiles import StaticFiles
import time
from urllib.parse import urlsplit

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
isolated = Path(tempfile.mkdtemp(prefix='ra-intake-browser-'))
os.environ['RA_DATA_DIR'] = str(isolated / 'data')
os.environ.pop('RA_DEEPSEEK_API_KEY', None)
import app.config
app.config.ROOT = isolated
import app.main as main
from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
from PIL import Image
from playwright.async_api import async_playwright, expect
from tests.ui02_fixture import model_server


async def run(dist, evidence):
    evidence.mkdir(parents=True, exist_ok=True)
    checks, errors = [], []
    completed = False
    server = thread = sock = None
    original_save = main.save_source

    def slow_save(*args, **kwargs):
        time.sleep(5)  # Controlled parser latency to inspect real XHR/reading UI.
        return original_save(*args, **kwargs)

    main.save_source = slow_save
    doc = Document()
    doc.add_paragraph('合成验收：联系人列表增加选填备注；只读角色不能修改。')
    for n in range(4):
        image = io.BytesIO()
        Image.frombytes('RGB', (480, 320), random.Random(n).randbytes(480 * 320 * 3)).save(image, format='PNG')
        doc.add_picture(io.BytesIO(image.getvalue()), width=Inches(4))
    file = io.BytesIO(); doc.save(file)
    (evidence / 'synthetic-intake.docx').write_bytes(file.getvalue())
    try:
        with model_server() as model:
            model['delay'] = .8
            application = main.create_app(isolated / 'case', access_token='off', env_path=isolated / 'absent.env')
            with TestClient(application) as client:
                csrf = client.post('/api/session/login', json={'key': 'off'}).json()['csrf']
                client.headers.update({'Origin': 'http://testserver', 'X-CSRF-Token': csrf})
                for slot in ('model', 'vision'):
                    assert client.put('/api/models/' + slot, json=model['config']).status_code == 200
                    assert client.put('/api/models/' + slot + '/key', json={'key': 'synthetic-local-key'}).status_code == 200
                project = client.post('/api/projects', json={'name': '合成导入与分析验收'}).json()
                root = '/api/projects/' + project['id']
                intake = {'platform': '合成平台：现有联系人列表', 'text': '合成诉求：联系人增加选填备注', 'preserved': '合成保持项：只读角色不能修改'}
                other = client.post('/api/projects', json={'name': '合成仅资料项目'}).json()
                assert client.post('/api/projects/' + other['id'] + '/sources/text', json={'text': '合成资料：联系人列表增加选填备注。', 'purpose': 'goal', 'expected_revision': other['revision']}).status_code == 200
                application.router.routes[:] = [r for r in application.router.routes if not (getattr(r, 'path', None) == '' and getattr(r, 'name', None) == 'web')]
                application.mount('/', StaticFiles(directory=dist, html=True), name='isolated-web')
                sock = socket.socket(); sock.bind(('127.0.0.1', 0))
                port = sock.getsockname()[1]
                server = uvicorn.Server(uvicorn.Config(application, log_level='error', access_log=False))
                thread = threading.Thread(target=lambda: server.run(sockets=[sock]), daemon=True); thread.start()
                for _ in range(100):
                    if server.started: break
                    await asyncio.sleep(.05)
                assert server.started, 'isolated test server did not start'
                async with async_playwright() as pw:
                    browser = await pw.chromium.launch(headless=True)
                    page = await browser.new_page(viewport={'width': 1440, 'height': 1000})
                    page.on('pageerror', lambda e: errors.append(str(e)))
                    intake_requests = []
                    page.on('request', lambda req: intake_requests.append(req.post_data_json) if req.method == 'POST' and urlsplit(req.url).path == root + '/intake' else None)

                    await page.goto(f'http://127.0.0.1:{port}/')
                    await page.get_by_label('本机访问令牌').fill('off')
                    await page.get_by_role('button', name='进入工作台').click()
                    await page.get_by_role('button', name='合成导入与分析验收', exact=True).click()
                    workspace = page.locator('.project-workspace:not([hidden])')
                    fields = {'platform': workspace.get_by_placeholder('例如：目前使用的平台、希望改动的页面…'), 'text': workspace.get_by_label('核心诉求', exact=True), 'preserved': workspace.get_by_placeholder('哪些内容保持原状，哪些暂时不做…')}
                    async def assert_intake(*, saved=False):
                        for key, value in intake.items():
                            await expect(fields[key]).to_have_value(value)
                        if saved:
                            snapshot = client.get(root).json()['intake']
                            assert {key: snapshot[key] for key in intake} == intake
                    for key, value in intake.items():
                        await fields[key].fill(value)
                    await page.get_by_role('button', name='项目资料 · 0', exact=True).click()
                    drawer = page.get_by_role('dialog', name='项目资料', exact=True)
                    await page.screenshot(path=str(evidence / 'import-desktop.png'))
                    await drawer.locator('input[type=file]').set_input_files({
                        'name': 'synthetic-intake.docx', 'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                        'buffer': file.getvalue()})
                    await expect(drawer.get_by_role('progressbar', name='文档读取进度')).to_be_visible()
                    await expect(drawer.get_by_role('button', name='关闭项目资料')).to_be_disabled()
                    await page.keyboard.press('Escape')
                    await expect(drawer).to_be_visible()
                    await page.screenshot(path=str(evidence / 'import-reading.png'))
                    checks.append('real XHR upload enters visible reading progress; close/Escape guarded')
                    await expect(drawer.get_by_role('button', name='关闭项目资料')).to_be_enabled(timeout=15000)
                    p = client.get(root).json()
                    assert len([s for s in p['sources'] if s.get('image_mime')]) == 4
                    assert any('选填备注' in e['text'] for s in p['sources'] for e in s['excerpts'])
                    assert not model['requests']
                    checks.append('Word text and 4 embedded images extracted; import makes no model requests')
                    await page.screenshot(path=str(evidence / 'import-result.png'))
                    await drawer.locator('input[type=file]').set_input_files({'name':'second.txt','mimeType':'text/plain','buffer':'合成追加说明：未知规则必须澄清。'.encode()})
                    await expect(drawer.get_by_role('progressbar', name='文档读取进度')).to_be_visible()
                    await expect(drawer.get_by_role('button', name='关闭项目资料')).to_be_enabled(timeout=15000)
                    assert len(client.get(root).json()['sources']) == 6
                    checks.append('selecting a second file directly after completion imports it once')
                    await assert_intake()
                    await drawer.get_by_role('button', name='关闭项目资料').click()
                    guard = page.get_by_role('dialog', name='未保存的修改', exact=True)
                    await expect(drawer).to_be_hidden()
                    await expect(guard).to_be_hidden()
                    await assert_intake()
                    assert not intake_requests and not client.get(root).json().get('intake')
                    checks.append('closing imported materials preserves all three unsaved intake fields without a global discard prompt')
                    await page.get_by_role('button', name='项目资料 ·', exact=False).click()
                    await drawer.get_by_text('粘贴文字或添加公开网页（辅助方式）', exact=True).click()
                    await drawer.get_by_label('粘贴材料').fill('合成资料草稿，关闭窗口也保留')
                    await drawer.get_by_label('公开网页 URL').fill('https://example.invalid/not-sent')
                    await page.keyboard.press('Escape')
                    await expect(drawer).to_be_hidden()
                    await expect(guard).to_be_hidden()
                    await assert_intake()
                    # A real project change still protects BOTH retained drafts.
                    await page.get_by_role('button', name='合成仅资料项目', exact=True).click()
                    await expect(guard).to_contain_text('第1步编辑')
                    await expect(guard).to_contain_text('项目资料')
                    await guard.get_by_role('button', name='留在当前页', exact=True).click()
                    await assert_intake()
                    await page.get_by_role('button', name='项目资料 ·', exact=False).click()
                    await expect(drawer.get_by_label('粘贴材料')).to_have_value('合成资料草稿，关闭窗口也保留')
                    await expect(drawer.get_by_label('公开网页 URL')).to_have_value('https://example.invalid/not-sent')
                    await drawer.get_by_label('粘贴材料').fill('')
                    await drawer.get_by_label('公开网页 URL').fill('')
                    await drawer.get_by_role('button', name='关闭项目资料').click()
                    await expect(drawer).to_be_hidden()
                    checks.append('Escape/reopen preserves intake and material drafts; actual project navigation still protects both')
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    plan = page.get_by_role('dialog', name='核对本次任务')
                    await expect(plan).to_be_visible(timeout=15000)
                    await assert_intake(saved=True)
                    assert len(intake_requests) == 1 and {key: intake_requests[0][key] for key in intake} == intake
                    await plan.get_by_role('button', name='返回保留输入', exact=True).click()
                    await assert_intake(saved=True)
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    await expect(plan).to_be_visible(timeout=20000)
                    assert len(intake_requests) == 1
                    checks.append('analyze posts the exact three fields once; cancel/reopen retains saved input without model requests')
                    assert not model['requests']
                    await plan.get_by_role('button', name='授权本次范围并运行').click()
                    await expect(page.get_by_role('progressbar', name='分析步骤进度')).to_be_visible()
                    await expect(page.get_by_role('button', name='正在分析现状与诉求…')).to_be_disabled()
                    await assert_intake(saved=True)
                    await page.screenshot(path=str(evidence / 'analysis-running.png'))
                    checks.append('explicit authorization before model; running status by analyze button; duplicate start disabled')
                    await expect(page.get_by_role('button', name='查看分析结果', exact=True)).to_be_visible(timeout=30000)
                    task = client.get(root + '/actions').json()[-1]
                    assert task['status'] == 'partial' and task['completed_steps'] == 3, task
                    assert task['calls'] == 3 and len(model['requests']) == 3
                    assert any(m['stage'] == 'ingest' and m['role'] == 'assistant' for m in client.get(root).json()['messages'])
                    await page.screenshot(path=str(evidence / 'analysis-complete.png'))
                    await assert_intake(saved=True)
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '0')
                    checks.append('large-image Word -> 2 vision batches -> ingest; partial result stays on intake with saved text, 3 synthetic requests')
                    await page.get_by_role('button', name='提供现状与诉求', exact=False).click()
                    await expect(page.locator('.workflow-page:not([hidden])')).to_have_attribute('data-phase', '0')
                    await page.wait_for_timeout(1200)
                    await expect(page.locator('.workflow-page:not([hidden])')).to_have_attribute('data-phase', '0')
                    await assert_intake(saved=True)
                    checks.append('partial analysis stays on intake until explicitly opening the result')

                    # Create realistic missing scope facts in the isolated store so the missing-field shortcuts render.
                    current = client.get(root).json()
                    with application.state.store.edit(project['id'], current['revision'], 'Synthetic missing scope facts') as (draft, _):
                        for key in ('value', 'priority'):
                            draft.setdefault('product_context_proposal', {})[key] = ''
                            draft.setdefault('product_context', {})[key] = ''
                    await page.reload()
                    await page.get_by_role('button', name='合成导入与分析验收', exact=True).click()
                    await page.get_by_role('button', name='核对现状、价值与改动范围', exact=False).click()
                    work = page.locator('.project-workspace:not([hidden])')
                    await expect(work.locator('.workflow-page:not([hidden])')).to_have_attribute('data-phase', '1')
                    missing_value = work.get_by_role('button', name='去填写问题与改动价值 →', exact=True)
                    missing_priority = work.get_by_role('button', name='去填写优先级与依据 →', exact=True)
                    await expect(missing_value).to_be_visible()
                    await expect(missing_priority).to_be_visible()
                    await page.set_viewport_size({'width': 390, 'height': 844})
                    await missing_value.click()
                    value_field = work.get_by_label('问题与改动价值', exact=True)
                    await expect(value_field).to_be_focused()
                    box = await value_field.bounding_box()
                    assert box and 0 <= box['y'] < 844 and box['y'] + box['height'] <= 844, box
                    assert await page.get_by_role('dialog', name='未保存的修改', exact=True).count() == 0
                    checks.append('missing value shortcut opens and keyboard-focuses its field in the narrow viewport without a global dirty guard')
                    await value_field.fill('合成价值：减少重复录入')
                    await expect(page.get_by_role('dialog', name='未保存的修改', exact=True)).to_have_count(0)
                    await work.get_by_role('button', name='去填写优先级与依据 →', exact=True).click()
                    priority_field = work.get_by_label('优先级与依据', exact=True)
                    await expect(priority_field).to_be_focused()
                    box = await priority_field.bounding_box()
                    assert box and 0 <= box['y'] < 844 and box['y'] + box['height'] <= 844, box
                    assert await page.get_by_role('dialog', name='未保存的修改', exact=True).count() == 0
                    await priority_field.fill('合成依据：先满足高频操作')
                    await expect(page.get_by_role('dialog', name='未保存的修改', exact=True)).to_have_count(0)
                    await work.get_by_role('button', name='保存补充信息', exact=True).click()
                    await expect(work.get_by_role('button', name='去填写问题与改动价值 →', exact=True)).to_have_count(0)
                    await expect(work.get_by_role('button', name='去填写优先级与依据 →', exact=True)).to_have_count(0)
                    saved_context = client.get(root).json()['product_context']
                    assert saved_context['value'] == '合成价值：减少重复录入', saved_context
                    assert saved_context['priority'] == '合成依据：先满足高频操作', saved_context
                    checks.append('value and priority shortcuts share one facts editor; saving posts both values and removes missing-field actions')

                    await page.screenshot(path=str(evidence / 'scope-fields-mobile.png'))
                    await page.set_viewport_size({'width': 1440, 'height': 1000})
                    # A genuine provider failure must remain visible after budgets are removed.
                    model['fail_next'] = True
                    for s in client.get(root).json()['sources']:
                        if s.get('image_mime'):
                            with application.state.store.edit(project['id'], client.get(root).json()['revision'], 'Synthetic failure observation') as (draft, _):
                                source = next(x for x in draft['sources'] if x['id'] == s['id'])
                                source.pop('vision_run_id', None)
                                source['parse_status'] = 'awaiting_vision'
                            break
                    await page.reload()
                    await page.get_by_role('button', name='合成导入与分析验收', exact=True).click()
                    await page.get_by_role('button', name='提供现状与诉求', exact=False).click()
                    await assert_intake(saved=True)
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    await page.get_by_role('dialog', name='核对本次任务').get_by_role('button', name='运行本次任务', exact=True).click()
                    await expect(page.get_by_text('本次已发出 1 次模型请求。', exact=False)).to_be_visible(timeout=15000)
                    status = page.locator('.step-check-0 .task-status')
                    await expect(status).to_contain_text('HTTP 401')
                    await expect(page.locator('.workflow-page:not([hidden])')).to_have_attribute('data-phase', '0')
                    assert await status.locator('details').get_attribute('open') is None
                    assert len(model['requests']) == 4
                    await assert_intake(saved=True)
                    checks.append('actual provider failure visible beside button with details closed; inputs preserved')
                    await page.screenshot(path=str(evidence / 'analysis-paused.png'))
                    await page.set_viewport_size({'width': 390, 'height': 844})
                    await page.get_by_role('button', name='项目资料 ·', exact=False).click()
                    await page.screenshot(path=str(evidence / 'import-mobile.png'))
                    await drawer.get_by_role('button', name='关闭项目资料').click()
                    await page.set_viewport_size({'width': 1440, 'height': 1000})
                    intake['platform'] += ' / 再编辑'
                    await fields['platform'].fill(intake['platform'])
                    await page.get_by_role('button', name='合成仅资料项目', exact=True).click()
                    await expect(guard).to_be_visible()
                    await guard.get_by_role('button', name='保存后离开', exact=True).click()
                    await expect(guard).to_be_hidden(timeout=15000)
                    await page.get_by_role('button', name='合成导入与分析验收', exact=True).click()
                    await assert_intake(saved=True)
                    assert len(intake_requests) == 2 and {key: intake_requests[-1][key] for key in intake} == intake
                    checks.append('saving on actual project departure retains all three fields after returning')
                    await page.get_by_role('button', name='合成仅资料项目', exact=True).click()
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    await expect(plan).to_be_visible(timeout=20000)
                    await plan.get_by_role('button', name='返回保留输入', exact=True).click()
                    for field in fields.values():
                        await expect(field).to_have_value('')
                    await expect(page.get_by_text('已保存资料范围，未填写文字', exact=True)).to_be_visible()
                    assert len(model['requests']) == 4
                    checks.append('materials-only analysis remains available and accurately says no text was entered')
                    # Pure text has no partial image read: verify a genuine succeeded action.
                    for key, value in intake.items():
                        await fields[key].fill(value)
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    await expect(plan).to_be_visible(timeout=20000)
                    await plan.get_by_role('button', name='授权本次范围并运行').click()
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '1', timeout=30000)
                    pure_root = '/api/projects/' + other['id']
                    pure = client.get(pure_root).json()
                    pure_task = client.get(pure_root + '/actions').json()[-1]
                    assert pure_task['status'] == 'succeeded' and pure_task['calls'] == 1, pure_task
                    assert {key: pure['intake'][key] for key in intake} == intake
                    assert '2' not in pure.get('stage_checks', {}), 'automatic navigation must not confirm scope'
                    assert len(model['requests']) == 5
                    await page.screenshot(path=str(evidence / 'automatic-second-step.png'))
                    await page.get_by_role('button', name='提供现状与诉求', exact=False).click()
                    await page.wait_for_timeout(1200)
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '0')
                    for key, value in intake.items():
                        await expect(fields[key]).to_have_value(value)
                    checks.append('successful pure-text analysis automatically opens step two with saved intake; no automatic scope approval or repeated navigation')

                    # A submitted job must not steal navigation after the user switches projects.
                    model['delay'] = 2
                    await page.get_by_role('button', name='✦ 分析现状与诉求 →', exact=True).click()
                    await plan.get_by_role('button', name='运行本次任务', exact=True).click()
                    await expect(page.get_by_role('button', name='正在分析现状与诉求…')).to_be_disabled()
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '0')
                    await page.get_by_role('button', name='合成导入与分析验收', exact=True).click()
                    for _ in range(150):
                        latest = client.get(pure_root + '/actions').json()[-1]
                        if latest['status'] not in ('queued', 'running'):
                            break
                        await asyncio.sleep(.1)
                    assert latest['status'] == 'succeeded', latest
                    await expect(workspace.get_by_role('heading', name='合成导入与分析验收', exact=True)).to_be_visible()
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '0')
                    await page.get_by_role('button', name='合成仅资料项目', exact=True).click()
                    await page.wait_for_timeout(1200)
                    await expect(workspace.locator('.workflow-page')).to_have_attribute('data-phase', '0')
                    assert len(model['requests']) == 6
                    checks.append('running task stays on intake; completed hidden-project task neither steals the current project nor auto-jumps on return')
                    assert not errors, errors
                    completed = True
                    await browser.close()
    finally:
        main.save_source = original_save
        if server: server.should_exit = True
        if thread: await asyncio.to_thread(thread.join, 5)
        if sock: sock.close()
        (evidence / 'browser-checks.json').write_text(json.dumps(dict(
            status='VERIFIED' if completed else 'FAILED',
            mode='built UI + actual isolated loopback HTTP app/model; ephemeral listener closed; no daily data',
            checks=checks, page_errors=errors, isolated_path=str(isolated)), ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps({'checks': len(checks), 'page_errors': errors}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--dist', type=Path, required=True)
    parser.add_argument('--evidence', type=Path, required=True)
    args = parser.parse_args()
    asyncio.run(run(args.dist.resolve(), args.evidence.resolve()))
